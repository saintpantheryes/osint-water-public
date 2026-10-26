"""阶段 D 第二步:生成报告(A14)与国别信息(A11)。用 test_jobs 的内存版 Teable 环境;Word 转换与世行、IMF 抓取换成假的,不联网。"""
from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from core import tstore
from core.jobs import countries, reports
from tests.test_jobs import TODAY, env  # noqa: F401  fixture

TZ = ZoneInfo("Asia/Shanghai")


def _rows(base, oid):
    return base.rows(oid)


def _val(base, oid, rec, prop):
    return base.get(oid, rec, prop)


def _fake_docx(md_path: Path, docx_path: Path, template=None) -> Path:
    docx_path.write_bytes(b"PK fake docx " + md_path.read_bytes()[:200])
    return docx_path


def _seed(base):
    """一个国别、两条线索、一个跟踪中的商机(带所属线索、时间节点、跟进)。"""
    c = base.rows("country")[0]
    o = base.create("opportunity", [{"code": "AZE-001", "name_zh": "甘贾污水厂", "country": c["id"], "stage": "跟踪中", "public_stage": "采购进行中",
                                     "financing": "亚行 · 50,000,000 USD [aaaaaaaa]"}])[0]
    l1, l2 = base.create("lead", [
        {"key": "aaaaaaaa", "title": "Ganja WWTP EOI", "url": "https://example.org/a", "relevant": True, "category": "意向书", "country": c["id"],
         "found_at": "2026-10-01T02:00:00.000Z", "opportunity": o["id"], "summary_zh": "甘贾污水厂意向书", "origin": "机器采集"},
        {"key": "bbbbbbbb", "title": "Baku water supply tender", "url": "https://example.org/b", "relevant": True, "category": "意向书",
         "country": c["id"], "found_at": "2026-10-03T02:00:00.000Z", "summary_zh": "巴库供水招标", "origin": "机器采集"},
    ])
    base.create("milestone", [{"opportunity": o["id"], "lead": l1["id"], "kind": "意向书截止", "day": "2026-10-20", "日期原文": "2026-10-20", "status": "官方计划"}])
    base.create("followup", [{"opportunity": o["id"], "day": "2026-10-02", "content": "阶段从「建议」改为「跟踪中」", "stage_change": "跟踪中"}])
    return c, o, l1, l2


def test_weekly_report_is_generated_uploaded_and_cited(env, monkeypatch):  # noqa: F811
    fake, admin, abase, ctx = env
    monkeypatch.setattr(reports, "_to_docx", _fake_docx)
    c, o, l1, l2 = _seed(abase)
    # 周一 09:00 自建上周周报;再跑一次不重复建
    now = datetime(2026, 10, 5, 9, 1, tzinfo=TZ)  # 2026-10-05 是周一
    assert reports.ensure_periodic(ctx, now) == {"自建报告": 1}  # 只有周报;招标监测不自动出,人点「生成」才有(10-16 定)
    assert reports.ensure_periodic(ctx, now) == {}
    by_kind = {_val(abase, "report", r, "kind"): r for r in _rows(abase, "report")}
    assert "招标监测" not in by_kind
    r = by_kind["周报"]
    assert _val(abase, "report", r, "period") == "2026-W40" and _val(abase, "report", r, "gen_status") == "待生成"
    assert reports.run(ctx) == {"报告完成": 1, "报告失败": 0}
    r = abase.record("report", r["id"])
    assert _val(abase, "report", r, "gen_status") == "完成"
    name = _val(abase, "report", r, "key")
    assert name.startswith("周报_") and name.endswith(".docx")
    files = _val(abase, "report", r, "file")
    assert files and files[0]["name"] == name
    # 引用的线索 = 正文列出的(本期两条)∪ 写进的商机的所属线索;写进的商机 = 阶段有变化的
    assert {x["id"] for x in _val(abase, "report", r, "cites_leads")} == {l1["id"], l2["id"]}
    assert [x["id"] for x in _val(abase, "report", r, "cites_opps")] == [o["id"]]
    md = (ctx.cache.root.parent / "reports" / name.replace(".docx", ".md")).read_text(encoding="utf-8")
    assert "aaaaaaaa" in md and "bbbbbbbb" in md and "AZE-001" in md and "意向书截止" in md and "本报告未使用模型" in md
    # 线索的派生属性「写进的报告」跟着有了
    assert _val(abase, "lead", abase.record("lead", l1["id"]), "in_reports")


def test_country_and_project_reports_and_failures(env, monkeypatch):  # noqa: F811
    fake, admin, abase, ctx = env
    monkeypatch.setattr(reports, "_to_docx", _fake_docx)
    c, o, l1, l2 = _seed(abase)
    abase.create("report", [{"kind": "国别报告", "country": c["id"], "day": "2026-10-05", "gen_status": "待生成"},
                            {"kind": "项目报告", "opportunity": o["id"], "day": "2026-10-05", "gen_status": "待生成"},
                            {"kind": "国别报告", "day": "2026-10-05", "gen_status": "待生成"}])  # 没选国别:失败
    assert reports.run(ctx) == {"报告完成": 2, "报告失败": 1}
    by = {_val(abase, "report", r, "kind") + str(_val(abase, "report", r, "gen_status")): r for r in _rows(abase, "report")}
    cr, pr = by["国别报告完成"], by["项目报告完成"]
    assert "国别报告失败" in by
    assert _val(abase, "report", cr, "key").startswith("国别报告_AZE_")
    assert {x["id"] for x in _val(abase, "report", cr, "cites_leads")} == {l1["id"]}  # 国别报告不列线索,只记商机的出处线索
    assert {x["id"] for x in _val(abase, "report", pr, "cites_leads")} == {l1["id"]}
    assert [x["id"] for x in _val(abase, "report", pr, "cites_opps")] == [o["id"]]
    md = (ctx.cache.root.parent / "reports" / _val(abase, "report", pr, "key").replace(".docx", ".md")).read_text(encoding="utf-8")
    assert "甘贾污水厂" in md and "一、核心结论" in md and "七、信息来源" in md and "Ganja WWTP EOI" in md  # 按《项目分析报告》模板
    # 重新生成:把完成的改回待生成,同一行再出一版(文件名不变、不新建行)
    abase.update("report", pr["id"], {"gen_status": "待生成"})
    assert reports.run(ctx) == {"报告完成": 1, "报告失败": 0}
    assert len(_rows(abase, "report")) == 3 and len(_val(abase, "report", abase.record("report", pr["id"]), "file")) == 2


def test_stale_generating_reports_are_reset(env):  # noqa: F811
    fake, admin, abase, ctx = env
    r = abase.create("report", [{"kind": "周报", "day": "2026-10-05", "gen_status": "生成中"}])[0]
    fake.records[abase.t["report"]][0]["lastModifiedTime"] = "2026-01-01T00:00:00.000Z"
    assert reports.reset_stale(ctx) == 1
    assert _val(abase, "report", abase.record("report", r["id"]), "gen_status") == "待生成"


class _V:
    def __init__(self, label, text, source, url):
        self.metric = type("M", (), {"label": label})()
        self._t, self.source, self.url = text, source, url

    def text(self, currency=""):
        return self._t


class _F:
    """load_facts 的最小替身。"""
    def __init__(self, iso):
        self.iso3 = iso
        self.values = {"gdp": _V("GDP", "720 亿美元(2024)", "WB-WDI", "https://data.worldbank.org/x")}
        self.meta = {"income": "Upper middle income"}
        self.market = None
        self.water = {"institutions": [{"name_zh": "阿塞拜疆水务局", "url": "https://water.gov.az"}], "enterprises": [], "laws": []}
        self.currency_code = "AZN"

    def income(self):
        return "中高收入"

    def lending(self):
        return "IBRD"

    def basics_rows(self, internal=True):
        return [("首都", "巴库", "维基", "")]


def test_country_info_fills_replaces_and_appends(env, monkeypatch):  # noqa: F811
    fake, admin, abase, ctx = env
    from core.country_facts import METRICS
    monkeypatch.setattr(countries, "_fetch", lambda ctx, d: False)
    monkeypatch.setattr(countries, "_facts", lambda d, m, isos: {i: _F(i) for i in isos})
    monkeypatch.setattr(countries, "METRICS", tuple(m for m in METRICS if m.key == "gdp"))
    monkeypatch.setattr(countries, "_basics", lambda: {"AZE": {"wiki_zh": "阿塞拜疆", "wiki_en": "Azerbaijan"}})
    c = abase.rows("country")[0]
    abase.update("country", c["id"], {"basics": "首都:巴库(维基)", "water_agencies": "旧机构 · https://old.example"})
    out = countries.run(ctx, refresh=True)
    assert out["国别"] == 1 and out["替换"] == 2 and out["追加"] == 1 and out["维基链接"] == 2  # 指标与收入组替换;机构追加;基础信息已有同一条不重复;维基百科链接填上
    c = abase.record("country", c["id"])
    assert _val(abase, "country", c, "region") and _val(abase, "country", c, "tier") == "重点"  # 人管的不动
    assert "GDP:720 亿美元(2024)(世行 WDI https://data.worldbank.org/x)" in _val(abase, "country", c, "indicators")
    assert _val(abase, "country", c, "income_group") == "中高收入;IBRD"
    assert _val(abase, "country", c, "water_agencies") == "旧机构 · https://old.example\n阿塞拜疆水务局 · https://water.gov.az"
    assert _val(abase, "country", c, "basics") == "首都:巴库(维基)"
    assert _val(abase, "country", c, "wiki_zh").endswith("%E7%96%86") and _val(abase, "country", c, "wiki_en") == "https://en.wikipedia.org/wiki/Azerbaijan"  # 维基百科链接(2026-10-26)
    # 只补空着的:都填了就什么也不做
    assert countries.run(ctx, only_missing=True) == {}
    # 新国别出现(线索建了国别行)→ only_missing 补它
    ctx.country_rid("MAR")
    out = countries.run(ctx, only_missing=True)
    assert out["国别"] == 1
    m = next(r for r in abase.rows("country") if _val(abase, "country", r, "key") == "MAR")
    assert _val(abase, "country", m, "name_zh") == "摩洛哥" and _val(abase, "country", m, "indicators") and not _val(abase, "country", m, "wiki_zh")  # 资料里没记条目名的不填搜索链接
    # 维基百科链接每次运行都只填空(不等每月 1 日):资料里补了摩洛哥的条目名 → 下次只补它,别的不动
    monkeypatch.setattr(countries, "_basics", lambda: {"AZE": {"wiki_zh": "阿塞拜疆", "wiki_en": "Azerbaijan"}, "MAR": {"wiki_zh": "摩洛哥", "wiki_en": "Morocco"}})
    assert countries.run(ctx, only_missing=True) == {"维基链接": 2}
    assert _val(abase, "country", abase.record("country", m["id"]), "wiki_zh") == "https://zh.wikipedia.org/zh-cn/%E6%91%A9%E6%B4%9B%E5%93%A5"


def test_report_uses_the_default_template_from_the_template_table(env, monkeypatch):  # noqa: F811
    """「报告模板」表里适用类型含该类型且勾了缺省的模板 → 下载到本机、转 Word 时套它;报告行记下实际用的模板。没有合适的 → 程序自带模板。"""
    fake, admin, abase, ctx = env
    used: list = []

    def fake_docx(md_path: Path, docx_path: Path, template=None) -> Path:
        used.append(template)
        return _fake_docx(md_path, docx_path)

    monkeypatch.setattr(reports, "_to_docx", fake_docx)
    monkeypatch.setattr(reports, "_download", lambda ctx, att: b"PK template " + att["name"].encode())
    _seed(abase)
    t_other = abase.create("template", [{"name": "只给月报", "kinds": ["月报"], "is_default": True}])[0]
    admin.upload_attachment(abase.t["template"], t_other["id"], abase.fid("template", "file"), "monthly.docx", b"PK m")
    t1 = abase.create("template", [{"name": "集团版式", "kinds": ["周报", "国别报告"], "is_default": True}])[0]
    admin.upload_attachment(abase.t["template"], t1["id"], abase.fid("template", "file"), "group.docx", b"PK g")
    # 没有文件的缺省模板不算
    abase.create("template", [{"name": "空的", "kinds": ["周报"], "is_default": True}])
    r = abase.create("report", [{"kind": "周报", "period": "2026-W40", "day": "2026-10-05", "edition": "内部版", "gen_status": "待生成"}])[0]
    assert reports.run(ctx) == {"报告完成": 1, "报告失败": 0}
    assert len(used) == 1 and used[0] is not None and used[0].suffix == ".docx" and used[0].read_bytes() == b"PK template group.docx"
    r = abase.record("report", r["id"])
    assert _val(abase, "report", r, "gen_status") == "完成" and tstore.link_id(_val(abase, "report", r, "template")) == t1["id"]
    # 报告行自己选了模板 → 用选的;项目报告没有合适的模板 → 自带(template=None)
    o = abase.rows("opportunity")[0]
    r2 = abase.create("report", [{"kind": "月报", "period": "2026-09", "day": "2026-10-05", "edition": "内部版", "gen_status": "待生成", "template": t1["id"]},
                                 {"kind": "项目报告", "opportunity": o["id"], "day": "2026-10-05", "edition": "内部版", "gen_status": "待生成"}])
    assert reports.run(ctx) == {"报告完成": 2, "报告失败": 0}
    assert len(used) == 2 and used[1].read_bytes() == b"PK template group.docx"  # 项目报告不走 pandoc,填自带的项目分析报告模板
    assert tstore.link_id(_val(abase, "report", abase.record("report", r2[0]["id"]), "template")) == t1["id"]
    assert _val(abase, "report", abase.record("report", r2[1]["id"]), "template") is None


def test_daily_brief_is_built_after_collection_and_shown_as_summary(env, monkeypatch):  # noqa: F811
    """日报(E10):采集完成后自建当天一行(同一天只建一次);生成时正文写进「简报正文」并传 Word;口径与首页一致。"""
    from core.jobs import brief
    fake, admin, abase, ctx = env
    monkeypatch.setattr(reports, "_to_docx", _fake_docx)
    c, o, l1, l2 = _seed(abase)
    today = TODAY.isoformat()
    abase.create("lead", [
        {"key": "cccccccc", "title": "Sumgait desalination tender", "url": "https://example.org/c", "relevant": True, "category": "招标", "sectors": ["海水淡化"],
         "country": c["id"], "found_at": f"{today}T01:00:00.000Z", "funder": "World Bank (IDA)", "amount": "50,000,000 USD", "status": "待处理", "origin": "机器采集", "summary_zh": "苏姆盖特海水淡化厂招标"},
        {"key": "dddddddd", "title": "Water policy news", "url": "https://example.org/d", "relevant": True, "category": "政策", "sectors": ["供水"],
         "country": c["id"], "found_at": f"{today}T01:10:00.000Z", "status": "待处理", "origin": "机器采集", "suggested_opp": o["id"]},
        {"key": "eeeeeeee", "title": "Unrelated page", "url": "https://example.org/e", "relevant": False, "found_at": f"{today}T01:20:00.000Z", "status": "待处理", "origin": "机器采集"},
    ])
    ctx.source_rid("fake_src", name="假信源", kind="接口")
    ctx.source_status("fake_src", "ok", f"{today}T00:40:00.000Z")
    assert reports.ensure_daily(ctx, TODAY) == {"自建报告": 1}
    assert reports.ensure_daily(ctx, TODAY) == {}  # 同一天不重复建
    rows = [r for r in abase.rows("report") if _val(abase, "report", r, "kind") == "日报"]
    assert len(rows) == 1 and _val(abase, "report", rows[0], "period") == today and _val(abase, "report", rows[0], "gen_status") == "待生成"
    assert reports.run(ctx) == {"报告完成": 1, "报告失败": 0}
    r = abase.record("report", rows[0]["id"])
    assert _val(abase, "report", r, "gen_status") == "完成" and _val(abase, "report", r, "key").startswith("日报_") and _val(abase, "report", r, "key").endswith(".docx")
    md = _val(abase, "report", r, "summary")
    assert md.startswith(f"# 线索日报 {today}")
    # 今日新增 3(相关 2、采购类 1);累计 5;待处理 = 处理状态为待处理的(与首页卡片同一条件);待确认归属 1(dddddddd 有建议所属)
    assert "线索累计 5 条;今日新增 3 条(相关 2 条,其中采购类 1 条)" in md
    assert "待处理 5 条,待确认归属 1 条;跟踪中的商机 1 个" in md  # 处理状态缺省待处理,5 条都是
    assert "资金来源:世界银行 1、未填 1" in md and "领域:海水淡化 1、供水 1" in md
    assert "- 苏姆盖特海水淡化厂招标〔阿塞拜疆 · 阶段:招标 · 领域:海水淡化 · 50,000,000 USD〕 [cccccccc]\n  原题:Sumgait desalination tender" in md  # 中文名在前,标签跟着;原题另起一行
    assert md.index("[cccccccc]") < md.index("[dddddddd]")
    assert "- 要点:政策〔阿塞拜疆 · 阶段:政策 · 领域:供水〕 [dddddddd]\n  原题:Water policy news" in md  # 没有中文名的:从原题挑词成「要点」
    assert "今日采集信源 1 个:正常 1" in md and "未使用模型" in md and "模型研判" not in md
    files = _val(abase, "report", r, "file")
    assert [f["name"] for f in files] == [_val(abase, "report", r, "key")]  # 有「简报正文」字段时只传 Word
    cited = {_val(abase, "lead", abase.record("lead", x["id"]), "key") for x in _val(abase, "report", r, "cites_leads")}
    assert cited == {"cccccccc", "dddddddd"}
    # 配了模型:多一段研判,编造的编号标「未核到」
    class W(type(ctx.clf)):
        label = "测试模型"
        def write(self, system, user, fields):
            return {"insight": "海水淡化厂招标值得跟 [cccccccc];另有一条编号不存在 [12345678]。"}
    ctx.clf = W()
    md2, _, _, _ = brief.build_daily(ctx, TODAY)
    assert "## 模型研判" in md2 and "[cccccccc]" in md2 and "12345678(未核到)" in md2 and "1 个编号在今日数据里没找到" in md2


def test_daily_brief_tells_sections_apart_shows_rmb_and_ai_suggested_opps(env, monkeypatch):  # noqa: F811
    """同一项目的不同标段只有结尾不同:截短时留尾;金额带折合人民币;机器推荐的商机和人立项的分开说。"""
    from datetime import date as _date
    from core.jobs import brief
    fake, admin, abase, ctx = env
    c, o, l1, l2 = _seed(abase)
    opp = abase.create("opportunity", [{"code": "MAR-099", "name_zh": "推荐的", "stage": "建议"}])[0]
    today = _date.fromisoformat(tstore.day(abase.record("opportunity", opp["id"])["createdTime"], abase.tz))  # 新建商机按创建时间算
    head = "TRAVAUX D'ELARGISSEMENT ET MISE A NIVEAU DE LA ROUTE D'OURIKA (RP 2017) - TRAVAUX DE VOIRIE, ASSAINISSEMENT ET GENIE CIVIL"
    abase.create("lead", [{"key": "ffffffff", "title": head + " - TRONÇON 2", "url": "https://example.org/f", "relevant": True, "category": "招标",
                           "country": c["id"], "found_at": f"{today.isoformat()}T01:00:00.000Z", "amount": "51,953,580 MAD", "amount_cny": 3700,
                           "status": "待处理", "origin": "机器采集"}])
    md, _, _, _ = brief.build_daily(ctx, today)
    lines = md.splitlines()
    i = next(k for k, x in enumerate(lines) if "[ffffffff]" in x)
    assert "51,953,580 MAD(约 3,700 万元)" in lines[i] and lines[i + 1].startswith("  原题:") and lines[i + 1].endswith("TRONÇON 2") and "…" in lines[i + 1]
    assert "今日 AI 新推荐商机 1 个(MAR-099 推荐的);新立项跟踪 " in md and "今日新建商机" not in md
    # 配了模型:没有中文名的原题请模型译成中文(有缓存),做标题
    class T:
        enabled, calls, label = True, 0, "测试模型"
        def ask(self, system, user, schema, guide):
            import json as _j
            return {"items": [{"i": x["i"], "zh": "奥里卡公路拓宽第二段"} for x in _j.loads(user)]}
    ctx.clf = T()
    md2, _, _, _ = brief.build_daily(ctx, today)
    assert "- 奥里卡公路拓宽第二段〔阿塞拜疆 · 阶段:招标" in md2


def test_daily_brief_compares_with_median_of_past_week(env, monkeypatch):  # noqa: F811
    """一次批量导入(新信源第一次跑进来上千条)不该让平常的一天看着像「明显低于」:基准是前 7 天的中位数,不含今天。"""
    from datetime import timedelta
    from core.jobs import brief
    fake, admin, abase, ctx = env
    per_day = {1: 20, 2: 13, 3: 3, 4: 300, 5: 9, 6: 28, 7: 20}  # 前 7 天;中位数 20
    rows = [{"key": f"k{i}-{j:03d}", "title": f"t{i}-{j}", "url": f"https://example.org/{i}/{j}", "found_at": f"{(TODAY - timedelta(days=i)).isoformat()}T01:00:00.000Z",
             "status": "待处理", "origin": "机器采集"} for i, n in per_day.items() for j in range(n)]
    rows += [{"key": f"t-{j:03d}", "title": f"today {j}", "url": f"https://example.org/t/{j}", "found_at": f"{TODAY.isoformat()}T01:00:00.000Z",
              "status": "待处理", "origin": "机器采集"} for j in range(16)]
    for i in range(0, len(rows), 100):
        abase.create("lead", rows[i:i + 100])
    md, _, _, _ = brief.build_daily(ctx, TODAY)
    assert "今日新增 16 条" in md and "低于前 7 天的一般水平(每天 20 条,取中位数)" in md and "明显低于" not in md


def test_daily_brief_requested_for_a_past_day_uses_that_day(env, monkeypatch):  # noqa: F811
    fake, admin, abase, ctx = env
    monkeypatch.setattr(reports, "_to_docx", _fake_docx)
    _seed(abase)
    r = abase.create("report", [{"kind": "日报", "period": "2026-10-03", "day": TODAY.isoformat(), "gen_status": "待生成"}])[0]
    assert reports.run(ctx) == {"报告完成": 1, "报告失败": 0}
    r = abase.record("report", r["id"])
    assert "# 线索日报 2026-10-03" in _val(abase, "report", r, "summary") and "今日新增 1 条" in _val(abase, "report", r, "summary")  # bbbbbbbb 是 10-03 发现的
    bad = abase.create("report", [{"kind": "日报", "period": "上周", "day": TODAY.isoformat(), "gen_status": "待生成"}])[0]
    assert reports.run(ctx) == {"报告完成": 0, "报告失败": 1}
    assert _val(abase, "report", abase.record("report", bad["id"]), "gen_status") == "失败"


def test_internal_country_and_project_reports_append_company_profile(env, monkeypatch):  # noqa: F811
    """A23:勾了「写进报告」的公司资料按适用国别与领域附在国别 / 项目报告的内部版后;公开版不附;没勾的不附;归属不是本单位的不附(v2.28)。"""
    fake, admin, abase, ctx = env
    monkeypatch.setattr(reports, "_to_docx", _fake_docx)
    c, o, l1, l2 = _seed(abase)
    abase.update("opportunity", o["id"], {"sectors": ["污水"]})
    other = abase.create("country", [{"name_zh": "格鲁吉亚", "ISO3": "GEO"}])[0]
    abase.create("company", [
        {"kind": "资质", "name": "市政公用工程施工总承包特级", "year": 2020, "for_report": True},                      # 全球通用、各领域
        {"kind": "业绩", "name": "巴库净水厂扩建", "countries": [c["id"]], "sectors": ["供水"], "year": 2023, "amount": "1.2 亿美元", "for_report": True},
        {"kind": "业绩", "name": "第比利斯污水厂", "countries": [other["id"]], "sectors": ["污水"], "year": 2022, "for_report": True},  # 别国
        {"kind": "其他", "name": "公司简介", "content": "中英文简介", "for_report": False},                               # 没勾
        {"kind": "业绩", "name": "兄弟单位的巴库污水厂", "countries": [c["id"]], "sectors": ["污水"], "year": 2021, "for_report": True, "owner": "兄弟单位"},  # 非本单位:只做记录
    ])
    abase.create("report", [{"kind": "国别报告", "country": c["id"], "day": "2026-10-05", "edition": "内部版", "gen_status": "待生成"},
                            {"kind": "国别报告", "country": c["id"], "day": "2026-10-05", "edition": "公开版", "gen_status": "待生成"},
                            {"kind": "项目报告", "opportunity": o["id"], "day": "2026-10-05", "gen_status": "待生成"}])
    assert reports.run(ctx) == {"报告完成": 3, "报告失败": 0}
    mds = {}
    for r in _rows(abase, "report"):
        key = _val(abase, "report", r, "key")
        mds[_val(abase, "report", r, "kind") + str(_val(abase, "report", r, "edition") or "")] = (ctx.cache.root.parent / "reports" / key.replace(".docx", ".md")).read_text(encoding="utf-8")
    internal = mds["国别报告内部版"]
    assert "## 附录 C 我方资质与相关业绩" in internal and "市政公用工程施工总承包特级" in internal and "巴库净水厂扩建" in internal
    assert "第比利斯污水厂" not in internal and "公司简介" not in internal and "兄弟单位的巴库污水厂" not in internal
    assert "我方资质" not in mds["国别报告公开版"]
    project = mds["项目报告"]  # 商机领域污水:供水的业绩不附,全球通用的资质附
    assert "市政公用工程施工总承包特级" in project and "巴库净水厂扩建" not in project and "第比利斯污水厂" not in project and "兄弟单位的巴库污水厂" not in project


class _FakeTranslator:
    enabled, label = True, "fake"

    def __init__(self):
        self.calls = 0

    def ask(self, system, user, schema, guide):
        import json
        self.calls += 1
        return {"items": [{"i": x["i"], "zh": "译:" + x["t"]} for x in json.loads(user)]}

    def write(self, system, user, fields):
        return None


def test_country_report_has_clickable_sources_chinese_titles_and_bank_names(env, monkeypatch):  # noqa: F811
    fake, admin, abase, ctx = env
    monkeypatch.setattr(reports, "_to_docx", _fake_docx)
    c, o, l1, l2 = _seed(abase)
    l3 = abase.create("lead", [{"key": "cccccccc", "title": "Construction of pumping station [Lot 2]", "url": "https://example.org/c", "relevant": True,
                                "category": "意向书", "country": c["id"], "found_at": "2026-10-04T02:00:00.000Z", "origin": "机器采集"}])[0]
    abase.update("country", c["id"], {"basics": "首都:巴库(维基百科 https://zh.wikipedia.org/zh-cn/%E5%B7%B4)\n世行区域 / 贷款类别:欧洲与中亚 / IBRD",
                                      "income_group": "中等偏上收入;IBRD", "water_agencies": "水务公司 · Azersu · https://www.azersu.az/en"})
    tr = _FakeTranslator()
    ctx.clf = tr
    md, *_ = reports.build_country(ctx, abase.get("country", abase.record("country", c["id"]), "key"))
    md = reports.expand_acronyms(md)
    # 基本情况成表,来源名是可点的链接,网址不再印成正文;机构成表,网址变成可点的网站名
    assert "| 首都 | 巴库 [维基百科](<https://zh.wikipedia.org/zh-cn/%E5%B7%B4>) |" in md
    assert "| 水务公司 | Azersu |  | [azersu.az](<https://www.azersu.az/en>) |" in md
    # 国别报告不列线索(使用方 2026-10-10),线索的标题与网址都不出现;出处线索号也去掉
    assert "example.org/a" not in md and "example.org/c" not in md and "[aaaaaaaa]" not in md and tr.calls == 0
    # 线索清单(周报等)标题中文在前、点到原文;没有中文的请模型译(只译一次,存本机缓存)
    rows = str(reports._lead_rows(ctx, [abase.record("lead", l1["id"]), abase.record("lead", l3["id"])]))
    assert "[甘贾污水厂意向书](<https://example.org/a>)(原文:Ganja WWTP EOI)" in rows
    assert r"[译:Construction of pumping station \\[Lot 2\\]](<https://example.org/c>)" in rows and tr.calls == 1
    reports._lead_rows(ctx, [abase.record("lead", l3["id"])])
    assert tr.calls == 1
    # 银行缩写第一次出现时补中文全称,后面不重复
    assert md.count("IBRD(国际复兴开发银行)") == 1 and "IBRD" in md.split("IBRD(国际复兴开发银行)", 1)[1]
    assert l3


def test_word_links_are_real_hyperlinks(tmp_path):
    import pytest
    pytest.importorskip("pypandoc")
    docx = pytest.importorskip("docx")
    md = tmp_path / "r.md"
    md.write_text("# 报告\n\n- 首都:巴库([维基百科](<https://zh.wikipedia.org/wiki/x>))\n\n| 标题 |\n|---|\n| [中文](<https://example.org/a>) |\n", encoding="utf-8")
    from core.watch import reporting
    out = reporting.to_docx_styled(md, tmp_path / "r.docx", template=tmp_path / "none.docx")
    d = docx.Document(str(out))
    rels = [r.target_ref for r in d.part.rels.values() if r.reltype.endswith("/hyperlink")]
    assert "https://zh.wikipedia.org/wiki/x" in rels and "https://example.org/a" in rels
    assert "https://" not in "\n".join(p.text for p in d.paragraphs)


def test_country_report_cover_and_toc_are_separate_pages_with_links(tmp_path, monkeypatch):
    """使用方 10-19:生成稿的目录和封面挤在一页、目录点不动,成稿不是这样。封面一页;「目　录」另起一页,每条是跳到该章的链接,
    行末虚线接页码(本机有 LibreOffice 时算出来写进去,没有留给 Word 打开时更新);摘要再另起一页。"""
    import pytest
    pytest.importorskip("pypandoc")
    docx = pytest.importorskip("docx")
    from docx.oxml.ns import qn
    from core.watch import reporting
    md = tmp_path / "r.md"
    md.write_text("---\ntitle: \"X 水务市场分析报告\"\nsubtitle: \"内部资料　仅供参考\"\ndate: \"2026 年 10 月\"\n---\n\n目　录\n\n# 摘要\n\n要点。\n\n"
                  "# 第 1 章 研究背景\n\n## 1.1 目的\n\n正文。\n\n# 第 2 章 国家概况\n\n正文。\n\n# 附录\n\n## 附录 A 数据来源登记\n\n来源。\n", encoding="utf-8")
    monkeypatch.setattr(reporting, "toc_pages", lambda path: {"_Toc001": 3, "_Toc002": 4, "_Toc003": 5, "_Toc004": 6})  # 像本机装了排版软件
    monkeypatch.setattr(reporting, "_org_of", lambda path: "能源环保事业部")  # 像项目报告模板写了编制单位
    out = reporting.to_docx_styled(md, tmp_path / "r.docx", template=tmp_path / "none.docx")
    d = docx.Document(str(out))
    paras = d.paragraphs
    toc_title = next(p for p in paras if p.text.strip() == "目　录")
    # 封面照成稿(使用方 10-19 对照截图):标题前空 5 行、黑体 26 号加粗居中;底部编制单位(加粗)、年月、版本三行居中;再往下才是目录页
    cover = paras[:paras.index(toc_title)]
    texts = [p.text for p in cover]
    assert texts[:6] == ["", "", "", "", "", "X 水务市场分析报告"] and texts[-3:] == ["能源环保事业部", "2026 年 10 月", "内部资料　仅供参考"]
    title_run = cover[5].runs[0]
    assert title_run.font.size.pt == 26 and title_run.font.bold and cover[5].paragraph_format.alignment == 1
    assert cover[-3].runs[0].font.bold and cover[-3].runs[0].font.size.pt == 12 and cover[-1].runs[0].font.size.pt == 9.5
    assert all(p.paragraph_format.alignment == 1 for p in cover[-3:]) and all(p.text == "" for p in cover[6:-3]) and len(cover) == 6 + 20 + 3
    assert toc_title.paragraph_format.page_break_before and toc_title.style.name != "Heading 1"  # 目录自己不进目录
    assert toc_title.paragraph_format.alignment == 0 and toc_title.runs[0].font.size.pt == 18  # 成稿的「目　录」靠左、18 号加粗
    entries = [p for p in paras if p._p.find(qn("w:hyperlink")) is not None and p._p.find(qn("w:hyperlink")).get(qn("w:anchor"))]
    assert [e.text for e in entries] == ["摘要\t3", "第 1 章 研究背景\t4", "第 2 章 国家概况\t5", "附录\t6"]  # 一级标题各一条(节不列),制表到行末接页码
    anchors = [e._p.find(qn("w:hyperlink")).get(qn("w:anchor")) for e in entries]
    marks = {b.get(qn("w:name")): b for b in d.element.body.iter(qn("w:bookmarkStart"))}
    assert all(a in marks for a in anchors)
    heads = [p for p in paras if p.style.name == "Heading 1"]
    assert [h.text for h in heads] == ["摘要", "第 1 章 研究背景", "第 2 章 国家概况", "附录"] and all(h.paragraph_format.page_break_before for h in heads)
    assert all(h._p.find(qn("w:bookmarkStart")) is not None for h in heads)  # 书签打在标题上,点目录就跳过去
    instr = " ".join(t.text or "" for t in d.element.body.iter(qn("w:instrText")))
    assert 'TOC \\o "1-1"' in instr and "PAGEREF _Toc001" in instr
    assert d.settings.element.find(qn("w:updateFields")) is not None  # Word 打开时更新域,页码跟着重算
    assert paras.index(toc_title) < paras.index(entries[0]) < paras.index(heads[0])


def test_fact_rows_parse_value_year_source_and_tone():
    ind = ("GDP(现价):1,470 亿美元(2025)(世行 WDI https://data.worldbank.org/x)\nGDP 实际增速:-0.3%(2025)(世行 WDI https://data.worldbank.org/y)\n"
           "淡水压力(取水/可再生资源):123.0%(2022)(世行 WDI https://data.worldbank.org/w)")
    rows = reports.fact_rows(ind)
    assert rows[0] == {"k": "GDP(现价)", "v": "1,470 亿美元", "year": "2025", "src": "世行 WDI", "url": "https://data.worldbank.org/x", "bad": None}
    assert rows[1]["bad"] == "负增长" and rows[2]["bad"] == "取水 ≥ 75% 可再生量"
    md = reports.kpi_table(reports._kpis(rows))
    assert '| 1,470 [亿美元]{custom-style="KPI Unit"} | -0.3% |' in md and "Table: 关键指标" in md
    assert "GDP 实际增速(2025) 超警戒线:负增长" in md and "▼" not in md  # 写明越过了哪条线,不用看不懂的符号


class _FakeWriter:
    enabled, label = True, "fake"

    def write(self, system, user, fields):
        return {k: f"〔{k}〕据数据分析 [aaaaaaaa]" for k in fields}

    def ask(self, *a):
        return None


def test_country_report_follows_the_azerbaijan_outline(env, monkeypatch):  # noqa: F811
    from core.jobs import country_report
    fake, admin, abase, ctx = env
    c, o, l1, l2 = _seed(abase)
    abase.update("country", c["id"], {"indicators": "GDP(现价):759 亿美元(2025)(世行 WDI https://data.worldbank.org/x)\n"
                                                    "GDP 实际增速:-0.3%(2025)(世行 WDI https://data.worldbank.org/y)",
                                      "basics": "首都:巴库(维基百科 https://zh.wikipedia.org/x)", "income_group": "中等偏上收入;IBRD"})
    iso = abase.get("country", abase.record("country", c["id"]), "key")
    md, cl, co, scope = reports.build_country(ctx, iso)  # 没配模型:只出数据,分析从略
    assert "本报告未使用模型" in md and "〔" not in md
    titles = [t for t, _ in country_report.OUTLINE]
    heads = re.findall(r"^# 第 \d+ 章 (.+)$", md, re.MULTILINE)
    assert heads and all(h in titles for h in heads) and heads == [t for t in titles if t in heads]  # 章的顺序与成稿一致
    assert md.index("# 摘要") < md.index("# 第 1 章") and "# 附录" in md and "\n目\u3000录\n" in md  # 目录只留标题,转 Word 时做成可点的目录页
    assert '1,470' not in md and '759 [亿美元]{custom-style="KPI Unit"}' in md
    assert "B-READY 2025 这一期没有覆盖" in md  # 营商环境(v2.23):没覆盖的写明
    abase.update("country", c["id"], {"bready_total": 66, "bready_water": 79.3, "bready_procure": 45.7, "water_tariff": 2, "water_tariff_cur": "AZN", "water_tariff_usd": 1.18,
                                      "foreign_bid": "部分招标有限制", "award_days": 365, "gov_pay_days": 7, "bready_edition": "B-READY 2025(数据版本 December 10, 2025)"})
    md, *_ = reports.build_country(ctx, iso)
    sec = md[md.index("外资与营商环境"):]
    assert "总分 66.0 分;需要留意的是公共采购得分、外企参标限制、大型工程授标天数" in sec and "约 1.18 美元/m³(2 AZN/m³)" in sec
    assert "45.7 分(超警戒线:低于 50 分,偏弱)" in sec and "365 天(超警戒线:超过 180 天,授标慢)" in sec and "7 天 |" in sec
    ctx.clf = _FakeWriter()
    md, *_ = reports.build_country(ctx, iso)
    # 配了模型:照成稿 72 节逐节写(字段 s章_节),程序填的几节不请模型;首都、我国换成实际说法;空节不出、编号连续
    assert "〔summary〕" in md and "〔s15_5〕" in md and "〔s11_2〕" in md and "## 2.3 宏观经济" in md and "〔s1_3〕" not in md
    assert "# 第 8 章 巴库及大巴库水务专题" in md and "## 2.6 中国与该国经贸关系与中国企业现状" in md and "## 12.5 中国企业在该国业绩与经验教训" in md
    n_secs = sum(len(secs) for _, secs in country_report.OUTLINE)
    assert len(re.findall(r"^## \d+\.\d+ ", md, re.MULTILINE)) == n_secs - 2  # 只少「首都在建与规划项目清单」(没挂首都的商机)与「国别备注」(没填)
    assert "## 附录 B 项目管道清单" in md and "## 附录 E 术语与缩略语" in md and "机构与企业名录" not in md  # 没机构、没中标单位的国家不出这一附录
    assert scope == iso and cl == {l1["id"]} and co == {o["id"]}  # 只记商机的出处线索
    assert "已立项跟踪的商机" in md and "AZE-001" in md and "example.org/b" not in md and "巴库供水招标" not in md


def test_template_outline_lists_every_chapter():
    from core.jobs import country_report
    md = country_report.outline_md()
    for n, (t, secs) in enumerate(country_report.OUTLINE, 1):
        assert f"# 第 {n} 章 {t}" in md and f"## {n}.1 {secs[0][0]}" in md


def test_negative_thresholds_match_the_country_page():
    import re
    from core.country_facts import TONES
    js = (Path(__file__).resolve().parents[1] / "portal" / "portal.js").read_text(encoding="utf-8")
    block = js[js.index("TONES: ["):js.index("\n    ],", js.index("TONES: ["))]
    assert re.findall(r'"([^"]+)"\]', block) == [why for _, _, why in TONES]


def test_kpi_cards_and_red_marks_in_word(tmp_path):
    import pytest
    pytest.importorskip("pypandoc")
    docx = pytest.importorskip("docx")
    from core.watch import reporting
    md = tmp_path / "r.md"
    md.write_text("# 报告\n\n" + reports.kpi_table([("GDP", {"v": "1,470 亿美元", "year": "2025", "bad": None}),
                                                     ("增速", {"v": "-0.3%", "year": "2025", "bad": "负增长"})])
                  + "\n| 指标 | 数值 |\n|---|---|\n| 外债/GNI | 60.5%(超警戒线:外债 ≥ 60% GNI) |\n", encoding="utf-8")
    d = docx.Document(str(reporting.to_docx_styled(md, tmp_path / "r.docx", template=tmp_path / "none.docx")))
    kpi, other = d.tables
    assert "关键指标" not in "\n".join(p.text for p in d.paragraphs)  # 表格标题段落删掉了
    big = [r for p in kpi.rows[1].cells[0].paragraphs for r in p.runs]
    assert big[0].font.size.pt == 18 and big[0].font.bold
    red = lambda c: {str(r.font.color.rgb) for p in c.paragraphs for r in p.runs}  # noqa: E731
    assert red(kpi.rows[1].cells[1]) == {"C00000"} and red(other.rows[1].cells[1]) == {"C00000"}


class _Capture:
    enabled, label = True, "fake"

    def __init__(self, answers=None):
        self.data, self.answers = [], answers or {}

    def write(self, system, user, fields):
        self.data.append(user)
        return {k: self.answers.get(k, "") for k in fields}


def test_project_report_internal_chinese_and_public_english(env):  # noqa: F811
    """项目报告按《项目分析报告》模板:内部版中文报领导;公开版英文发伙伴,不给模型人写的内容(R9),没有「请领导决策」。"""
    from core.jobs import opp_report
    fake, admin, abase, ctx = env
    c, o, l1, l2 = _seed(abase)
    abase.update("opportunity", o["id"], {"internal_note": "建议动作:约见业主(到期 2026-10-31)", "name_local": "Ganja Wastewater Treatment Plant",
                                          "sectors": ["污水"]})
    cap = _Capture({"project": "甘贾新建污水厂,亚行贷款 5,000 万美元 [1]。", "phases": "\n".join(f"{p} | 未开始 | — | —" for p in opp_report.PR.PHASES)})
    ctx.clf = cap
    zh, cited, code = opp_report.build(ctx, o["id"], internal=True)
    assert zh["lang"] == "zh" and zh["to"] == "公司领导" and code == "AZE-001" and cited == {l1["id"]}
    assert zh["core"][0][1] == "甘贾新建污水厂,亚行贷款 5,000 万美元 [1]。" and zh["sources"][0][1]["url"] == "https://example.org/a"
    assert zh["actions"][0][:3] == ("约见业主", "【待定】", "2026-10-31") and zh["leader"]  # 人写的下一步进内部版
    assert "内部备注" in cap.data[0] and "[aaaaaaaa]" not in zh["basics"][4][1][0]  # 线索号换成来源编号
    cap.data.clear()
    en, _, _ = opp_report.build(ctx, o["id"], internal=False)
    assert en["lang"] == "en" and en["to"] == "Partners in Azerbaijan" and en["field"] == "Wastewater" and en["leader"] == []
    assert en["basics"][0][0] == "Name and ID" and en["phases"][0][0] == "Preparation"
    assert all("内部" not in d and "约见业主" not in d for d in cap.data)  # 公开版不给模型人写的内容
    tpl = opp_report.TEMPLATES["zh"]
    if tpl.exists():  # 私有仓的模板在时连 Word 一起出(中英文各一份)
        import docx
        from scripts.make_project_template_en import main as make_en
        en_tpl = Path(ctx.cache.root.parent / "en.docx")
        make_en(str(tpl), str(en_tpl))
        for content, t, want in ((zh, tpl, "一、核心结论"), (en, en_tpl, "1. Key Findings")):
            out = opp_report.PR.to_docx(content, ctx.cache.root.parent / f"{content['lang']}.out.docx", t)
            text = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
            assert want in text and "【" not in text.split(want)[1]


def _seed_tenders(base, c):
    """乌兹别克斯坦式的招标监测数据:两个在招标段(一个有原文摘录、一个挂了商机)、一个本期已截止、一个授标线索、一条已中标项目、一条过期旧标。"""
    o = base.create("opportunity", [{"code": "AZE-002", "name_zh": "努库斯污水厂", "country": c["id"], "stage": "跟踪中", "eligibility": "可直接投标",
                                     "eligibility_basis": "世行采购规则,各国企业均可"}])[0]
    ls = base.create("lead", [
        {"key": "t1t1t1t1", "title": "Construction of the WWTP in the city of Nukus", "url": "https://projects.worldbank.org/n/1", "relevant": True,
         "category": "招标", "country": c["id"], "published": "2026-04-06", "found_at": "2026-04-06T02:00:00.000Z", "funder": "世界银行",
         "amount": "USD 45,000,000", "amount_cny": 32000, "project_name": "努库斯污水处理厂建设", "opportunity": o["id"], "origin": "机器采集",
         "excerpt": "Bidders shall have an average annual turnover of USD 30 million. Participation in a joint venture is allowed. "
                    "A bid security of USD 400,000 is required. Bidding documents may be purchased for a non-refundable fee of USD 100. Bids will be opened on 18 May 2026."},
        {"key": "t2t2t2t2", "title": "Package 4: Water Supply Networks in Syrdarya", "url": "https://projects.worldbank.org/n/2", "relevant": True,
         "category": "招标", "country": c["id"], "published": "2026-03-20", "found_at": "2026-03-20T02:00:00.000Z", "funder": "世界银行", "origin": "机器采集",
         "summary_zh": "锡尔河州供水管网 4 标段"},
        {"key": "t3t3t3t3", "title": "Sewage pumping station in Nukus Central", "url": "https://projects.worldbank.org/n/3", "relevant": True,
         "category": "招标", "country": c["id"], "published": "2026-03-05", "found_at": "2026-03-05T02:00:00.000Z", "origin": "机器采集"},
        {"key": "t4t4t4t4", "title": "Contract award: Kuwanish Jarma Canal", "url": "https://projects.worldbank.org/n/4", "relevant": True,
         "category": "授标或短名单", "country": c["id"], "published": "2026-04-01", "found_at": "2026-04-01T02:00:00.000Z", "origin": "机器采集",
         "summary_zh": "库瓦尼什贾尔马运河改造授标"},
        {"key": "t5t5t5t5", "title": "Old tender 2025", "url": "https://projects.worldbank.org/n/5", "relevant": True, "category": "招标", "country": c["id"],
         "published": "2025-11-01", "found_at": "2025-11-01T02:00:00.000Z", "origin": "机器采集"},
    ])
    base.create("milestone", [
        {"lead": ls[0]["id"], "opportunity": o["id"], "kind": "投标截止", "day": "2026-05-18", "日期原文": "2026-05-18", "status": "官方计划"},
        {"lead": ls[1]["id"], "kind": "投标截止", "day": "2026-06-10", "日期原文": "2026-06-10", "status": "官方计划"},
        {"lead": ls[2]["id"], "kind": "投标截止", "day": "2026-04-07", "日期原文": "2026-04-07", "status": "官方计划"},
        {"lead": ls[4]["id"], "kind": "投标截止", "day": "2025-12-23", "日期原文": "2025-12-23", "status": "已发生"},
    ])
    base.create("award", [{"name": "Modernization of the Kuwanish Jarma Canal", "name_zh": "库瓦尼什贾尔马运河现代化改造", "sector": "水利", "price": 25720221,
                           "currency": "USD", "price_cny": 18500, "awarded": "2026-04-02", "winner": "China Gezhouba Group", "winner_country": "中国",
                           "funder": "世界银行", "project": "P504600 节水与保护项目", "country": c["id"], "url": "https://projects.worldbank.org/p/P504600"}])
    return o, ls


def test_tender_monitor_report_lists_open_lots_with_superscript_links(env, monkeypatch):  # noqa: F811
    fake, admin, abase, ctx = env
    monkeypatch.setattr(reports, "_to_docx", _fake_docx)
    c = abase.rows("country")[0]
    o, ls = _seed_tenders(abase, c)
    ctx.today = date(2026, 4, 10)
    abase.create("report", [{"kind": "招标监测", "country": c["id"], "day": "2026-04-10", "period": "2026-03-11~2026-04-10", "gen_status": "待生成"},
                            {"kind": "招标监测", "day": "2026-04-10", "gen_status": "待生成"}])  # 没选国别:失败
    assert reports.run(ctx) == {"报告完成": 1, "报告失败": 1}
    r = next(x for x in _rows(abase, "report") if _val(abase, "report", x, "gen_status") == "完成")
    assert _val(abase, "report", r, "period") == "AZE 2026-03-11~2026-04-10"
    name = _val(abase, "report", r, "key")
    assert name == "招标监测_AZE_20260311-0410_2026-04-10.docx"
    md = (ctx.cache.root.parent / "reports" / name.replace(".docx", ".md")).read_text(encoding="utf-8")
    # 概览数字:在招 2(5-18、6-10)、本期新发布 2(4-06、3-20)、已截止等结果 1(4-07)、授标 2(线索 + 已中标项目);旧标(2025)不出现
    assert "在招的水务标段 2 个,最近的 2026-05-18 截止(剩 38 天)" in md and "新发布 2 个、已截止等结果 1 个、授标 2 项" in md
    assert "t5t5t5t5" not in md
    # 角标 = 上标超链接,回到原公告;同一网址一个编号;文末列全
    assert "^[1](<https://projects.worldbank.org/n/1>)^" in md and md.count("^[1](<https://projects.worldbank.org/n/1>)^") >= 2
    assert "## 7. 原公告链接" in md and "1. projects.worldbank.org:" in md
    # 中文名在前、原文在后;金额带折算;已中标项目带中标单位与国别
    assert "努库斯污水处理厂建设" in md and "(原文:Construction of the WWTP in the city of Nukus)" in md and "USD 45,000,000(约 32,000 万元)" in md and "2026-05-18(剩 38 天)" in md
    assert "库瓦尼什贾尔马运河现代化改造" in md and "China Gezhouba Group(中国)" in md and "25,720,221 USD(约 18,500 万元)" in md
    # 资格要求:没配模型 → 按关键句规则抽英文原句,表里写明;商机表里的资格判断一并列出
    assert "按关键句规则" in md and "average annual turnover of USD 30 million" in md and "joint venture is allowed" in md
    assert "bid security of USD 400,000" in md and "可直接投标 世行采购规则,各国企业均可" in md
    # 引用:列出的线索(不含旧标),写进的商机
    assert {x["id"] for x in _val(abase, "report", r, "cites_leads")} == {x["id"] for x in ls[:4]}
    assert [x["id"] for x in _val(abase, "report", r, "cites_opps")] == [o["id"]]
    # 第二次生成同一国别:资格要求走缓存,不再抽
    cache = ctx.cache.root.parent / "tender_elig.json"
    assert cache.exists() and "t1t1t1t1" in cache.read_text(encoding="utf-8")


def test_tender_monitor_window_defaults_to_last_30_days():
    from core.jobs import tender_monitor as tm
    assert tm.window("", date(2026, 10, 16)) == (date(2026, 9, 17), date(2026, 10, 16))
    assert tm.window("UZB 2026-03-11~2026-04-10", date(2026, 10, 16)) == (date(2026, 3, 11), date(2026, 4, 10))
    assert tm.window("2026-04-10~2026-03-11", date(2026, 10, 16)) == (date(2026, 9, 17), date(2026, 10, 16))  # 起晚于止:当没填
    assert tm.period_text("UZB", date(2026, 3, 11), date(2026, 4, 10)) == "UZB 2026-03-11~2026-04-10"


def test_country_report_carries_the_other_country_page_tabs(env, monkeypatch):  # noqa: F811
    """2026-10-16 使用方:国别概况页有的(已中标项目、营商细项、法规原文与排放标准、本国定义、大事记、片区、人均 GDP、研究资料)报告模板也要有。"""
    from core.jobs import country_report
    fake, admin, abase, ctx = env
    c, o, l1, l2 = _seed(abase)
    abase.update("country", c["id"], {"gdp_pc": 11879, "water_stress": 3.63, "note": "2026-10 [情报] 水务公司改制中", "bready_total": 68.8, "bready_procure": 45.7,
                                      "bready_edition": "B-READY 2025", "bready_items": "[采购] 外企参标或授标受限:是\n[采购] 有电子采购平台:是\n[供水] 现行水价在网上公布:否"})
    law = abase.create("law", [{"name": "水法典", "kind": "水法与水资源", "number": "O'RQ-1076", "issued": "2025-07-30", "status": "现行有效", "country": c["id"], "url": "https://lex.uz/docs/7655343"}])[0]
    abase.create("discharge_limit", [{"param": "BOD5", "value": 120, "unit": "mg/L", "basis": "出水浓度上限", "receiving": "地表水", "country": c["id"], "law": law["id"], "verified": "待核实"}])
    abase.create("local_term", [{"concept": "生活污水", "term_zh": "家庭污水", "relation": "更窄", "difference": "不含商业排水", "country": c["id"], "law": law["id"]}])
    abase.create("water_event", [{"day": "2025-10-31", "title": "新水法典施行", "kind": "规划与政策", "country": c["id"], "source": "lex.uz", "url": "https://lex.uz/docs/7655343"}])
    abase.create("award", [
        {"name": "Tashkent WWTP rehabilitation", "name_zh": "塔什干污水厂改造", "sector": "污水", "price": 50000000, "currency": "USD", "price_cny": 35500, "awarded": "2026-03-01",
         "winner": "China Railway X", "winner_country": "中国", "country": c["id"], "url": "https://projects.worldbank.org/a/1", "buyer": "Uzsuvtaminot"},
        {"name": "Nukus water network", "sector": "供水", "price": 8000000, "currency": "USD", "price_cny": 5680, "awarded": "2025-06-01",
         "winner": "Local Build LLC", "winner_country": "乌兹别克斯坦", "country": c["id"], "url": "https://projects.worldbank.org/a/2"},
        {"name": "Old pipes 2015", "sector": "供水", "price_cny": 100, "awarded": "2015-01-01", "winner": "Old Co", "winner_country": "土耳其", "country": c["id"]},
    ])
    abase.create("study", [{"name": "Water Sector Review", "name_zh": "水务行业评估", "kind": "行业研究", "publisher": "世界银行", "published": "2024-05-01", "country": c["id"],
                            "url": "https://documents.worldbank.org/s/1"}])
    iso = abase.get("country", abase.record("country", c["id"]), "key")
    md, *_ = reports.build_country(ctx, iso)
    # 竞争格局:近三年的中标项目按金额列,角标回原公告;中国企业份额;主要中标单位与国别分布;2015 年的旧记录只计入总数
    assert "已中标项目 3 份(2015–2026 年),其中近 3 年(2023 年起)2 份,写明中标价的合计约 4.12 亿元;中国企业中标 1 份,占 50%" in md
    assert "| 塔什干污水厂改造^[2](<https://projects.worldbank.org/a/1>)^ | 污水 | China Railway X(中国) | 3.55 亿元 | 2026-03-01 | Uzsuvtaminot |" in md
    assert "Old pipes 2015" not in md and "| 中国 | 1 | 50% |" in md and "| China Railway X | 中国 | 1 | 3.55 亿元 |" in md
    # 营商环境细项:对外企不利的标超警戒线
    assert "| 外企参标或授标受限 | 是(超警戒线:对外企投标不利) |" in md and "| 有电子采购平台 | 是 |" in md and "| 现行水价在网上公布 | 否 |" in md
    # 法规原文、排放标准、本国定义、大事记各成一节,角标按出现顺序编号,同一网址一个号(法规在第 5 章先出现,所以是 1)
    assert "| 水法典^[1](<https://lex.uz/docs/7655343>)^ | 水法与水资源 | O'RQ-1076 | 2025-07-30 | 现行有效 |" in md
    assert "| BOD5 | 120.00 mg/L | 出水浓度上限 | 地表水 |  | 水法典^[1](<https://lex.uz/docs/7655343>)^ | 待核实 |" in md and "1 条还没人工核实" in md
    assert "| 生活污水 | 家庭污水 | 更窄 | 不含商业排水 | 水法典^[1]" in md
    assert "| 2025-10-31 | 规划与政策 | 新水法典施行^[1](<https://lex.uz/docs/7655343>)^ |  | lex.uz |" in md
    # 人均 GDP(PPP)进宏观经济表;总体水风险表;国别备注进内部版、不进公开版;研究资料与原文链接进附录
    assert "| 人均 GDP(PPP) | 11,879 美元(2021 年国际美元) | 2024 | [Kummu 等(2025)](<https://doi.org/10.1038/s41597-025-04487-x>) |" in md
    assert "| 用水压力(当前) | 3.6 分(高) 超警戒线(≥ 3 分) |" in md
    assert re.search(r"^## 2\.\d 国别备注与近期动态$", md, re.MULTILINE) and "- 2026-10 [情报] 水务公司改制中" in md
    assert re.search(r"^## \d+\.\d 中国企业在该国业绩与经验教训$", md, re.MULTILINE) and "库里登记中国企业在该国的已中标项目 1 份" in md
    assert "## 附录 F 参考研究资料" in md and "水务行业评估^[4](<https://documents.worldbank.org/s/1>)^ | 行业研究 | 世界银行 | 2024-05-01 |" in md
    assert "## 附录 G 原文链接" in md and "2. projects.worldbank.org:[塔什干污水厂改造](<https://projects.worldbank.org/a/1>)(原文:Tashkent WWTP rehabilitation)" in md
    heads = re.findall(r"^# 第 \d+ 章 (.+)$", md, re.MULTILINE)
    assert heads == [t for t, _ in country_report.OUTLINE if t in heads]
    pub, *_ = reports.build_country(ctx, iso, internal=False)
    assert "国别备注与近期动态" not in pub and "塔什干污水厂改造" in pub
