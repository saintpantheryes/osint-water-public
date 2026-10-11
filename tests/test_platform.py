"""平台功能:国别识别、领域分类、全球信源采集器、全球范围运行、人工入库、筛选与处理状态、报告生成。
全部离线:网页与模型接口都用假对象;网址为虚构,国名识别用真实国名。"""
import json
from datetime import date
from types import SimpleNamespace

import httpx
import pytest
import yaml

from core import geo, storage
from core.profile import load_profile
from core.watch import fetch, manual
from core.watch.collectors import KINDS, Ctx
from core.watch.llm import Classifier, DeepSeekClient, LLMConfig
from core.watch.manual import SubmitError, submit
from core.watch.outputs import sectors_of
from core.watch.publish import write_outputs
from core.watch.query import Filter, apply, delete_view, load_rows, load_views, rename_view, same_filter, save_view, set_status
from core.watch.reporting import TEMPLATE, Spec, build, template_spec, to_docx_styled
from core.watch.rules import evaluate
from core.watch.run import DEFAULTS, GLOBAL, WatchConfigError, load_watch, run_scope, save_user_sources, scopes
from core.watch.store import Item, Store
from core.watch.text import parse_date, scrub_contacts
from tests.test_watch import FX, OK, ROOT, FakeWeb, _docx, _ds_resp

TODAY = date(2026, 9, 26)
CTX = Ctx(iso3="", names=[], wb_code="", since=date(2026, 7, 28), today=TODAY)


# ---------------------------------------------------------------- 国别与领域


@pytest.mark.parametrize("text,expect", [
    ("Water supply in Azerbaijan and Georgia", ["AZE", "GEO"]),
    ("Водоснабжение в Азербайджане", ["AZE"]),  # 西里尔字母带格变化词尾
    ("乌兹别克斯坦塔什干污水处理厂", ["UZB"]),
    ("Papua New Guinea: rural water", ["PNG"]),  # 长的名字优先,不配成几内亚
    ("Congo, Democratic Republic of: Kinshasa water", ["COD"]),  # 世行的写法
    ("Niger Delta water supply", ["NGA"]),
    ("Multicountry", []),
])
def test_geo_detect(text, expect):
    assert geo.detect(text) == expect


def test_geo_codes_names_and_regions():
    assert [geo.iso3(v) for v in ("AZ", "aze", "ZR", "YF", "XK", "Türkiye", "UZBEKISTAN", "N/A", "", None)] == [
        "AZE", "AZE", "COD", "SRB", "XKX", "TUR", "UZB", None, None, None]  # 世行国别码 ZR、YF、XK 与 ISO 不同
    assert (geo.name_zh("AZE"), geo.continent("AZE"), geo.region("AZE")) == ("阿塞拜疆", "亚洲", "西亚")
    assert (geo.name_zh(None), geo.continent(None), geo.region("")) == ("未识别", "其他", "其他")
    regions = geo.regions_by_continent()
    assert list(regions)[:5] == geo.CONTINENTS[:5] and "中亚" in regions["亚洲"] and "东非" in regions["非洲"]


@pytest.mark.parametrize("raw,iso", [
    ("25-Sep-26", "2026-09-25"),  # 联合国开发署
    ("Sep 15,2026", "2026-09-15"),  # 亚投行
    ("July 27,2026", "2026-07-27"),
    ("23 Sentyabr,2026", "2026-09-23"),
    ("Phase 2,2026", None),
])
def test_parse_date_compact_forms(raw, iso):
    assert parse_date(raw) == iso


@pytest.mark.parametrize("raw,clean", [
    ("Contact: Mr. A. Aliyev, tel: +994 12 598 02 00, e-mail: a.aliyev@azersu.test",
     "Contact: Mr. A. Aliyev, tel: [电话已略], e-mail: [邮箱已略]"),
    ("Tel.: (012) 555-12-34; Fax: 012 555 12 35", "Tel.: [电话已略]; Fax: [电话已略]"),
    ("Телефон: +7 (495) 123-45-67", "Телефон: [电话已略]"),
    ("Tel/faks: 012 555 55 55; məlumat üçün telefon: (012) 598-01-30", "Tel/faks: [电话已略]; məlumat üçün telefon: [电话已略]"),
    ("Call +1 202 473 1000 or 0044 20 7946 0958", "Call [电话已略] or [电话已略]"),
    # 日期、时间、金额、贷款号、标准号不动
    ("Deadline 2026-10-07 17:00; USD 11,000,000.00; 500 000 m3/day; 25/09/2026 00:48; 2026-09-24T10:00:00+04:00",
     "Deadline 2026-10-07 17:00; USD 11,000,000.00; 500 000 m3/day; 25/09/2026 00:48; 2026-09-24T10:00:00+04:00"),
    ("Loan 9876-AZ, P176429, OP00456235, ISO 9001:2015, Phase 2 (2026-2028)", "Loan 9876-AZ, P176429, OP00456235, ISO 9001:2015, Phase 2 (2026-2028)"),
])
def test_scrub_contacts(raw, clean):
    assert scrub_contacts(raw) == clean


@pytest.mark.parametrize("title,sectors", [
    ("Sumgayit wastewater treatment plant", ["wastewater"]),
    ("Seawater desalination plant 100,000 m3/day", ["desalination"]),
    ("Irrigation canal rehabilitation", ["hydraulic"]),
    ("Storm water drainage in Baku", ["stormwater"]),
    ("Drinking water supply network", ["water_supply"]),
    ("Lake restoration and river pollution control", ["environment"]),
    ("Kanalizasiya şəbəkəsinin tikintisi", ["wastewater"]),  # kanalizasiya 是排水,不是"渠"
    ("Magistral kanalın bərpası", ["hydraulic"]),
    ("Строительство канализации", ["wastewater"]),
    ("Реконструкция канала", ["hydraulic"]),
    ("Ministry of Water Resources staff training", []),  # 机构名里的"水资源"不定领域(仍送模型判定)
])
def test_rule_sectors(title, sectors):
    r = evaluate(Item(key=title, url="https://ex.test/", title=title), {}, [])
    assert r["candidate"] and r["sectors"] == sectors


def test_sectors_of_prefers_model_then_rules_then_old_terms():
    sig = {"rules": {"sectors": ["hydraulic"], "water_terms": ["canal"]}, "llm": {"relevant": True, "sectors": ["water_supply", "bogus"]}}
    assert sectors_of(sig) == ["water_supply"]
    assert sectors_of({**sig, "llm": {"_error": "x"}}) == ["hydraulic"]
    assert sectors_of({"rules": {"water_terms": ["WWTP", "desalination"]}}) == ["wastewater", "desalination"]  # 旧线索按命中词补
    assert sectors_of({"rules": {}}) == ["other"]


# ---------------------------------------------------------------- 全球信源采集器


def _page(url, body, ctype="text/html; charset=utf-8"):
    return fetch.Page(url, 200, body.encode("utf-8") if isinstance(body, str) else body, ctype)


def _serve(monkeypatch, pages: dict, calls: list | None = None):
    def get(url, params=None, **kw):
        if calls is not None:
            calls.append((url, params))
        return _page(url, pages[url])

    monkeypatch.setattr(fetch, "get", get)


def test_html_list_country_from_title(monkeypatch):
    html = """<table id="t"><tbody>
      <tr><td><a href="/n?id=1">Uzbekistan: Tashkent Wastewater - Construction of treatment plant</a></td><td>Invitation for Tenders</td>
          <td>24/09/2026</td><td>30/10/2026</td></tr>
      <tr><td><a href="/n?id=2">Regional: Water loss programme</a></td><td>General Procurement Notice</td><td>23/09/2026</td><td></td></tr>
    </tbody></table>"""
    _serve(monkeypatch, {"https://mdb.test/list": html})
    src = {"id": "ebrd", "kind": "html_list", "url": "https://mdb.test/list", "item": "#t tbody tr", "title": "td:nth-of-type(1) a",
           "date": "td:nth-of-type(3)", "fields": {"notice_type": "td:nth-of-type(2)", "closing": "td:nth-of-type(4)"},
           "country_from": {"field": "title", "regex": "^([^:]+):"}, "deadline_field": "closing"}
    a, b = KINDS["html_list"](src, CTX)
    assert (a.url, a.published, a.fields["country"], a.fields["deadline"]) == ("https://mdb.test/n?id=1", "2026-09-24", "Uzbekistan", "2026-10-30")
    assert b.fields["country"] == "Regional" and geo.iso3(b.fields["country"]) is None  # 多国项目:国别留空,交给标题识别或模型


AIIB_PPO = ('var ppoData = [ {id:"Sep 14,2026",cd:"Sep 30,2026",mb:"Uzbekistan",pj:"Bukhara Region Water Supply and Sewerage Project",'
            'ds:"Addendum 1",cr:"",sd:"",pc:"",st:"Water ",ct:"Notices",tp:"Addenda to Request for Expressions of Interest",'
            'dc:"/en/projects/details/2026/_download/Uzbekistan/reoi-addendum.pdf"},'
            ' {id:" Sep 1,2026",cd:"",mb:"TÃ¼rkiye",pj:"Istanbul Metro",ds:"",cr:"",sd:"",pc:"",st:"Transport",ct:"Procurement Plan",'
            'tp:"Procurement Plan",dc:"/en/x.xlsx"}, ];')
AIIB_PROJECTS = ('var data=[{"date":"-","economy":"Azerbaijan","sector":"Water","financing_type":"Nonsovereign",'
                 '"name":"Azerbaijan: Caspian Sea Desalination Project","approved_funding":"","committed_funding":"",'
                 '"proposed_funding":"USD45 million","status":"Proposed","path":"/en/projects/details/2026/proposed/caspian.html"},'
                 '{"date":"2025","economy":"India","sector":"Energy","financing_type":"Sovereign","name":"India: Solar",'
                 '"approved_funding":"USD100 million","status":"Approved","path":"/en/projects/details/2025/approved/solar.html"},];')


def test_aiib_notices_and_projects(monkeypatch):
    _serve(monkeypatch, {"https://aiib.test/ppo.js": AIIB_PPO, "https://aiib.test/projects.js": AIIB_PROJECTS})
    water, = KINDS["aiib"]({"id": "a", "url": "https://aiib.test/ppo.js", "sectors": ["Water"]}, CTX)  # 键没加引号、末尾多逗号也能读
    assert water.title == "Addenda to Request for Expressions of Interest: Bukhara Region Water Supply and Sewerage Project"
    assert water.url == "https://www.aiib.org/en/projects/details/2026/_download/Uzbekistan/reoi-addendum.pdf"
    assert (water.published, water.fields["deadline"], water.fields["country"]) == ("2026-09-14", "2026-09-30", "Uzbekistan")
    both = KINDS["aiib"]({"id": "a", "url": "https://aiib.test/ppo.js"}, CTX)
    assert both[1].fields["country"] == "Türkiye" and both[1].published == "2026-09-01"  # 修好读错编码的国名
    proj, = KINDS["aiib"]({"id": "p", "url": "https://aiib.test/projects.js", "dataset": "projects", "sectors": ["water"]}, CTX)
    assert proj.title == "Azerbaijan: Caspian Sea Desalination Project" and proj.fields["funding"] == "USD45 million"
    assert proj.fingerprint  # 状态或金额变了会出一条"变化"
    _serve(monkeypatch, {"https://aiib.test/projects.js": AIIB_PROJECTS.replace('"Proposed"', '"Approved"')})
    again, = KINDS["aiib"]({"id": "p", "url": "https://aiib.test/projects.js", "dataset": "projects", "sectors": ["water"]}, CTX)
    assert again.key == proj.key and again.fingerprint != proj.fingerprint


def test_ckan_collector(monkeypatch):
    records = [{"noticeid": "IDB-1", "noticetitle": "Construcción de planta de tratamiento de aguas residuales", "countryname": "PARAGUAY",
                "publicationdate": "2026-09-24T00:00:00", "deadline": "2026-10-30", "documenturl": "https://idb.test/n/1",
                "type": "SPECIFIC", "process_desc": "Obras"}]
    calls: list = []
    _serve(monkeypatch, {"https://idb.test/api": json.dumps({"success": True, "result": {"records": records}})}, calls)
    src = {"id": "idb", "kind": "ckan", "url": "https://idb.test/api", "resource_id": "r-1", "sort": "publicationdate desc",
           "filters": {"type": ["SPECIFIC", "EOI"]}, "fields": {"notice_type": "type"},
           "map": {"key": "noticeid", "title": "noticetitle", "url": "documenturl", "published": "publicationdate",
                   "country": "countryname", "deadline": "deadline", "text": "process_desc"}}
    it, = KINDS["ckan"](src, CTX)
    assert calls[0][1] == {"resource_id": "r-1", "limit": 500, "sort": "publicationdate desc", "filters": '{"type": ["SPECIFIC", "EOI"]}'}
    assert (it.key, it.url, it.published, it.text) == ("idb:IDB-1", "https://idb.test/n/1", "2026-09-24", "Obras")
    assert it.fields == {"notice_type": "SPECIFIC", "country": "PARAGUAY", "deadline": "2026-10-30"} and geo.iso3("PARAGUAY") == "PRY"
    _serve(monkeypatch, {"https://idb.test/api": json.dumps({"success": False, "error": {"message": "bad resource"}})})
    with pytest.raises(RuntimeError, match="bad resource"):
        KINDS["ckan"](src, CTX)


LINKS_HTML = """<html><body><nav><a href="/">Home page of the ministry of water</a></nav>
<div class="news">
  <div class="item"><span>24 Sentyabr 2026</span><h3><a href="/news/101">Suvarma kanalının tikintisi üzrə tender elan edilir</a></h3></div>
  <div class="item"><span>19.09.2026</span><a href="/news/100">Sumqayıt kanalizasiya təmizləyici qurğusu layihəsi</a>
      <a href="/files/100.pdf">PDF</a></div>
  <div class="item"><a href="/news/99">Su təchizatı sisteminin yenidən qurulması başa çatıb</a></div>
  <div class="item"><a href="/news/archive">Daha çox</a></div>
  <div class="item"><a href="/about/contacts-and-structure">Ministry structure and contacts page</a></div>
  <div class="item"><a href="mailto:info@x.test">info@x.test mailbox for citizen enquiries</a></div>
</div>
<footer><a href="/privacy">Privacy policy and terms of use of the site</a></footer></body></html>"""


def test_links_collector(monkeypatch):
    _serve(monkeypatch, {"https://ministry.test/az/news": LINKS_HTML})
    src = {"id": "u", "kind": "links", "url": "https://ministry.test/az/news", "exclude": "/about/"}
    items = KINDS["links"](src, CTX)
    assert [(i.url, i.published) for i in items] == [
        ("https://ministry.test/news/101", "2026-09-24"),
        ("https://ministry.test/news/100", "2026-09-19"),  # 同块里的附件短链接不妨碍找日期
        ("https://ministry.test/news/99", None),  # 没有日期的不拿别条的日期
    ]  # 导航、页脚、"更多"、邮件链接、排除的地址都不收
    assert [i.url for i in KINDS["links"]({**src, "selector": ".news", "include": "news/10"}, CTX)] == [
        "https://ministry.test/news/101", "https://ministry.test/news/100"]
    with pytest.raises(RuntimeError, match="选择器"):
        KINDS["links"]({**src, "selector": ".missing"}, CTX)


def test_real_global_config_loads():
    wc = load_watch(ROOT / "global" / "watch.yaml")
    ids = [s["id"] for s in wc["sources"]]
    assert "wb_notices" in ids and len(ids) == len(set(ids))
    for s in wc["sources"]:
        assert s.get("name") and (s.get("url") or s["kind"] == "wb_procnotices"), s["id"]
    by_id = {s["id"]: s for s in wc["sources"]}  # 海湾国家的水务 PPP 采购公司(海水淡化、污水、输水)
    assert (by_id["sharakat_news"]["country"], by_id["sharakat_news"]["topic"]) == ("SAU", "water")
    assert by_id["nama_pwp_news"]["country"] == "OMN" and "topic" not in by_id["nama_pwp_news"]  # 电力项目不收:要命中水务词


NAMA_HTML = """<html><body><nav><a href="/news">News</a></nav><div class="row">
<div class="news-item"><figure><img src="/x.png"></figure><div class="news-body"><div class="news-date">27-Sep-26</div>
  <h3> Nama PWP Launches Qualification Phase for Ghubrah 3 IWP </h3><p>Seawater reverse osmosis plant of 300,000 m3/day...</p>
  <a class="btn-com news-nav" href="https://www.omanpwp.om/news-details/ghubrah-3-iwp"><span>Read More</span></a></div></div>
<div class="news-item"><div class="news-body"><div class="news-date">10-May-26</div><h3> Adam Solar IPP bids received </h3>
  <p>Solar photovoltaic plant</p><a class="btn-com news-nav" href="https://www.omanpwp.om/news-details/adam-solar"><span>Read More</span></a>
</div></div></div></body></html>"""


def test_nama_pwp_news_config_parses(monkeypatch):
    """阿曼 Nama PWP 新闻栏按真实配置解析:标题、日期、摘要;水务的过初筛,光伏项目不过。"""
    src = next(s for s in load_watch(ROOT / "global" / "watch.yaml")["sources"] if s["id"] == "nama_pwp_news")
    _serve(monkeypatch, {src["url"]: NAMA_HTML})
    items = KINDS[src["kind"]](src, CTX)
    assert [(i.title, i.published, i.url) for i in items] == [
        ("Nama PWP Launches Qualification Phase for Ghubrah 3 IWP", "2026-09-27", "https://www.omanpwp.om/news-details/ghubrah-3-iwp"),
        ("Adam Solar IPP bids received", "2026-05-10", "https://www.omanpwp.om/news-details/adam-solar")]
    assert "reverse osmosis" in items[0].text
    assert [evaluate(i, src, [])["candidate"] for i in items] == [True, False]


# ---------------------------------------------------------------- 全球范围运行


MDB_HTML = """<table>
<tr><th>Date</th><th>Title</th><th>Closing</th></tr>
<tr><td>20/09/2026</td><td><a href="/n/1">Uzbekistan: Bukhara water supply and sewerage - Invitation for Bids</a></td><td>30/10/2026</td></tr>
<tr><td>21/09/2026</td><td><a href="/n/2">Kenya: Road maintenance works</a></td><td>01/11/2026</td></tr>
<tr><td>22/09/2026</td><td><a href="/n/3">Peru: Lima desalination plant - Request for Expressions of Interest</a></td><td>15/10/2026</td></tr>
<tr><td>01/06/2026</td><td><a href="/n/4">Ghana: Accra sewerage network</a></td><td>01/07/2026</td></tr>
</table>"""
WB_JSON = {"procnotices": [
    {"id": "OP00999", "notice_type": "Invitation for Bids", "project_ctry_code": "KE", "bid_description": "Nairobi sewerage network extension",
     "noticedate": "24-Sep-2026", "submission_deadline_date": "2026-11-05T00:00:00Z",
     "notice_text": "<p>Sewerage works. Contact: Ms. J. Doe, tel: +254 20 123 4567, e-mail: j.doe@nws.test</p>"},
    {"id": "OP01000", "notice_type": "Request for Expression of Interest", "project_ctry_code": "ZR",
     "bid_description": "Kinshasa water supply rehabilitation", "noticedate": "23-Sep-2026", "notice_text": ""},
]}
GSOURCES = [
    {"id": "mdb_list", "name": "测试银行采购(全部国别)", "kind": "html_list", "url": "https://mdb.test/notices", "item": "table tr",
     "title": "td:nth-of-type(2) a", "date": "td:nth-of-type(1)", "fields": {"closing": "td:nth-of-type(3)"},
     "country_from": {"field": "title", "regex": "^([^:]+):"}, "deadline_field": "closing", "lookback_days": 30,
     "seed_from": ["old_notices"]},
    {"id": "wb_global", "name": "世行采购公告(全球)", "kind": "wb_procnotices", "rows": 50, "qterm": "water OR sewerage"},
]


@pytest.fixture
def genv(tmp_path, monkeypatch):
    """全球范围的临时 profile:虚构国别 ZZZ(有台账、没有国别信源)+ 全球配置 + 假网站与假世行接口;不读真实凭据。"""
    for k in ("ANTHROPIC_API_KEY", "ANTHROPIC_AUTH_TOKEN", "ANTHROPIC_PROFILE", "DEEPSEEK_API_KEY"):
        monkeypatch.delenv(k, raising=False)
    (tmp_path / "global").mkdir()
    (tmp_path / "global" / "watch.yaml").write_text(yaml.safe_dump(
        {"schema": "osint-water/watch@1", "country": "GLOBAL", "sources": GSOURCES}, allow_unicode=True), encoding="utf-8")
    prof = yaml.safe_load((FX / "profile.yaml").read_text(encoding="utf-8"))
    prof["countries"][0].update(ledger=str(FX / "ZZZ" / "ledger.yaml"), overlay=str(FX / "overlay.yaml"))
    prof.update(procurement_rules=str(FX / "rules.yaml"), data_dir="data", watch={"global": "global/watch.yaml", "inbox_days": 14})
    (tmp_path / "profile.yaml").write_text(yaml.safe_dump(prof, allow_unicode=True), encoding="utf-8")
    profile = load_profile(tmp_path / "profile.yaml")
    monkeypatch.setattr(storage, "DATA_DIR", profile.data_dir)
    web = FakeWeb({"https://mdb.test/notices": MDB_HTML})
    calls: list = []

    def get(url, params=None, **kw):
        calls.append((url, params))
        if "worldbank" in url:
            return fetch.Page(url, 200, json.dumps(WB_JSON).encode(), "application/json")
        return web.get(url, params, **kw)

    monkeypatch.setattr(fetch, "get", get)
    return SimpleNamespace(profile=profile, web=web, calls=calls, store=Store(profile.data_dir), data=profile.data_dir)


def test_scopes_merge_global_and_user_sources(genv):
    sc, = scopes(genv.profile, None, genv.data)
    assert (sc.key, sc.iso3, sc.label) == (GLOBAL, None, "全球") and [s["id"] for s in sc.wc["sources"]] == ["mdb_list", "wb_global"]
    assert sc.led.items[0]["id"] == "ZZZ-001"  # 全球范围用各国台账合并
    user = {"schema": "osint-water/watch@1", "sources": [{"id": "u_1", "name": "用户网站", "kind": "links", "url": "https://user.test/news"},
                                                         {"id": "u_2", "kind": "rss", "url": "https://user.test/rss", "disabled": True}],
            "disabled": ["wb_global"]}
    save_user_sources(genv.data, user)
    sc, = scopes(genv.profile, "GLOBAL", genv.data)
    assert [(s["id"], s.get("user")) for s in sc.wc["sources"]] == [("mdb_list", None), ("u_1", True)]  # 停用的不跑
    save_user_sources(genv.data, {**user, "sources": [{"id": "mdb_list", "kind": "links", "url": "https://user.test/x"}], "disabled": []})
    with pytest.raises(WatchConfigError, match="重复"):
        scopes(genv.profile, None, genv.data)


def test_global_run_countries_seed_and_inbox(genv):
    genv.store.src("old_notices")["seen"]["https://mdb.test/n/1"] = {"first": "2026-09-20", "fp": ""}  # 原国别信源已提醒过
    sc, = scopes(genv.profile, None, genv.data)
    res = run_scope(genv.profile, sc, genv.store, None, TODAY, {**DEFAULTS, "inbox_days": 14})
    genv.store.save()
    got = {s["key"]: s for s in res.signals}
    # n/1 沿用旧信源的已见记录;n/2 不是水务(全球信源只存候选);n/4 早于回溯期,首次只记为已见
    assert set(got) == {"https://mdb.test/n/3", "wb:OP00999", "wb:OP01000"}
    assert [got[k]["country"] for k in ("https://mdb.test/n/3", "wb:OP00999", "wb:OP01000")] == ["PER", "KEN", "COD"]
    assert got["https://mdb.test/n/3"]["scope"] == GLOBAL and got["https://mdb.test/n/3"]["fields"]["deadline"] == "2026-10-15"
    assert got["https://mdb.test/n/3"]["rules"]["sectors"] == ["desalination"] and got["wb:OP00999"]["rules"]["sectors"] == ["wastewater"]
    ex = got["wb:OP00999"]["excerpt"]
    assert "Sewerage works" in ex and "j.doe@" not in ex and "123 4567" not in ex and "[邮箱已略]" in ex  # 联系方式不入库
    wb_params = next(p for u, p in genv.calls if "worldbank" in u)
    assert "project_ctry_code_exact" not in wb_params and wb_params["qterm"] == "water OR sewerage"  # 全球范围不按国别过滤
    assert res.sources["mdb_list"]["baseline"] is True and res.sources["wb_global"]["new"] == 2
    write_outputs(genv.profile, [res], None, TODAY)
    page = (genv.data / "inbox" / "GLOBAL.md").read_text(encoding="utf-8")
    assert page.startswith("# 全球 · 系统提醒") and "### 美洲(1 条)" in page and "### 非洲(2 条)" in page
    assert "Lima desalination" in page and "秘鲁" in page and "## 三、全球信源状态" in page
    overview = (genv.data / "README.md").read_text(encoding="utf-8")
    assert "全球(其他国别)" in overview and "[inbox/GLOBAL.md](inbox/GLOBAL.md)" in overview
    assert "全球信源:本次新增线索 3 条" in (genv.data / "watch" / "summary.md").read_text(encoding="utf-8")
    # 同一份摘要的数字另存 watch/summary.json(表格「每日摘要」用),自动入账后补上条数
    from core.watch.publish import read_summary, update_summary
    d = read_summary(genv.data)
    assert d["date"] == TODAY.isoformat() and d["new"] == 3 and "autoledger" not in d
    assert f"本次新增线索 {d['new']} 条,判定相关 {d['new_relevant']} 条" in (genv.data / "watch" / "summary.md").read_text(encoding="utf-8")
    assert isinstance(d["due_14d"], int) and isinstance(d["source_errors"], list)
    update_summary(genv.data, autoledger=2)
    assert read_summary(genv.data)["autoledger"] == 2 and read_summary(genv.data)["new"] == 3
    update_summary(genv.data / "nowhere", autoledger=1)
    assert not (genv.data / "nowhere").exists()


# ---------------------------------------------------------------- 人工入库


def test_manual_submit_text_file_and_duplicates(genv):
    p, st = genv.profile, genv.store
    text = ("Republic of Uzbekistan. Ministry of Water Resources. Invitation for Bids for the Bukhara wastewater treatment plant "
            "rehabilitation. Deadline for submission of bids: 30 November 2026. Contact: tel. +998 71 123 45 67, piu@uzsuv.test")
    sig, new = submit(p, st, None, text=text, note="同事转来", today=TODAY)
    assert new and (sig["source"], sig["scope"], sig["country"]) == ("manual", "manual", "UZB")
    assert sig["rules"]["candidate"] and sig["rules"]["sectors"] == ["wastewater"] and sig["rules"]["stage"] == "tender"
    assert sig["title"].startswith("Republic of Uzbekistan") and sig["fields"] == {"note": "同事转来"}
    cached = st.doc_path(sig["key"]).read_text(encoding="utf-8")
    assert "piu@" not in sig["excerpt"] + cached and "123 45 67" not in sig["excerpt"] + cached
    again, new = submit(p, st, None, text=text, today=TODAY)
    assert not new and again["id"] == sig["id"]  # 同一段文字不重复入库
    doc = _docx("Kenya: Mombasa seawater desalination plant PPP. Request for expressions of interest.")
    sig2, new = submit(p, st, None, file=("mombasa.docx", doc), today=TODAY)
    assert new and sig2["country"] == "KEN" and sig2["rules"]["sectors"] == ["desalination"] and sig2["fields"]["file"] == "mombasa.docx"
    assert {s["id"] for s in st.load_signals()} == {sig["id"], sig2["id"]}  # 人工提交的也在线索库里
    for kw, msg in [({"file": ("x.xlsx", b"PK")}, "不支持"), ({"url": "ftp://x.test/a"}, "http"), ({}, "请给出"),
                    ({"text": "too short"}, "没有读到正文")]:
        with pytest.raises(SubmitError, match=msg):
            submit(p, st, None, today=TODAY, **kw)


def test_manual_submit_url_with_attachments_and_search_result(genv):
    genv.web.pages.update({
        "https://owner.test/news/1": "<html><head><title>Owner</title></head><body><h1>Tashkent water supply tender announced</h1>"
                                     "<div>The utility announced a tender. <a href='/f/tor.docx'>ToR</a> <a href='/f/map.pdf'>Map</a></div>"
                                     "</body></html>",
        "https://owner.test/f/tor.docx": _docx("Terms of reference: water supply network rehabilitation, Uzbekistan."),
        "https://owner.test/f/map.pdf": fetch.FetchError("HTTP 404 from https://owner.test/f/map.pdf"),
    })
    sig, new = submit(genv.profile, genv.store, None, url="https://owner.test/news/1", today=TODAY)
    assert new and sig["title"] == "Tashkent water supply tender announced" and sig["country"] == "UZB"  # 国名在附件里
    assert [d["url"] for d in sig["docs"]] == ["https://owner.test/news/1", "https://owner.test/f/tor.docx", "https://owner.test/f/map.pdf"]
    assert "error" in sig["docs"][2] and "[附件]" in sig["excerpt"]
    hits = len(genv.web.hits)
    sig2, new = submit(genv.profile, genv.store, None, url="https://projects.test/x", title="Invitation for Bids: Kinshasa water",
                       text="Congo, Democratic Republic of. Kinshasa water supply rehabilitation. Bids are invited.", today=TODAY)
    assert new and len(genv.web.hits) == hits and sig2["country"] == "COD"  # 检索结果入库:带了正文就不再抓取


def _ds(handler):
    return Classifier(LLMConfig(provider="deepseek"), client=DeepSeekClient("sk-test", transport=httpx.MockTransport(handler), sleep=lambda s: None))


def test_manual_submit_uses_model_even_without_ledger(genv, monkeypatch):
    monkeypatch.setattr(manual, "merged_ledger", lambda profile: None)
    seen = []

    def handler(req):
        seen.append(json.loads(req.content))
        return _ds_resp(json.dumps({**OK, "sectors": ["wastewater"], "country": "Georgia", "ledger_links": []}, ensure_ascii=False))

    sig, new = submit(genv.profile, genv.store, _ds(handler), text="Invitation for bids: construction of a sewage pumping station "
                      "and pressure main in the capital city.", today=TODAY)
    assert "(暂无台账)" in seen[0]["messages"][0]["content"]
    assert sig["llm"]["sectors"] == ["wastewater"] and sig["country"] == "GEO"  # 文中没有国名时用模型判断的国别
    assert genv.store.submitted(sig["id"])["llm"]["_model"] == "deepseek-v4-pro"


# ---------------------------------------------------------------- 筛选、处理状态与保存的筛选


def _signal(sid, title, country, sectors, first="2026-09-24", category="tender", deadline=None, relevant=True, source="s1"):
    return {"id": sid, "country": country, "scope": GLOBAL, "source": source, "source_name": f"测试源{source[-1]}", "kind": "html_list",
            "key": sid, "url": f"https://ex.test/{sid}", "title": title, "published": None, "first_seen": f"{first}T01:00:00+00:00",
            "change": "new", "fields": {}, "rules": {"candidate": True, "stage": category}, "docs": [], "links": [], "excerpt": "",
            "llm": {"relevant": relevant, "category": category, "summary_zh": title, "sectors": list(sectors), "country": country,
                    "project": None, "funder": "World Bank" if source == "s1" else None, "amount": None,
                    "deadlines": [{"type": "bid_deadline", "date": deadline, "evidence": ""}] if deadline else [],
                    "ledger_links": [], "new_opportunity": False, "evidence": []}}


@pytest.fixture
def qstore(tmp_path):
    st = Store(tmp_path)
    st.append_signals([
        _signal("a1" * 8, "Bukhara wastewater plant", "UZB", ["wastewater"], deadline="2026-10-10"),
        _signal("b2" * 8, "Nairobi water supply", "KEN", ["water_supply"], first="2026-09-01", category="project_news", source="s2"),
        _signal("c3" * 8, "Lima desalination", "PER", ["desalination"], first="2026-06-01", deadline="2026-12-01"),
        _signal("d4" * 8, "Road works", "KEN", [], category="other", relevant=False),
    ])
    return st


def test_query_rows_and_filters(qstore):
    rows = load_rows(qstore, TODAY)
    a = next(r for r in rows if r["no"] == "a1a1a1a1")
    assert (a["country_zh"], a["continent"], a["region"], a["sectors_zh"], a["category_zh"]) == ("乌兹别克斯坦", "亚洲", "中亚", "污水", "招标")
    assert (a["deadline"], a["deadline_type"], a["status"], a["judged_by"], a["funder"]) == ("2026-10-10", "投标截止", "new", "模型", "World Bank")

    def nos(**kw):
        return [r["no"][:2] for r in apply(rows, Filter(**kw), TODAY)]

    assert nos() == ["a1", "b2"]  # 默认:最近 30 天、待处理、判定相关
    assert nos(days=None, continents=["美洲"]) == ["c3"]
    assert nos(days=None, sectors=["water_supply", "desalination"]) == ["b2", "c3"]
    assert nos(days=None, keyword="world BANK bukhara") == ["a1"]  # 空格分隔,全部命中,不分大小写
    assert nos(days=None, deadline_within=30) == ["a1"]
    assert nos(days=None, procurement_only=True) == ["a1", "c3"]
    assert nos(days=None, relevant_only=False, countries=["KEN"]) == ["d4", "b2"]
    assert nos(days=None, sources=["s2"], regions=["东非"]) == ["b2"]
    assert nos(days=0) == [] and nos(days=2) == ["a1"]  # 0 为"今天"(不是"不限")
    assert nos(days=None, keyword="C3C3C3C3") == ["c3"]  # 可按线索号搜(报告里引用的 [xxxxxxxx])


def test_triage_and_saved_views(qstore, tmp_path):
    set_status(qstore, ["a1" * 8], "done", "已报名")
    tri = (qstore.root / "triage.yaml").read_text(encoding="utf-8")
    assert tri.startswith("# 人工处理记录") and yaml.safe_load(tri) == {"a1a1a1a1": {"status": "done", "note": "已报名"}}
    rows = load_rows(qstore, TODAY)
    assert [r["no"][:2] for r in apply(rows, Filter(), TODAY)] == ["b2"]  # 处理过的不再出现在待处理里
    assert [(r["status"], r["note"]) for r in apply(rows, Filter(days=None, status=["done"]), TODAY)] == [("done", "已报名")]
    set_status(qstore, ["a1" * 8], "new")
    assert yaml.safe_load((qstore.root / "triage.yaml").read_text(encoding="utf-8")) is None
    f = Filter(days=90, regions=["中亚"], sectors=["wastewater"], keyword="Bukhara", deadline_within=60)
    save_view(tmp_path, "中亚污水", f)
    save_view(tmp_path, "全部", Filter(days=None, status=[]))
    views = load_views(tmp_path)
    assert list(views) == ["中亚污水", "全部"] and Filter.from_dict(views["中亚污水"]) == f
    assert Filter.from_dict({"days": 7, "bogus": 1}).days == 7  # 旧文件里多出的键忽略
    delete_view(tmp_path, "全部")
    assert list(load_views(tmp_path)) == ["中亚污水"]
    # 改名:位置与条件不变;重名、空名、不存在的方案报错
    save_view(tmp_path, "拉美海淡", Filter(days=None, continents=["美洲"], sectors=["desalination"]))
    save_view(tmp_path, "肯尼亚", Filter(countries=["KEN"]))
    rename_view(tmp_path, "拉美海淡", " 美洲海水淡化 ")
    views = load_views(tmp_path)
    assert list(views) == ["中亚污水", "美洲海水淡化", "肯尼亚"] and Filter.from_dict(views["美洲海水淡化"]).sectors == ["desalination"]
    for old, new, msg in (("肯尼亚", "中亚污水", "已经有"), ("肯尼亚", "  ", "不能为空"), ("没有这个", "x", "没有名为")):
        with pytest.raises(ValueError, match=msg):
            rename_view(tmp_path, old, new)
    save_view(tmp_path, "中亚污水", Filter(days=7, regions=["中亚"]))  # 用新条件覆盖:位置不变
    views = load_views(tmp_path)
    assert list(views) == ["中亚污水", "美洲海水淡化", "肯尼亚"] and views["中亚污水"]["days"] == 7
    assert same_filter(Filter(sectors=["a", "b"], status=["new", "done"]), Filter(sectors=["b", "a"], status=["done", "new"]))
    assert not same_filter(Filter(sectors=["a"]), Filter(sectors=["a"], days=7))


# ---------------------------------------------------------------- 报告


def test_reports_build_without_model(genv, qstore, monkeypatch):
    brief = build(genv.profile, qstore, Spec("brief", Filter(days=None, status=[], relevant_only=False)), TODAY)
    assert brief.startswith('---\ntitle: "水务项目机会线索简报"') and "# 摘要" in brief and "- 本期线索 4 条" in brief
    for h in ("# 第 1 章 报告说明", "## 1.2 信源与方法", "# 第 2 章 线索汇总", "## 2.1 按国别", "## 附录 A 本期线索清单"):
        assert h in brief
    assert "本报告未使用模型" in brief and "[a1a1a1a1]" in brief  # 线索号与分析段落的引用写法一致
    sector = build(genv.profile, qstore, Spec("sector", Filter(days=None, status=[], sectors=["wastewater", "desalination"])), TODAY)
    assert "领域:污水、海水淡化" in sector and "## 2.2 亚洲(1 条)" in sector and "## 2.3 美洲(1 条)" in sector
    country = build(genv.profile, qstore, Spec("country", Filter(days=None, status=[]), country="ZZZ"), TODAY)
    assert "# 第 2 章 国家概况与宏观经济" in country and "本期没有宏观指标快照" in country
    assert "ZZZ-001" in country and "甲市污水处理厂" in country and "未来 60 天" in country


def test_reports_count_done_but_not_ignored(genv, qstore):
    """报告统计"待处理 + 已处理",不含忽略的;筛选方案里存的"只看待处理"不影响报告。"""
    set_status(qstore, ["a1" * 8], "done", "已报名")
    set_status(qstore, ["c3" * 8], "ignore")
    md = build(genv.profile, qstore, Spec("brief", Filter(days=None, status=["new"])), TODAY)
    assert "[a1a1a1a1]" in md and "[b2b2b2b2]" in md and "[c3c3c3c3]" not in md and "- 本期线索 2 条" in md
    assert "统计待处理与已处理的线索,不含标为忽略(无关或重复)的" in md


def test_report_analysis_by_model(genv, qstore):
    seen = []
    fields = {"summary": "- 乌兹别克斯坦布哈拉污水厂招标 10 月 10 日截止 [a1a1a1a1]", "dynamics": "本期动态…… [a1a1a1a1, 99999999]",
              "opportunities": "- 布哈拉污水厂 [a1a1a1a1]", "actions": "- 10 月 1 日前决定是否投标", "risks": "- 资格条件待核实"}

    def handler(req):
        seen.append(json.loads(req.content))
        return _ds_resp(json.dumps(fields, ensure_ascii=False))

    md = build(genv.profile, qstore, Spec("brief", Filter(days=None, status=[])), TODAY, _ds(handler))
    user = seen[0]["messages"][1]["content"]
    assert "线索(共 3 条)" in user and "[a1a1a1a1] | 2026-09-24 | 乌兹别克斯坦 | 污水 | 招标" in user
    assert "不要编造" in seen[0]["messages"][0]["content"] and '"summary"' in seen[0]["messages"][0]["content"]
    assert fields["summary"] in md and "## 3.2 重点机会" in md and "分析段落由模型(deepseek-v4-pro(DeepSeek))" in md
    assert "[a1a1a1a1, 99999999(未核到)]" in md and "1 个引用编号在本期数据里没有找到" in md  # 模型编的引用会被标出
    failed = build(genv.profile, qstore, Spec("brief", Filter(days=None, status=[])), TODAY,
                   _ds(lambda req: httpx.Response(402, json={"error": {"message": "Insufficient Balance"}})))
    assert "模型分析未成功" in failed and "余额不足" in failed and "- 本期线索 3 条" in failed  # 失败时退回统计摘要


def test_word_output_takes_template_layout(tmp_path):
    pytest.importorskip("pypandoc")
    from docx import Document
    from docx.shared import Cm

    if not TEMPLATE.exists():
        pytest.skip("没有报告模板")
    md = tmp_path / "r.md"
    md.write_text('---\ntitle: "测试报告"\n---\n\n# 摘要\n\n- 要点一\n\n# 第 1 章 概况\n\n## 1.1 小节\n\n正文段落。\n\n'
                  "| 列一 | 列二 |\n|---|---|\n| 甲 | 乙 |\n", encoding="utf-8")
    out = to_docx_styled(md, tmp_path / "r.docx")
    spec = template_spec()
    d = Document(str(out))
    sec = d.sections[0]
    assert abs(sec.page_width - Cm(21)) < Cm(0.1) and sec.page_height > sec.page_width  # A4 竖版
    assert sec.left_margin == spec["page"]["left_margin"] and sec.top_margin == spec["page"]["top_margin"]
    assert d.styles["Normal"].font.size == spec["styles"]["Normal"]["size"]
    assert d.styles["Heading 1"].font.size == spec["styles"]["Heading 1"]["size"]
    assert len(d.tables) == 1 and d.tables[0].cell(1, 1).text == "乙"
