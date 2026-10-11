"""项目报告(core/project_report.py):在使用方模板上填 Word、来源按首次引用排号、挂到项目的线索进来源、
内部叠加的判断进报告而情报不进、模型写的来源编号逐个核对。用测试国台账(虚构),不联网。"""
from __future__ import annotations

import json
import re
from datetime import date
from pathlib import Path

import pytest
from docx import Document

from core import project_report as pr

pytestmark = pytest.mark.skipif(not pr.TEMPLATE.exists(), reason="使用方的 Word 模板在私有配置仓(private/templates);公开核心仓里没有")
from core.profile import load_profile
from core.watch import flow
from core.watch.store import Store

FX = Path(__file__).parent / "fixtures"
DAY = date(2026, 9, 29)
SECRET = "虚构的非公开信息"  # 内部叠加里的情报,不能出现在报告里


def _store(tmp_path) -> Store:
    s = Store(tmp_path / "data")
    s.append_signals([{
        "id": "c0ffee0012345678", "title": "Invitation for Bids: WWTP design-build", "url": "https://example.org/t/c0ff", "published": "2026-09-20",
        "first_seen": "2026-09-20T08:00:00+00:00", "source": "zz_notices", "source_name": "虚构采购网", "kind": "html_list", "country": "ZZZ",
        "links": [{"item": "ZZZ-001", "confidence": "high", "by": "rule"}],
        "llm": {"relevant": True, "category": "tender", "sectors": ["wastewater"], "summary_zh": "甲市污水厂设计施工招标", "deadlines": [],
                "ledger_links": []}}])
    return s


def _content(tmp_path, clf=None, org="测试部门") -> dict:
    prof = load_profile(FX / "profile.yaml")
    return pr.compose(pr.load_project(prof, _store(tmp_path), "ZZZ", "ZZZ-001", DAY), DAY, clf, org)


def test_money():
    assert pr.money(57_000_000, "EUR") == "5,700 万欧元"
    assert pr.money(410_000_000, "EUR") == "4.1 亿欧元"
    assert pr.money(1_138_000, "USD") == "113.8 万美元"
    assert pr.money(None, "USD") == "金额未公布"


def test_sources_are_numbered_by_first_citation(tmp_path):
    c = _content(tmp_path)
    nums = [n for n, _ in c["sources"]]
    assert nums == list(range(1, len(nums) + 1))
    blob = json.dumps({k: v for k, v in c.items() if k != "sources"}, ensure_ascii=False)
    assert "⟦" not in blob
    cited = {int(x) for x in re.findall(r"\[(\d+)\]", blob)}
    assert cited == set(nums)  # 文中用到的都登记了,登记的都用到了
    first = dict(c["core"])["项目"]
    assert first.endswith("[1]") and c["sources"][0][1]["publisher"] == "虚构银行" and c["sources"][0][1]["grade"] == "官方一手"
    by_title = {r["title"]: r for _, r in c["sources"]}
    assert by_title["虚构新闻"]["grade"] == "媒体与企业"  # C 级:媒体与企业
    lead = next(r for t, r in by_title.items() if "系统线索 c0ffee00" in t)
    assert lead["publisher"] == "虚构采购网" and lead["url"] == "https://example.org/t/c0ff" and c["leads"] == ["c0ffee00"]


def test_internal_judgement_in_but_intel_out(tmp_path):
    c = _content(tmp_path)
    core = dict(c["core"])
    assert "A 类" in core["建议"] and "设计施工总承包" in core["建议"] and "准备资审文件" in core["建议"]
    assert c["actions"][0][:2] == ("准备资审文件", "【待定】")
    assert SECRET not in json.dumps(c, ensure_ascii=False) and "虚构的国别级非公开信息" not in json.dumps(c, ensure_ascii=False)
    phases = {ph: (st, what) for ph, st, _, what, _ in c["phases"]}
    assert phases["施工招标"][0] == "进行中" and "系统采集" in phases["施工招标"][1]  # GPN 已发、资审截止未到,另有系统采集的招标公告
    assert phases["详细设计"][0] == "待核实"  # 没有信息的阶段不留空
    assert dict((k, v) for k, v in c["pros"])["采购"].startswith("EBRD 采购规则,各国企业均可")


def test_word_follows_the_template(tmp_path):
    prof = load_profile(FX / "profile.yaml")
    md, docx = pr.write(prof, _store(tmp_path), "ZZZ", "ZZZ-001", DAY, org="测试部门", out_dir=tmp_path / "out")
    assert flow.kind_of(docx.stem) == "项目报告" and "c0ffee00" in md.read_text(encoding="utf-8")  # 首页能按线索号关联
    d = Document(str(docx))
    text = "\n".join(p.text for p in d.paragraphs) + "\n" + "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
    assert d.paragraphs[0].text == "测试国甲市污水处理厂" and "填写说明" not in text
    left = set(re.findall(r"【[^】]*】", text))
    assert left <= {"【待定】"}, left  # 模板的占位都填上了(只剩责任方待定)
    head = "".join(p.text for p in d.sections[0].header.paragraphs)
    assert "甲市污水处理厂" in head and "中文" in head
    assert d.core_properties.author == "测试部门" and d.core_properties.title == "测试国甲市污水处理厂项目分析报告"
    progress = next(t for t in d.tables if t.rows[0].cells[0].text == "阶段")
    fills = {r.cells[0].text: r.cells[0]._tc.tcPr.find(f"{{{pr.W}}}shd").get(f"{{{pr.W}}}fill") for r in progress.rows[1:]}
    assert fills["融资签约"] == pr.SHADE["待核实"] and fills["施工招标"] == pr.SHADE["进行中"]
    srcs = next(t for t in d.tables if t.rows[0].cells[0].text == "编号")
    assert [r.cells[0].text for r in srcs.rows[1:]] == [f"[{i}]" for i in range(1, len(srcs.rows))]
    assert SECRET not in text


class _FakeModel:
    label = "测试模型"
    fatal = None

    def __init__(self):
        self.user = ""

    def write(self, system: str, user: str, fields: dict) -> dict:
        self.user = user
        n = next(line for line in user.splitlines() if line.startswith("来源编号:"))
        first = n.split("[")[1].split("]")[0]
        out = {k: f"模型写的{k}。" for k in fields}
        out["project"] = f"**甲市污水厂**,EBRD 贷款 5,000 万欧元 [{first}];另一说 [99]。"
        out["leader"] = "- 是否列为跟踪项目\n- 前期费用\n- 授权\n- 资源"
        return out


def test_model_draft_cites_are_checked(tmp_path):
    m = _FakeModel()
    c = _content(tmp_path, clf=m)
    core = dict(c["core"])
    assert core["项目"].startswith("甲市污水厂,EBRD 贷款 5,000 万欧元 [1]") and "[99]" not in core["项目"]  # 粗体去掉,编造的编号删去
    assert "1 处来源编号" in c["note"] and "测试模型" in c["note"]
    assert c["leader"] == ["是否列为跟踪项目", "前期费用", "授权", "资源"]
    assert SECRET not in m.user and "设计施工总承包" in m.user  # 模型看得到内部判断,看不到情报


def _edited_fixture(tmp_path) -> Path:
    """测试国台账的副本:ZZZ-001 加一条只写一句话的备注、一条只有内部来源的备注、一个咨询意向书截止;来源表加内部成稿与一条没用到的来源。"""
    import shutil

    import yaml

    prof = tmp_path / "prof"
    shutil.copytree(FX, prof)
    led_fp = prof / "ZZZ" / "ledger.yaml"
    led = yaml.safe_load(led_fp.read_text(encoding="utf-8"))
    it = led["opportunities"][0]
    it["notes"] = ["付费机制未公开(只写一句话的备注)", {"text": "工期 18 个月", "src": ["R01"]}]
    it["milestones"].append({"type": "eoi_deadline", "date": "2026-10-12", "status": "planned", "note": "监理咨询意向书", "src": ["S01"]})
    led_fp.write_text(yaml.safe_dump(led, allow_unicode=True, sort_keys=False), encoding="utf-8")
    src_fp = prof / "ZZZ" / "sources.yaml"
    srcs = yaml.safe_load(src_fp.read_text(encoding="utf-8"))
    srcs["sources"]["R01"] = {"title": "内部成稿", "publisher": "使用方内部成稿", "date": "2026-09", "grade": "B", "via": "report_baseline"}
    srcs["sources"]["S99"] = {"title": "与本项目无关的来源", "publisher": "别处", "date": "2026-01-01", "grade": "A", "via": "primary_read",
                              "url": "https://example.org/other"}
    src_fp.write_text(yaml.safe_dump(srcs, allow_unicode=True, sort_keys=False), encoding="utf-8")
    return prof / "profile.yaml"


def test_ledger_quirks_and_what_may_be_cited(tmp_path):
    """只写一句话的备注不出错;只有内部来源的事实标"待核实"、内部成稿不登记;咨询意向书截止归"咨询招标";
    标为忽略的线索、没有网页地址的人工提交(可能是内部材料)不当来源。"""
    s = _store(tmp_path)
    base = {"url": "", "published": "2026-09-21", "first_seen": "2026-09-21T08:00:00+00:00", "country": "ZZZ", "kind": "manual",
            "links": [{"item": "ZZZ-001", "confidence": "high", "by": "rule"}],
            "llm": {"relevant": True, "category": "tender", "sectors": ["wastewater"], "summary_zh": "……", "deadlines": [], "ledger_links": []}}
    s.append_signals([{**base, "id": "dead000011112222", "title": "我司投标分析(内部)", "source": "manual", "source_name": "人工提交"},
                      {**base, "id": "beef000011112222", "title": "重复公告", "url": "https://example.org/dup", "source": "zz_notices",
                       "source_name": "虚构采购网", "kind": "html_list"}])
    from core.watch.query import set_status

    set_status(s, ["beef000011112222"], "ignore")
    prof = load_profile(_edited_fixture(tmp_path))
    c = pr.compose(pr.load_project(prof, s, "ZZZ", "ZZZ-001", DAY), DAY, None, "测试部门")
    assert c["leads"] == ["c0ffee00"]  # 忽略的、人工粘贴或上传的都不在
    titles = [r["title"] for _, r in c["sources"]]
    assert not any("内部" in t or "重复公告" in t or "无关" in t for t in titles)
    period = dict((k, v) for k, v, _ in c["basics"])["工期"]
    assert period == ["工期 18 个月(待核实:仅有内部来源)"]
    phases = {ph: what for ph, _, _, what, _ in c["phases"]}
    assert "监理咨询意向书" in phases["咨询招标"] and "监理咨询意向书" not in phases["施工招标"]
    pr.write(prof, s, "ZZZ", "ZZZ-001", DAY, org="测试部门", out_dir=tmp_path / "out")  # 一句话备注不出错


def test_word_is_well_formed(tmp_path):
    """Word 里:书签不重复;底色 w:shd 在 tcMar、vAlign 之前;链接的字体设置按 color、sz、u 的顺序(Word 对顺序严格)。"""
    import zipfile

    prof = load_profile(FX / "profile.yaml")
    s = _store(tmp_path)
    s.append_signals([{"id": f"{i:x}c0ffee12345678"[:16], "title": f"Notice {i}", "url": f"https://example.org/n/{i}", "published": f"2026-09-1{i}",
                       "first_seen": f"2026-09-1{i}T08:00:00+00:00", "source": "zz_notices", "source_name": "虚构采购网", "kind": "html_list",
                       "country": "ZZZ", "links": [{"item": "ZZZ-001", "confidence": "high", "by": "rule"}],
                       "llm": {"relevant": True, "category": "tender", "sectors": ["wastewater"], "summary_zh": "……", "deadlines": [],
                               "ledger_links": []}} for i in range(1, 6)])
    _, docx = pr.write(prof, s, "ZZZ", "ZZZ-001", DAY, org="测试部门", out_dir=tmp_path / "out")
    xml = zipfile.ZipFile(docx).read("word/document.xml").decode("utf-8")
    srcs = next(t for t in Document(str(docx)).tables if t.rows[0].cells[0].text == "编号")
    assert len(srcs.rows) - 1 > 3  # 来源表比模板的 3 行多:有复制出来的行
    ids = re.findall(r'<w:bookmarkStart[^>]*w:id="(\d+)"', xml)
    names = re.findall(r'<w:bookmarkStart[^>]*w:name="([^"]+)"', xml)
    assert len(ids) == len(set(ids)) and len(names) == len(set(names))
    for tcpr in re.findall(r"<w:tcPr>(.*?)</w:tcPr>", xml, re.S):
        tags = re.findall(r"<w:(\w+)", tcpr)
        if "shd" in tags:
            assert all(tags.index("shd") < tags.index(t) for t in ("tcMar", "vAlign", "noWrap") if t in tags), tags
    for rpr in re.findall(r"<w:hyperlink [^>]*>\s*<w:r>\s*<w:rPr>(.*?)</w:rPr>", xml, re.S):
        assert re.findall(r"<w:(\w+)", rpr) == ["color", "sz", "u"]


def test_template_without_the_howto_box_and_without_sources(tmp_path):
    """模板里已经删了"填写说明"框也能出;没有可引用来源的项目,来源表整行写明,不留占位;Word 失败时不留 Markdown。"""
    from docx import Document

    tpl = tmp_path / "tpl.docx"
    d = Document(str(pr.TEMPLATE))
    box = next(t for t in d.tables if t.cell(0, 0).text.startswith("填写说明"))
    box._tbl.getparent().remove(box._tbl)
    d.save(str(tpl))
    prof = load_profile(FX / "profile.yaml")
    md, docx = pr.write(prof, _store(tmp_path), "ZZZ", "ZZZ-004", DAY, org="测试部门", out_dir=tmp_path / "out", template=tpl)
    out = Document(str(docx))
    text = "\n".join(c.text for t in out.tables for r in t.rows for c in r.cells)
    assert "(没有可引用的公开来源)" in text and not re.findall(r"【(?!待定)[^】]*】", text)
    broken = tmp_path / "broken.docx"
    Document().save(str(broken))  # 空文档:认不出模板的表
    try:
        pr.write(prof, _store(tmp_path / "b"), "ZZZ", "ZZZ-001", DAY, out_dir=tmp_path / "out2", template=broken)
    except ValueError as e:
        assert "找不到" in str(e)
    assert not (tmp_path / "out2").exists() or not list((tmp_path / "out2").glob("*.md"))


def test_model_sees_only_sources_used_by_the_facts(tmp_path):
    m = _FakeModel()
    prof = load_profile(_edited_fixture(tmp_path))
    pr.compose(pr.load_project(prof, _store(tmp_path), "ZZZ", "ZZZ-001", DAY), DAY, m, "测试部门")
    listed = next(line for line in m.user.splitlines() if line.startswith("来源编号:"))
    assert "与本项目无关的来源" not in listed and "内部成稿" not in listed and "虚构银行" in listed
