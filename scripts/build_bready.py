"""把世界银行 Business Ready(B-READY)的数据包整理成国别表要的一张表(A36 的数据文件)。
数据包从 https://www.worldbank.org/en/businessready 下载(B-READY_ALL_DATA_<年份>.zip,里面两个 Excel:
01_…PILLAR-TOPIC-SCORES.xlsx 各级得分,02_…EconomyAnswer.xlsx 各题的国别答案)。只取和水务、投标有关的约 20 项:
总分、供水得分、公共采购得分、纠纷解决得分、商业水价、企业供水可靠比例、接水天数、外企参标限制、大型工程授标天数、政府合同回款天数,
外加一张细项清单(监管机构定价权、漏损目标、污水须处理、须与本地企业联合、电子投标……)。
商业水价按当天汇率折成美元(core.fx;取不到用内置近似表)。中国香港、台湾不收(本系统不建这两处的行)。
用法:pip install openpyxl;python scripts/build_bready.py B-READY_ALL_DATA_2025.zip [--out private/data/bready.csv]
新一期出来后重跑一次、换掉数据文件,后台按数据版本整组替换。"""
from __future__ import annotations

import argparse
import csv
import io
import re
import statistics
import sys
import unicodedata
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from core.jobs.bready import COLS  # noqa: E402

SKIP = {"HKG", "TWN", "MAC"}
# 各经济体的本币(B-READY 的金额题都按本币填)
CUR = {
    "AGO": "AOA", "ARM": "AMD", "AZE": "AZN", "BEL": "EUR", "BEN": "XOF", "BFA": "XOF", "BGD": "BDT", "BGR": "BGN", "BHR": "BHD", "BIH": "BAM",
    "BRB": "BBD", "BTN": "BTN", "BWA": "BWP", "CAF": "XAF", "CAN": "CAD", "CIV": "XOF", "CMR": "XAF", "COD": "CDF", "COG": "XAF", "COL": "COP",
    "CPV": "CVE", "CRI": "CRC", "CYP": "EUR", "CZE": "CZK", "ECU": "USD", "ESP": "EUR", "EST": "EUR", "GBR": "GBP", "GEO": "GEL", "GHA": "GHS",
    "GMB": "GMD", "GNQ": "XAF", "GRC": "EUR", "HRV": "EUR", "HUN": "HUF", "IDN": "IDR", "IRL": "EUR", "ISL": "ISK", "ISR": "ILS", "ITA": "EUR",
    "JAM": "JMD", "JOR": "JOD", "KAZ": "KZT", "KGZ": "KGS", "KHM": "KHR", "KOR": "KRW", "LAO": "LAK", "LSO": "LSL", "LVA": "EUR", "MAR": "MAD",
    "MDA": "MDL", "MDG": "MGA", "MEX": "MXN", "MKD": "MKD", "MLI": "XOF", "MLT": "EUR", "MNE": "EUR", "MUS": "MUR", "MYS": "MYR", "NAM": "NAD",
    "NPL": "NPR", "NZL": "NZD", "PAK": "PKR", "PER": "PEN", "PHL": "PHP", "PNG": "PGK", "POL": "PLN", "PRT": "EUR", "PRY": "PYG", "PSE": "ILS",
    "ROU": "RON", "RWA": "RWF", "SEN": "XOF", "SGP": "SGD", "SLE": "SLE", "SLV": "USD", "SRB": "RSD", "SSD": "SSP", "SVK": "EUR", "SVN": "EUR",
    "SWE": "SEK", "SWZ": "SZL", "SYC": "SCR", "TCD": "XAF", "TGO": "XOF", "TJK": "TJS", "TKM": "TMT", "TLS": "USD", "TON": "TOP", "TTO": "TTD",
    "TUN": "TND", "TUR": "TRY", "TZA": "TZS", "URY": "UYU", "USA": "USD", "UZB": "UZS", "VNM": "VND", "VUT": "VUV", "WSM": "WST",
}
# 汇率接口与内置表都没有的币种:1 单位 = 多少人民币(近似)
EXTRA_CNY = {"BGN": 4.0, "BAM": 4.0, "BWP": 0.53, "NAD": 0.40, "LSL": 0.40, "SZL": 0.40, "RWF": 0.0049, "SLE": 0.31, "CDF": 0.0025, "AOA": 0.0078,
             "BHD": 19.0, "BBD": 3.57, "BTN": 0.082, "CVE": 0.071, "CRC": 0.014, "GMD": 0.10, "ISK": 0.052, "ILS": 1.95, "KRW": 0.0051, "LAK": 0.00033,
             "MDL": 0.40, "MGA": 0.0016, "MKD": 0.127, "MUR": 0.155, "NZD": 4.2, "SSP": 0.0016, "SEK": 0.74, "SCR": 0.50, "TTD": 1.05}
# 商业水价折成美元后的合理区间(每立方米):超出的多半是填错币种或单位,不收,细项里注明
TARIFF_OK = (0.03, 15.0)


def norm(name: str) -> str:
    s = unicodedata.normalize("NFKD", str(name)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]", "", s)


def _num(v) -> float | None:
    """数,或「Calculate median; 50% Skip Logic: 120」里的数;-777 = 没有数据。"""
    if isinstance(v, (int, float)):
        return None if v == -777 else float(v)
    m = re.search(r"(-?\d+(?:\.\d+)?)\s*$", str(v or ""))
    if not m or float(m.group(1)) == -777:
        return None
    return float(m.group(1))


def half(v) -> bool:
    return "skip logic" in str(v or "").lower()


def yes_no(v) -> str | None:
    """问卷答案 → 中文;没有答案(多数专家跳过)的返回 None。"""
    s = str(v or "").strip()
    low = s.lower()
    if not s or low == "skip logic":
        return None
    if low.startswith("multiple modes") or low.startswith("even split"):
        return "说法不一"
    tail = "(半数专家未答)" if "skip logic" in low else ""
    if low.startswith("yes, in all"):
        return "全部招标都有" + tail
    if low.startswith("yes, in some"):
        return "部分招标有" + tail
    if low.startswith("yes, fully"):
        return "是(全程线上)" + tail
    if low.startswith("yes, but"):
        return "部分(仍要纸件或到场)" + tail
    if low.startswith("yes, for all"):
        return "是(全部)" + tail
    if low.startswith("yes, except"):
        return "是(小额合同除外)" + tail
    if low.startswith("yes"):
        return "是" + tail
    if low.startswith("no"):
        return "否" + tail
    return None


FOREIGN = {"否": "无限制", "部分招标有": "部分招标有限制", "全部招标都有": "全部招标有限制", "说法不一": "说法不一"}

# 细项:(分组, 中文说法, 变量名, 哪张答案表);是非题
ITEMS = [
    ("供水", "供水监管机构有最终定价权", "UW_1_1_2", "03"),
    ("供水", "监管机构须制定服务质量标准", "UW_1_1_3_1", "03"),
    ("供水", "法规设了降低漏损的目标", "UW_3_1_5", "03"),
    ("供水", "未达漏损目标有经济处罚", "UW_3_1_7", "03"),
    ("供水", "污水须处理达标后才能排放", "UW_3_2_2", "03"),
    ("供水", "有污水回用(再生水)规定", "UW_3_2_4", "03"),
    ("供水", "要求企业节约用水", "UW_3_1_9", "03"),
    ("供水", "有节水的财政激励", "UW_3_1_13", "03"),
    ("供水", "现行水价在网上公布", "UW_5_1_1", "03"),
    ("供水", "每年在网上公布供水可靠性指标", "UW_4_3_2_R2", "03"),
    ("供水", "有供水管网数据库", "UW_6_1_1_1", "03"),
    ("供水", "可在网上全程申请接水", "UW_6_2_1", "03"),
    ("供水", "有独立于水司的投诉渠道", "UW_5_4_2_R2", "03"),
    ("采购", "外企参标或授标受限", "MCP_22", "09"),
    ("采购", "外企须与本地企业联合才能投标", "MCP_23", "09"),
    ("采购", "外企须在当地设子公司才能投标", "MCP_24", "09"),
    ("采购", "有专留给本地企业的合同", "MCP_24_1_R2", "09"),
    ("采购", "有全国统一的电子采购平台", "MCP_4", "09"),
    ("采购", "可在网上提交全部投标文件", "MCP_7_10_r_a5", "09"),
    ("采购", "可在网上看授标决定及理由", "MCP_7_10_r_a10", "09"),
    ("采购", "授标结果须通知全部投标人", "MCP_35", "09"),
    ("采购", "授标与签约之间有异议期", "MCP_36", "09"),
    ("采购", "法规规定收到发票后的付款时限", "MCP_37", "09"),
    ("采购", "鼓励按综合评分而非最低价评标", "MCP_31_2", "09"),
]


def _cell_text(v) -> str:
    return " ".join(str(v or "").split())


def read_zip(path: Path):
    import openpyxl  # 只有整理数据时要
    z = zipfile.ZipFile(path)
    names = z.namelist()
    sc = next(n for n in names if "SCORES" in n.upper())
    an = next(n for n in names if "ANSWER" in n.upper())
    return openpyxl.load_workbook(io.BytesIO(z.read(sc)), read_only=True, data_only=True), openpyxl.load_workbook(io.BytesIO(z.read(an)), read_only=True, data_only=True)


def edition(wb) -> str:
    about = [_cell_text(r[0]) for r in wb["ABOUT"].iter_rows(values_only=True) if r and r[0]] if "ABOUT" in wb.sheetnames else []
    year = next((m.group(1) for t in about for m in [re.search(r"Business Ready (\d{4})", t)] if m), "")
    ver = next((t.split(":", 1)[1].strip() for t in about if t.lower().startswith("file version")), "")
    return ("B-READY " + year).strip() + (f"(数据版本 {ver})" if ver else "")


def sheet(wb, prefix: str):
    return wb[next(n for n in wb.sheetnames if n.startswith(prefix))]


def topic(ws, cats: list[str] | None) -> dict[str, float]:
    """一个专题表 → {ISO3: 分};cats 给了就把这几个类别(各满分 33.3)相加,不给取专题总分。表头四行:支柱 / 类别 / 子类 / 列名。"""
    rows = list(ws.iter_rows(values_only=True))
    head, data = rows[:4], [r for r in rows[4:] if r and r[1]]
    if cats is None:
        cols = [2]
    else:
        cols = [j for j, h in enumerate(head[3]) if str(h or "").startswith("Category") and _cell_text(head[1][j]) in cats]
        if len(cols) != len(cats):
            raise SystemExit(f"{ws.title}:没找到类别 {cats}(表头变了?)")
    return {str(r[1]): sum(float(r[j] or 0) for j in cols) for r in data}


def build(zpath: Path, rates=None) -> tuple[list[dict], str]:
    sc, an = read_zip(zpath)
    ed = edition(sc)
    pil = {str(r[1]): r for r in sheet(sc, "00_").iter_rows(min_row=2, values_only=True) if r and r[1]}
    iso_of = {norm(r[0]): iso for iso, r in pil.items()}
    total = {iso: sum(float(x) for x in r[2:5]) / 3 for iso, r in pil.items()}
    water = topic(sheet(sc, "03_"), ["1.2 Water", "2.2 Water", "3.2 Water"])
    procure = topic(sheet(sc, "09_"), ["1.3 Bidding For Public Contracts", "2.3 E-Procurement", "3.3 Public Procurement"])
    dispute = topic(sheet(sc, "08_"), None)
    ans: dict[str, dict[str, object]] = {}
    for pre in ("03_", "09_"):
        for r in sheet(an, pre).iter_rows(min_row=2, values_only=True):
            if r and r[0] and r[1]:
                iso = iso_of.get(norm(r[0]))
                if iso:
                    ans.setdefault(iso, {})[str(r[1])] = r[3]
    if rates is None:
        from core import fx
        rates = fx.Rates(log=lambda *_: None)
    out = []
    for iso in sorted(pil):
        if iso in SKIP:
            continue
        a = ans.get(iso, {})
        items = []
        for grp, label, var, _ in ITEMS:
            v = yes_no(a.get(var))
            if v:
                items.append(f"[{grp}] {label}:{v}")
        cur = CUR.get(iso, "")
        bill = _num(a.get("UW_7_1_33_R2"))
        tariff = bill / 1000 if bill is not None and bill > 0 else None
        usd = None
        if tariff is not None and cur:
            usd = rates.to_usd(tariff, cur) if rates.rate(cur) else None
            if usd is None and cur in EXTRA_CNY and rates.rate("USD"):
                usd = tariff * EXTRA_CNY[cur] / rates.rate("USD")
        if usd is not None and not (TARIFF_OK[0] <= usd <= TARIFF_OK[1]):
            items.append(f"[供水] 商业水价:填报每月 1000 m³ 付 {bill:g} {cur},折合每立方米约 {usd:.3f} 美元,明显不合常理,未采用")
            tariff = usd = None
        small = _num(a.get("UW_7_1_34_R2"))
        if small is not None and small > 0 and cur:
            items.append(f"[供水] 小商户每月 20 m³ 水费+污水费:{small:,.6g} {cur}")
        conn = _num(a.get("UW_7_1_17_R2"))
        if conn is not None and conn > 0 and cur:
            items.append(f"[供水] 新接一个商业供水口的费用:{conn:,.6g} {cur}")
        award = _num(a.get("MCP_29_1"))
        if award is not None and half(a.get("MCP_29_1")):
            items.append(f"[采购] 大型工程授标天数 {award:g} 天是半数专家的估计(另一半未答)")
        rel = _num(a.get("in18"))
        fb = yes_no(a.get("MCP_22"))
        row = {
            "iso3": iso,
            "bready_total": round(total[iso], 1),
            "bready_water": round(water[iso], 1) if iso in water else None,
            "bready_procure": round(procure[iso], 1) if iso in procure else None,
            "bready_dispute": round(dispute[iso], 1) if iso in dispute else None,
            "water_tariff": round(tariff, 4) if tariff is not None else None,
            "water_tariff_cur": cur if tariff is not None else None,
            "water_tariff_usd": round(usd, 2) if usd is not None else None,
            "water_reliable": round(rel, 1) if rel is not None else None,
            "water_conn_days": _num(a.get("in5")),
            "foreign_bid": FOREIGN.get((fb or "").replace("(半数专家未答)", "")) if fb else None,
            "award_days": award,
            "gov_pay_days": round(_num(a.get("reg9")), 0) if _num(a.get("reg9")) is not None else None,
            "bready_items": "\n".join(items) or None,
            "bready_edition": ed,
        }
        out.append(row)
    return out, ed


def write(rows: list[dict], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=COLS)
        w.writeheader()
        for r in rows:
            w.writerow({k: ("" if r.get(k) is None else (f"{r[k]:g}" if isinstance(r[k], float) else r[k])) for k in COLS})


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(prog="python scripts/build_bready.py", description="整理 B-READY 数据包成国别表的数据文件")
    ap.add_argument("zip", type=Path, help="B-READY_ALL_DATA_<年份>.zip")
    ap.add_argument("--out", type=Path, default=ROOT / "private" / "data" / "bready.csv")
    a = ap.parse_args(argv)
    rows, ed = build(a.zip)
    write(rows, a.out)
    usd = [r["water_tariff_usd"] for r in rows if r["water_tariff_usd"] is not None]
    print(f"{ed}:{len(rows)} 个经济体 → {a.out};有商业水价的 {len(usd)} 个,中位数约 {statistics.median(usd):.2f} 美元/m³" if usd else f"{ed}:{len(rows)} 个经济体")
    return 0


if __name__ == "__main__":
    sys.exit(main())
