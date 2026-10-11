"""按模板出 Word 报告:国别报告、领域专题报告、线索简报(任意筛选或保存的视图)。

结构仿阿塞拜疆成稿:封面标题 → 摘要 → 第 N 章(N.M 节)→ 附录;样式、页面设置与页脚页码取自成稿模板
(private/templates/报告模板_阿塞拜疆成稿.docx,由 scripts/make_report_template.py 生成,不含成稿正文)。
数据章节(线索表、台账表、截止日历、宏观指标)由系统生成;分析段落(摘要、动态分析、机会研判、行动建议、风险)
由模型只基于本期数据撰写,每条结论标线索号或台账编号;没有模型时只出数据章节,摘要按统计生成。
"""
from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass, field, replace
from datetime import date, timedelta
from pathlib import Path

from core import geo
from core.country_facts import background_md, facts_line, load_facts, overview_md
from core.ledger import ELIG, MILESTONES, ROOT, STAGES, derive, load_ledger, load_rules, template_file
from core.markets import load_markets
from core.profile import Profile
from core.report import _background, _fit_columns, _nobreak_hyphens, cell, table
from core.watch.llm import Classifier
from core.watch.outputs import CAT_ZH, SECTOR_ZH, deadline_rows
from core.watch.query import COUNTED, Filter, apply, load_rows
from core.watch.store import Store

TEMPLATE = template_file("报告模板_阿塞拜疆成稿.docx", "国别报告模板.docx")
KINDS = {"country": "国别报告", "sector": "领域专题报告", "brief": "线索简报"}
FIELDS = {
    "summary": "摘要:4–8 条要点,Markdown 列表,每条一句话,最重要的机会与截止事项放在前面",
    "dynamics": "本期动态分析:300–600 字,按领域或国别归纳本期出现的采购、融资、项目进展与政策信号,指出趋势;可分段",
    "opportunities": "重点机会研判:最值得跟进的 3–8 个机会,Markdown 列表;每个写明是什么、为什么值得跟进、关键日期、资格限制或风险",
    "actions": "行动建议:按紧迫程度排序的具体行动,Markdown 列表,写明截止或建议完成时间",
    "risks": "风险提示:3–6 条,Markdown 列表,涉及投标资格、融资落实、竞争、政治经济与执行风险",
}
SYSTEM = """你是投标人国籍为 {bidder} 的水务工程承包与投资企业的市场分析师,为管理层撰写{kind}。
只根据用户给出的数据写:不要编造数字、项目、机构或日期;引用具体事实时在句末用方括号注明线索号或台账编号,如 [e81ade2d]、[AZE-006];
数据不足以下结论时直接说明。用简洁、专业的中文书面语,不写客套话。"""


@dataclass
class Spec:
    kind: str  # country | sector | brief
    filter: Filter
    title: str = ""
    country: str = ""  # 国别报告的国别 ISO3
    note: str = ""  # 写在"报告说明"里的补充说明(如所用的保存视图名)
    use_llm: bool = True
    extra: dict = field(default_factory=dict)


def _period(f: Filter, today: date) -> str:
    if not f.days:
        return f"截至 {today}"
    return f"{today - timedelta(days=f.days)} 至 {today}"


def _scope_text(spec: Spec) -> str:
    f = spec.filter
    parts = []
    if spec.country:
        parts.append(f"国别:{geo.name_zh(spec.country)}")
    if f.continents:
        parts.append("大洲:" + "、".join(f.continents))
    if f.markets:
        parts.append(spec.extra.get("markets_label", "经营区域") + ":" + "、".join("全部市场国别" if m == "*" else m for m in f.markets))
    if f.regions:
        parts.append("地理分区:" + "、".join(f.regions))
    if f.countries and not spec.country:
        parts.append("国别:" + "、".join(geo.name_zh(c) for c in f.countries))
    if f.sectors:
        parts.append("领域:" + "、".join(SECTOR_ZH.get(s, s) for s in f.sectors))
    if f.categories:
        parts.append("类别:" + "、".join(CAT_ZH.get(c, c) for c in f.categories))
    if f.keyword:
        parts.append(f"关键词:{f.keyword}")
    if f.procurement_only:
        parts.append("只含采购与融资公告")
    if f.relevant_only:
        parts.append("只含判定为相关的线索")
    return ";".join(parts) or "不限"


def _sig_table(rows: list[dict], with_country: bool = True, limit: int = 200) -> str:
    """线索表(A4 竖版放得下的列):线索号放在摘要末尾,与分析段落里的引用写法一致。"""
    head = ["发布", *(["国别"] if with_country else []), "类别", "公告/项目", "截止", "摘要"]
    body = [[r["published"] or r["first_seen"], *([r["country_zh"]] if with_country else []), r["category_zh"],
             f"[{cell(r['title'])[:70]}]({r['url']})" if r["url"] else cell(r["title"])[:70], r["deadline"] or "—",
             f"{r['summary']} [{r['no']}]"]
            for r in rows[:limit]]
    return table(head, body) + (f"\n\n另有 {len(rows) - limit} 条,见系统线索看板。" if len(rows) > limit else "")


def _ledger_part(profile: Profile, iso3: str, today: date) -> tuple[object | None, list, str]:
    """→ (台账, 派生结果, 台账项目一览表 Markdown)。"""
    c = next((x for x in profile.countries if x.iso3 == iso3 and x.ledger), None)
    if not c:
        return None, [], "该国尚无项目台账。"
    rules = load_rules(profile.rules) if profile.rules else load_rules()
    led = load_ledger(c.ledger, rules)
    ds = derive(led, rules, profile.bidder, profile.lead_times, today)
    body = []
    for d in ds:
        it, t = d.item, d.timing
        funders = "、".join(dict.fromkeys(str(f.get("funder")) for f in it.get("financing") or [])) or "—"
        node = (f"{MILESTONES.get(t.node.type, t.node.type)} {t.node.when.raw}" if t.node else "—") + (
            f"(最晚 {t.latest_start} 启动)" if t.latest_start else "")
        body.append([it["id"], it["name_zh"], STAGES.get(it.get("stage"), it.get("stage")), funders, node, ELIG[d.elig.status]])
    return led, ds, table(["编号", "项目", "阶段", "出资方", "下一节点", "投标资格"], body)


def load_indicators(data_dir: Path, iso3: str) -> dict | None:
    """最近一期宏观指标快照里该国的行(没有快照时 None)。"""
    snaps = sorted((Path(data_dir) / "snapshots").glob("*/country_indicator.parquet"))
    if not snaps:
        return None
    import duckdb

    rel = duckdb.sql(f"select * from '{snaps[-1]}' where iso3 = '{iso3}'")
    cols = rel.columns
    return {"as_of": snaps[-1].parent.name, "rows": [dict(zip(cols, r)) for r in rel.fetchall()]}


def _dataset(spec: Spec, rows: list[dict], ds: list, profile: Profile, today: date, facts: dict | None = None) -> str:
    """给模型的数据:线索、台账、截止、国别概况(纯文本,一行一条)。"""
    lines = [f"报告类型:{KINDS[spec.kind]};范围:{_scope_text(spec)};期间:{_period(spec.filter, today)};今天:{today};"
             f"投标人国籍:{(profile.bidder or {}).get('nationality') or '未设'}", "", f"线索(共 {len(rows)} 条):"]
    for r in rows[:150]:
        lines.append(" | ".join(x for x in [f"[{r['no']}]", r["published"] or r["first_seen"], r["country_zh"], r["sectors_zh"],
                                            r["category_zh"], r["title"][:120], f"截止 {r['deadline']}" if r["deadline"] else "",
                                            f"出资方 {r['funder']}" if r["funder"] else "", f"金额 {r['amount']}" if r["amount"] else "",
                                            f"台账 {r['links']}" if r["links"] else "", r["summary"][:200]] if x))
    if ds:
        lines += ["", f"台账项目(共 {len(ds)} 项):"]
        for d in ds:
            it, t = d.item, d.timing
            lines.append(" | ".join(x for x in [f"[{it['id']}]", it["name_zh"], STAGES.get(it.get("stage"), ""),
                                                "出资 " + "、".join(dict.fromkeys(str(f.get("funder")) for f in it.get("financing") or [])),
                                                f"下一节点 {t.node.type} {t.node.when.raw}" if t.node else "",
                                                f"最晚启动 {t.latest_start}" if t.latest_start else "", f"投标资格 {ELIG[d.elig.status]}"] if x))
    if facts:
        lines += ["", "国别概况(最新值,括号内为数据年份;来源:世行 WDI、IMF《世界经济展望》;引用时写\"据世行\"或\"据 IMF\",不要加方括号):"]
        lines += [facts_line(f) for f in facts.values() if f.values]
    return "\n".join(lines)


def _analysis(spec: Spec, rows: list[dict], ds: list, profile: Profile, clf: Classifier | None, today: date,
              facts: dict | None = None) -> tuple[dict, str]:
    """→ (分析段落, 说明)。没有模型或调用失败时返回空段落与原因。"""
    if not (spec.use_llm and clf is not None):
        return {}, "本报告未使用模型,分析段落从略;摘要按统计生成。"
    if not rows and not ds:
        return {}, "本期没有线索与台账数据,分析段落从略。"
    system = SYSTEM.format(bidder=(profile.bidder or {}).get("nationality") or "未设", kind=KINDS[spec.kind])
    out = clf.write(system, _dataset(spec, rows, ds, profile, today, facts), FIELDS)
    if not out or out.get("_error"):
        why = (out or {}).get("_error") or clf.fatal or ";".join(clf.errors[-1:]) or "模型未返回"
        return {}, f"模型分析未成功({why}),分析段落从略。"
    out, bad = _check_cites(out, {r["no"] for r in rows[:150]} | {d.item["id"] for d in ds})
    note = f"分析段落由模型({clf.label})根据本期数据撰写,方括号内为线索号或台账编号,请以原文为准。"
    if bad:
        note += f"其中 {len(bad)} 个引用编号在本期数据里没有找到,已标注\"未核到\",相关表述须核实。"
    return out, note


_CITE = re.compile(r"\[([^\[\]]{8,80})\]")
_CITE_ID = re.compile(r"\b(?:[0-9a-f]{8}|[A-Z]{3}-\d{3})\b")


def _check_cites(ana: dict, known: set[str]) -> tuple[dict, list[str]]:
    """模型段落里方括号中的线索号、台账编号逐个核对:本期数据里没有的标"未核到"(防止模型编造引用)。"""
    bad: list[str] = []

    def mark(m: re.Match) -> str:
        ids = _CITE_ID.findall(m.group(1))
        miss = [i for i in ids if i not in known]
        bad.extend(miss)
        return m.group(0) if not miss else "[" + _CITE_ID.sub(lambda x: x.group(0) + ("(未核到)" if x.group(0) in miss else ""), m.group(1)) + "]"

    fixed = {k: _CITE.sub(mark, v) for k, v in ana.items()}
    return fixed, sorted(set(bad))


def _stat_summary(rows: list[dict], dl: list[list]) -> str:
    if not rows:
        return "- 本期范围内没有线索。"
    by_cat = Counter(r["category_zh"] for r in rows)
    by_cty = Counter(r["country_zh"] for r in rows)
    by_sec = Counter(s for r in rows for s in r["sectors_zh"].split("、"))
    out = [f"- 本期线索 {len(rows)} 条,其中采购与融资公告 {sum(r['procurement'] for r in rows)} 条。",
           "- 按类别:" + ";".join(f"{k} {v} 条" for k, v in by_cat.most_common()),
           "- 按领域:" + ";".join(f"{k} {v} 条" for k, v in by_sec.most_common()),
           "- 按国别:" + ";".join(f"{k} {v} 条" for k, v in by_cty.most_common(10))]
    if dl:
        out.append("- 最近的截止:" + ";".join(f"{r[0]} {_plain(r[3])[:40]}" for r in dl[:5]))
    return "\n".join(out)


def _plain(md: object) -> str:
    """Markdown 链接 [文字](网址) → 文字。"""
    return re.sub(r"\[([^]]*)\]\([^)]*\)", r"\1", str(md))


def build(profile: Profile, store: Store, spec: Spec, today: date, clf: Classifier | None = None) -> str:
    """→ 报告 Markdown(pandoc 格式:YAML 标题块 + # 为章、## 为节)。"""
    f = replace(spec.filter, countries=[spec.country]) if spec.kind == "country" and spec.country else spec.filter
    f = replace(f, status=list(COUNTED))  # 报告统计待处理与已处理的线索,不含忽略的;筛选方案里存的"只看待处理"不适用于报告
    markets = load_markets(profile.markets)
    if markets:
        spec = replace(spec, extra={"markets_label": markets.label, **spec.extra})
    rows = apply(load_rows(store, today, markets=markets.region_map() if markets else None), f, today)
    led, ds, ledger_md = _ledger_part(profile, spec.country, today) if spec.kind == "country" else (None, [], "")
    dl = deadline_rows(led, ds, [], today, 60) if led else []
    dl += [[r["deadline"], (date.fromisoformat(r["deadline"]) - today).days, r["deadline_type"] or "截止",
            f"[{cell(r['title'])[:50]}]({r['url']})" if r["url"] else cell(r["title"])[:50], "—", f"线索 {r['no']}"]
           for r in rows if r["deadline"] and r["deadline"] <= (today + timedelta(days=60)).isoformat()]
    dl.sort(key=lambda r: r[0])
    # 国别概况:国别报告用本国;简报与专题用线索最多的前 15 个国别(报告附录与模型数据共用)
    top = [spec.country] if spec.kind == "country" and spec.country else [
        c for c, _ in Counter(r["country"] for r in rows if r["country"]).most_common(15)]
    facts = {k: v for k, v in load_facts(store.root.parent, markets, top).items() if v.values or v.basics} if top else {}
    ana, ana_note = _analysis(spec, rows, ds, profile, clf, today, facts)
    scope = geo.name_zh(spec.country) if spec.country else _scope_text(spec)
    title = spec.title or {"country": f"{scope}水务市场动态与项目机会报告", "sector": f"{scope}专题报告",
                           "brief": "水务项目机会线索简报"}[spec.kind]
    L = ["---", f'title: "{title}"', f'subtitle: "{_period(f, today)} · Water Tender Watch 自动生成"', f'date: "{today}"', "---", ""]
    L += ["# 摘要\n", ana.get("summary") or _stat_summary(rows, dl), ""]
    ch = iter(range(1, 20))

    def chapter(t: str) -> tuple[int, str]:
        n = next(ch)
        return n, f"# 第 {n} 章 {t}\n"

    n, h = chapter("报告说明")
    srcs = Counter(r["source_name"] for r in rows)
    L += [h, f"## {n}.1 范围与期间\n", f"范围:{_scope_text(spec)}。期间:{_period(f, today)}(按线索{'发布' if f.date_field == 'published' else '入库'}日期)。"
          "统计待处理与已处理的线索,不含标为忽略(无关或重复)的。" + (f"{spec.note}" if spec.note else ""), "",
          f"## {n}.2 信源与方法\n",
          "线索由系统每天从多边开发银行采购与项目接口、业主网站公告栏、当地招标栏目与新闻源自动采集,经多语种关键词初筛后由大模型判定"
          "相关性、类别、领域、国别并抽取截止日、金额、出资方等字段;项目台账由人工核实维护。截止日、资格条件等以原文为准。", "",
          ("本期线索来源:" + ";".join(f"{k} {v} 条" for k, v in srcs.most_common()) + "。") if srcs else "本期没有线索。", "",
          ana_note, ""]
    if spec.kind == "country":
        n, h = chapter("国家概况与宏观经济")
        f = facts.get(spec.country)
        bg = background_md(f, table) if f else _background(load_indicators(store.root.parent, spec.country))
        L += [h, *(bg or ["本期没有宏观指标快照(每月 1 日自动更新)。\n"])]
        n, h = chapter("本期市场动态")
        proc = [r for r in rows if r["procurement"]]
        other = [r for r in rows if not r["procurement"]]
        L += [h, f"## {n}.1 概述\n", ana.get("dynamics") or _stat_summary(rows, dl), "",
              f"## {n}.2 采购与融资公告\n", _sig_table(proc, with_country=False) if proc else "本期没有。", "",
              f"## {n}.3 项目与政策动态\n", _sig_table(other, with_country=False) if other else "本期没有。", ""]
        n, h = chapter("项目机会清单")
        L += [h, f"## {n}.1 台账项目一览\n", ledger_md, "",
              *([f"## {n}.2 重点机会研判\n", ana["opportunities"], ""] if ana.get("opportunities") else [])]
        n, h = chapter("截止事项与行动建议")
        L += [h, f"## {n}.1 未来 60 天截止日历\n",
              table(["日期", "剩余天数", "事项", "项目/线索", "投标资格", "来源"], dl) if dl else "未来 60 天没有截止事项。", "",
              *([f"## {n}.2 行动建议\n", ana["actions"], ""] if ana.get("actions") else [])]
        if ana.get("risks"):
            n, h = chapter("风险提示")
            L += [h, ana["risks"], ""]
    elif spec.kind == "sector":
        n, h = chapter("分区域动态")
        L += [h, f"## {n}.1 概述\n", ana.get("dynamics") or _stat_summary(rows, dl), ""]
        if markets:  # 按使用方的经营区域分述,表外国别另列;没有市场分布时按大洲
            parts = [(g, [r for r in rows if r["market"] == g]) for g in markets.regions]
            parts.append(("市场外国别", [r for r in rows if not r["market"]]))
        else:
            parts = [(c, [r for r in rows if r["continent"] == c]) for c in geo.CONTINENTS]
        for i, (name, sub) in enumerate([x for x in parts if x[1]], 2):
            L += [f"## {n}.{i} {name}({len(sub)} 条)\n", _sig_table(sub), ""]
        n, h = chapter("重点机会")
        proc = [r for r in rows if r["procurement"]]
        L += [h, f"## {n}.1 采购与融资公告\n", _sig_table(proc) if proc else "本期没有。", "",
              *([f"## {n}.2 机会研判\n", ana["opportunities"], ""] if ana.get("opportunities") else [])]
        if ana.get("actions") or ana.get("risks"):
            n, h = chapter("行动建议与风险")
            L += [h, *([f"## {n}.1 行动建议\n", ana["actions"], ""] if ana.get("actions") else []),
                  *([f"## {n}.2 风险提示\n", ana["risks"], ""] if ana.get("risks") else [])]
    else:
        n, h = chapter("线索汇总")
        by_c: dict[str, list[dict]] = {}
        for r in rows:
            by_c.setdefault(r["country_zh"], []).append(r)
        sec_count = Counter(s for r in rows for s in r["sectors_zh"].split("、"))
        L += [h, f"## {n}.1 按国别\n",
              table(["国别", "大洲", "线索", "采购与融资", "最近截止"],
                    [[c, rs[0]["continent"], len(rs), sum(r["procurement"] for r in rs), min((r["deadline"] for r in rs if r["deadline"]), default="—")]
                     for c, rs in sorted(by_c.items(), key=lambda x: -len(x[1]))]) if rows else "本期没有。", "",
              f"## {n}.2 按领域\n", table(["领域", "线索"], [[k, v] for k, v in sec_count.most_common()]) if rows else "本期没有。", "",
              f"## {n}.3 线索明细\n", _sig_table(rows) if rows else "本期没有。", ""]
        if ana:
            n, h = chapter("分析与建议")
            L += [h, *([f"## {n}.1 动态分析\n", ana["dynamics"], ""] if ana.get("dynamics") else []),
                  *([f"## {n}.2 重点机会\n", ana["opportunities"], ""] if ana.get("opportunities") else []),
                  *([f"## {n}.3 行动建议\n", ana["actions"], ""] if ana.get("actions") else []),
                  *([f"## {n}.4 风险提示\n", ana["risks"], ""] if ana.get("risks") else [])]
    L += ["# 附录\n", "## 附录 A 本期线索清单\n", _sig_table(rows, limit=400) if rows else "本期没有。", ""]
    if spec.kind != "country" and facts:
        as_of = next(iter(facts.values())).as_of
        L += ["## 附录 B 涉及国别概况\n", overview_md(list(facts.values()), table), "",
              f"本期线索最多的前 15 个国别中有数据的(数据覆盖台账国别与市场分布表里的国别)。指标取自 {as_of or '最近一期'} 的数据快照"
              "(世行 WDI、IMF《世界经济展望》接口),括号内为数据年份;各国完整指标与来源链接见本机界面\"国别概况\"页。", ""]
    return "\n".join(L) + "\n"


KPI_CAPTION = "关键指标"  # 表格标题是它的,按关键指标卡排:名称行小字、数值行大号加粗、不要框线(core.jobs.reports.kpi_table)
BAD_MARK = "超警戒线"  # 表格里带这几个字的格子整格标红(负面指标;使用方 2026-10-10:不用 ▼ 这种看不懂的符号,直接写明)

STYLE_MAP = {  # 输出文档的样式 ← 模板里对应的样式
    "Normal": "Normal", "Body Text": "Normal", "First Paragraph": "Normal", "Heading 1": "Heading 1", "Heading 2": "Heading 2",
    "Heading 3": "Heading 3", "Title": "Title",
}


def template_spec(template: Path = TEMPLATE) -> dict:
    """读模板(成稿)的版式:页面尺寸与页边距、各样式的字体、字号、加粗、段前段后、行距、首行缩进。"""
    from docx import Document
    from docx.oxml.ns import qn

    spec: dict = {"styles": {}}
    if not Path(template).exists():
        return spec
    d = Document(str(template))
    sec = d.sections[-1]
    spec["page"] = {k: getattr(sec, k) for k in ("page_width", "page_height", "top_margin", "bottom_margin", "left_margin", "right_margin")}
    for name in set(STYLE_MAP.values()):
        try:
            st = d.styles[name]
        except KeyError:
            continue
        rpr = st.element.rPr
        fonts = rpr.find(qn("w:rFonts")) if rpr is not None else None
        pf = st.paragraph_format
        spec["styles"][name] = {
            "ascii": fonts.get(qn("w:ascii")) if fonts is not None else None,
            "east": fonts.get(qn("w:eastAsia")) if fonts is not None else None,
            "size": st.font.size, "bold": st.font.bold, "before": pf.space_before, "after": pf.space_after,
            "line": pf.line_spacing, "first": pf.first_line_indent, "align": pf.alignment,
        }
    # 样式没写的继承 Normal
    base = spec["styles"].get("Normal", {})
    for name, v in spec["styles"].items():
        for k, x in base.items():
            if v.get(k) is None and k in ("ascii", "east", "size", "line") and name != "Normal":
                v[k] = x
    return spec


def to_docx_styled(md_path: Path, docx_path: Path, template: Path = TEMPLATE) -> Path:
    """Markdown → Word:pandoc 出结构,再按模板(成稿)的版式设样式、页面与页脚页码;表格统一版式;封面单独成页。
    (直接拿成稿作 pandoc 的 reference-doc,在部分软件里表格会错位,所以只取其版式参数。)"""
    import pypandoc

    pypandoc.convert_file(str(md_path), "docx", format="markdown", outputfile=str(docx_path), extra_args=["--quiet", "-M", "lang=zh-CN"])
    _polish(docx_path, template_spec(template))
    return docx_path


def _set_fonts(style, ascii_: str | None, east: str | None) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for k in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):  # 主题字体会盖过具体字体,去掉
        fonts.attrib.pop(qn(f"w:{k}"), None)
    if ascii_:
        fonts.set(qn("w:ascii"), ascii_)
        fonts.set(qn("w:hAnsi"), ascii_)
    if east:
        fonts.set(qn("w:eastAsia"), east)


def _page_number_footer(section) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def _polish(path: Path, spec: dict) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from docx.text.run import Run

    doc = Document(str(path))
    names = {s.name for s in doc.styles}
    for out_name, tpl_name in STYLE_MAP.items():
        v = spec.get("styles", {}).get(tpl_name)
        if not v or out_name not in names:
            continue
        st = doc.styles[out_name]
        _set_fonts(st, v["ascii"], v["east"])
        st.font.size = v["size"]
        st.font.bold = v["bold"]
        st.font.color.rgb = RGBColor(0, 0, 0)
        pf = st.paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing, pf.first_line_indent = v["before"], v["after"], v["line"], v["first"]
        if out_name == "Title":
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for name in ("Subtitle", "Date", "Author"):
        if name in names:
            st = doc.styles[name]
            st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            st.paragraph_format.first_line_indent = Cm(0)
            st.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    if "Hyperlink" in names:  # 链接:蓝字下划线,看得出能点(模板没设链接样式时也一样)
        st = doc.styles["Hyperlink"]
        st.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
        st.font.underline = True
    if "Compact" in names:  # 表格与紧凑列表:不缩进、行距紧
        normal = spec.get("styles", {}).get("Normal", {})
        st = doc.styles["Compact"]
        _set_fonts(st, normal.get("ascii"), normal.get("east"))
        pf = st.paragraph_format
        pf.first_line_indent, pf.space_before, pf.space_after, pf.line_spacing = Cm(0), Pt(0), Pt(2), 1.15
    page = spec.get("page")
    for sec in doc.sections:
        if page:
            for k, v in page.items():
                setattr(sec, k, v)
        else:
            sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
            for k in ("left_margin", "right_margin", "top_margin", "bottom_margin"):  # pandoc 的版面不带页边距
                if getattr(sec, k) is None:
                    setattr(sec, k, Cm(2.5))
        _page_number_footer(sec)
    first_h1 = next((p for p in doc.paragraphs if p.style.name == "Heading 1"), None)
    if first_h1 is not None:  # 封面(标题、副标题、日期)单独一页
        first_h1.paragraph_format.page_break_before = True
    for p in doc.paragraphs:  # 成稿式的「摘要」「第 N 章」「附录」各起一页
        if p.style.name == "Heading 1" and re.match(r"^(摘要|第 \d+ 章|附录)", p.text.strip()):
            p.paragraph_format.page_break_before = True
    sec = doc.sections[-1]
    text_cm = (sec.page_width - sec.left_margin - sec.right_margin) / 360000
    after_borders = {qn(f"w:{t}") for t in ("shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription")}
    for tb in doc.tables:
        tblpr = tb._tbl.tblPr
        for old in tblpr.findall(qn("w:tblBorders")):
            tblpr.remove(old)
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), "8EAADB")
            borders.append(el)
        nxt = next((c for c in tblpr if c.tag in after_borders), None)  # OOXML 规定 tblBorders 在这些元素之前
        if nxt is None:
            tblpr.append(borders)
        else:
            nxt.addprevious(borders)
        _fit_columns(tb, text_cm, Cm, qn)
        for i, row in enumerate(tb.rows):
            for c in row.cells:
                if i == 0:
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:fill"), "D9E2F3")
                    c._tc.get_or_add_tcPr().append(shd)
                for p in c.paragraphs:
                    for r in (Run(e, p) for e in p._p.iter(qn("w:r"))):  # 连链接里的字(w:hyperlink 下的 run)一起
                        r.font.size = Pt(9)
                        if i == 0:
                            r.font.bold = True
                        _nobreak_hyphens(r, OxmlElement, qn)
    _style_kpis(doc, OxmlElement, qn, Pt, RGBColor)
    doc.save(str(path))


def _style_kpis(doc, OxmlElement, qn, Pt, RGBColor) -> None:
    """负面标红;关键指标卡:去掉框线与表头底色,名称行 9 磅灰字,数值行 18 磅加粗深蓝,负面的标红;表格标题段落删掉(上面已有小标题)。"""
    from docx.text.run import Run

    red, navy, grey = RGBColor(0xC0, 0x00, 0x00), RGBColor(0x1F, 0x38, 0x64), RGBColor(0x59, 0x59, 0x59)
    runs = lambda c: [Run(e, p) for p in c.paragraphs for e in p._p.iter(qn("w:r"))]  # noqa: E731
    for tb in doc.tables:
        cap = tb._tbl.tblPr.find(qn("w:tblCaption"))
        is_kpi = cap is not None and cap.get(qn("w:val")) == KPI_CAPTION
        bad_cols: set[int] = set()
        for i, row in enumerate(tb.rows):
            if is_kpi and i % 2 == 0:  # 名称行写明「超警戒线:…」,下面数值行同一列跟着标红
                bad_cols = {j for j, c in enumerate(row.cells) if BAD_MARK in c.text}
            for j, c in enumerate(row.cells):
                bad = BAD_MARK in c.text or (is_kpi and j in bad_cols)
                if is_kpi:
                    tcpr = c._tc.get_or_add_tcPr()
                    for old in tcpr.findall(qn("w:shd")):
                        tcpr.remove(old)
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:fill"), "F3F6FB")
                    tcpr.append(shd)
                for r in runs(c):
                    if is_kpi:
                        label = i % 2 == 0
                        unit = "Unit" in str(r.style.style_id if r.style is not None else "")  # 数值后的单位(KPI Unit 样式)小一号
                        r.font.size = Pt(9) if label else (Pt(10) if unit else Pt(18))
                        r.font.bold = not label
                        r.font.color.rgb = red if bad else (grey if label else navy)
                    elif bad:
                        r.font.color.rgb = red
        if is_kpi:
            borders = tb._tbl.tblPr.find(qn("w:tblBorders"))
            if borders is not None:
                for el in borders:
                    el.set(qn("w:val"), "single")
                    el.set(qn("w:color"), "FFFFFF")
                    el.set(qn("w:sz"), "12")
            n = len(tb.columns) or 1
            width = int((doc.sections[-1].page_width - doc.sections[-1].left_margin - doc.sections[-1].right_margin) / n)
            for row in tb.rows:  # 各卡一样宽(单元格与表格网格都要改,不同软件各读一处)
                for c in row.cells:
                    c.width = width
            grid = tb._tbl.find(qn("w:tblGrid"))
            for col in grid.findall(qn("w:gridCol")) if grid is not None else []:
                col.set(qn("w:w"), str(int(width / 635)))
            for sib in (tb._tbl.getprevious(), tb._tbl.getnext()):  # 表格标题段落(样式 TableCaption)
                style = sib.find(f"{qn('w:pPr')}/{qn('w:pStyle')}") if sib is not None and sib.tag == qn("w:p") else None
                if style is not None and "Caption" in str(style.get(qn("w:val"))) and KPI_CAPTION in "".join(sib.itertext()):
                    sib.getparent().remove(sib)


def slug(s: str) -> str:
    return re.sub(r"[^\w一-鿿-]+", "_", s).strip("_")[:40]


def write(profile: Profile, store: Store, spec: Spec, today: date, clf: Classifier | None = None,
          out_dir: Path | None = None) -> list[Path]:
    """生成报告 → [md, docx],存在数据目录 reports/custom/(或 out_dir)。"""
    md = build(profile, store, spec, today, clf)
    out_dir = out_dir or store.root.parent / "reports" / "custom"
    out_dir.mkdir(parents=True, exist_ok=True)
    f = spec.filter
    short = (geo.name_zh(spec.country) if spec.country else "、".join(SECTOR_ZH.get(x, x) for x in f.sectors)
             or "、".join(geo.name_zh(c) for c in f.countries[:3]) or "、".join(f.continents or f.regions) or "")
    name = slug(spec.title or "_".join(x for x in (KINDS[spec.kind], short) if x))
    fp = out_dir / f"{name}_{today}.md"
    fp.write_text(md, encoding="utf-8")
    return [fp, to_docx_styled(fp, fp.with_suffix(".docx"))]
