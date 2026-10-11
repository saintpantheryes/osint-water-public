"""国别报告渲染:分析单位是国别,正文以项目为粒度。Markdown 为源,pandoc 同源出 Word。

只读台账(公开事实)、规则表、profile(投标人国籍与内部周期)和 overlay(私有叠加);
public=True 时只输出公开事实:不含 overlay、不含针对投标人的资格结论与最晚启动日期。
"""
from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from core.ledger import (
    ELIG, FUNDING_STATUS, GRADES, INSTRUMENTS, MILESTONES, MS_STATUS, NATIONALITY, PKG_STATUS, PKG_TYPES, SEGMENTS, STAGES,
    VERIFY_STATUS, VIA,
    Derived, Ledger, cited, parse_pdate,
)
from core.profile import Country, Profile

CUR_ZH = {"USD": "美元", "EUR": "欧元", "AZN": "马纳特", "SAR": "沙特里亚尔", "KWD": "科威特第纳尔", "JPY": "日元", "CNY": "人民币"}
STAGE_ORDER = list(STAGES)
# 报告里的排序:离投标越近越靠前;已定标/在建只剩分包与供货,排在规划之后
REPORT_ORDER = ["procurement", "financing_signed", "preparation", "plan", "awarded", "construction", "on_hold", "completed", "cancelled"]


def _num(x: float, nd: int) -> str:
    s = f"{x:,.{nd}f}"
    return s.rstrip("0").rstrip(".") if "." in s else s


def money(amount: float | None, currency: str | None) -> str:
    if amount is None:
        return "未公开"
    cur = CUR_ZH.get(str(currency).upper(), str(currency))
    a = float(amount)
    if a >= 1e8:
        return f"{_num(a / 1e8, 2)} 亿{cur}"
    if a >= 1e4:
        return f"{_num(a / 1e4, 0 if a >= 1e6 else 1)} 万{cur}"
    return f"{_num(a, 0)} {cur}"


def usd(v: float | None) -> str:
    return "未公开" if v is None else "约 " + money(v, "USD")


def cell(x: object) -> str:
    s = "" if x is None else str(x)
    return s.replace("|", "/").replace("\n", " ").strip() or "—"


def refs(ids: list | None) -> str:
    return ", ".join(ids) if ids else "—"


def _w(s: str) -> int:
    return sum(2 if ord(ch) > 0x2E80 else 1 for ch in s)


def table(head: list[str], body: list[list[object]]) -> str:
    """管道表。分隔行的短横数按列内容宽度给,pandoc 转 Word 时据此分配列宽。"""
    cells = [[cell(c) for c in r] for r in body]
    widths = [max([_w(h)] + [_w(r[i]) for r in cells if i < len(r)]) for i, h in enumerate(head)]
    out = ["| " + " | ".join(head) + " |", "|" + "|".join("-" * max(4, min(w, 32)) for w in widths) + "|"]
    out += ["| " + " | ".join(r) + " |" for r in cells]
    return "\n".join(out)


def _node_text(d: Derived) -> str:
    n = d.timing.node
    if not n:
        return "—"
    return f"{MILESTONES.get(n.type, n.type)} {n.when.raw}({MS_STATUS.get(n.status, n.status)})"


def _latest_text(d: Derived) -> str:
    t = d.timing
    if t.latest_start is None:
        return "—"
    if t.days_left is not None and t.days_left < 0:
        return f"{t.latest_start}(已过 {-t.days_left} 天)"
    return f"{t.latest_start}(剩 {t.days_left} 天)"


def _pri(v: object) -> tuple[int, float]:
    """内部优先级 → 排序键:表格里定的 A/B/C(项目表人管列,拉回 overlay)在前,手写的数字(1 最先)随后,没写的最后。"""
    s = str(v).strip().upper() if v is not None else ""
    if s in ("A", "B", "C"):
        return 0, "ABC".index(s)
    try:
        return 1, float(s)
    except ValueError:
        return 2, 0


def _order(ds: list[Derived], overlay: dict) -> list[Derived]:
    pri = {k: _pri((v or {}).get("priority")) for k, v in (overlay.get("items") or {}).items()}
    return sorted(
        ds,
        key=lambda d: (
            not d.item.get("in_scope", True),
            pri.get(d.item["id"], (2, 0)),
            REPORT_ORDER.index(d.item["stage"]) if d.item.get("stage") in REPORT_ORDER else 99,
            -(d.size_usd or 0),
        ),
    )


def render_country(
    profile: Profile, country: Country, led: Ledger, rules: dict, ds: list[Derived], overlay: dict, today: date,
    public: bool = False, indicators: dict | None = None, facts=None,
) -> str:
    ov_items = {} if public else (overlay.get("items") or {})
    ds = _order(ds, {} if public else overlay)
    nat = (profile.bidder or {}).get("nationality") or "未设"
    core = [d for d in ds if d.item.get("in_scope", True)]
    L: list[str] = []
    num = iter("一二三四五六七八九十")
    sec = lambda t: f"## {next(num)}、{t}\n"  # noqa: E731  公开版少一节,序号顺延
    title = led.meta.get("title") or f"{country.name_zh or country.name_en} 水务项目机会分析"
    L.append(f"# {title}\n")
    L.append(
        f"> 台账截至 {led.as_of} · 报告生成 {today} · 共 {len(ds)} 项(核心范围 {len(core)} 项)"
        + ("" if public else f" · 投标人国籍假设 {nat} · 含使用方内部信息,仅供内部参考")
    )
    if led.meta.get("baseline"):
        L.append(f">\n> {led.meta['baseline']}")
    L.append("")

    # 本期要点
    L.append(sec("本期要点"))
    for note in [] if public else overlay.get("country_notes") or []:
        L.append(f"- {note}")
    if not public and overlay.get("country_notes"):
        L.append("")
    by_stage: dict[str, int] = {}
    for d in core:
        by_stage[d.item["stage"]] = by_stage.get(d.item["stage"], 0) + 1
    L.append("- 核心范围按阶段:" + ";".join(f"{STAGES[s]} {n} 项" for s, n in sorted(by_stage.items(), key=lambda x: STAGE_ORDER.index(x[0]))))
    if not public:
        by_el: dict[str, int] = {}
        for d in core:
            by_el[d.elig.status] = by_el.get(d.elig.status, 0) + 1
        L.append(f"- 投标资格(投标人国籍 {nat}):" + ";".join(f"{ELIG[s]} {by_el[s]} 项" for s in ELIG if s in by_el))
        soon = sorted((d for d in ds if d.timing.latest_start), key=lambda d: d.timing.latest_start)
        # 只提示核心范围内、资格不受限的项目;参考项与受限项在时间窗口表里仍可查
        urgent = [
            d for d in soon
            if d.timing.days_left is not None and d.timing.days_left <= 30
            and d.item.get("in_scope", True) and d.elig.status != "restricted"
        ]
        if urgent:
            L.append(
                "- 最晚启动日期已过或 30 天内到期:"
                + ";".join(f"{d.item['id']} {d.item['name_zh']}({_latest_text(d)},{_node_text(d)})" for d in urgent)
            )
    by_v: dict[str, int] = {}
    for d in ds:
        s = (d.item.get("verification") or {}).get("status", "unverified")
        by_v[s] = by_v.get(s, 0) + 1
    n_stale = sum(len(d.stale) for d in ds)
    L.append(
        "- 核实状态:" + ";".join(f"{VERIFY_STATUS[s].split('(')[0]} {by_v[s]} 项" for s in VERIFY_STATUS if s in by_v)
        + (f";过期未更新的计划节点 {n_stale} 处(见\"待核实与过期提醒\")" if n_stale else "")
    )
    L.append("")
    intel = [] if public else overlay.get("intel") or []
    if intel:
        L.append("**非公开信息(内部,未经公开来源证实)**\n")
        L += [f"- {x.get('date', '')} {x.get('source', '')}:{x.get('content')}" for x in intel]
        L.append("")

    # 国别背景(基本资料与宏观、人口、水资源指标,来自 ingest 快照;facts 为 core.country_facts 的国别概况)
    if facts is not None and (facts.values or facts.basics):
        from core.country_facts import background_md

        bg = background_md(facts, table, internal=not public)
    else:
        bg = _background(indicators)
    if bg:
        L.append(sec("国别背景(宏观与供排水指标)"))
        L += bg

    # 总表
    L.append(sec("项目机会台账(总表)"))
    head = ["编号", "项目", "领域", "阶段", "体量", "出资方", "下一节点"]
    head += [] if public else ["投标资格", "最晚启动"]
    head += ["核实"]
    body = []
    for d in ds:
        it = d.item
        r = [
            it["id"], it["name_zh"] + ("" if it.get("in_scope", True) else "(参考)"),
            "、".join(SEGMENTS[s] for s in it.get("segments") or []),
            STAGES[it["stage"]], usd(d.size_usd), "、".join(dict.fromkeys(_funder(f.get("funder"), rules) for f in it.get("financing") or [])),
            _node_text(d),
        ]
        r += [] if public else [ELIG[d.elig.status], _latest_text(d)]
        r += [VERIFY_STATUS[(it.get("verification") or {}).get("status", "unverified")].split("(")[0]]
        body.append(r)
    L.append(table(head, body))
    L.append("\n体量:优先取总投资,否则取已批准/已签署资金之和,按附录 B 汇率折美元;不含备忘录与仅见报道的金额。\n")

    # 时间窗口
    if not public:
        L.append(sec("时间窗口(按最晚启动日期排序)"))
        timed = sorted((d for d in ds if d.timing.latest_start), key=lambda d: d.timing.latest_start)
        L.append(
            table(
                ["最晚启动", "编号", "项目", "下一节点", "投标资格", "内部周期构成"],
                [[_latest_text(d), d.item["id"], d.item["name_zh"] + ("" if d.item.get("in_scope", True) else "(参考)"), _node_text(d),
                  ELIG[d.elig.status], d.timing.lead_basis] for d in timed],
            )
            if timed else "暂无带日期的投标人节点。"
        )
        untimed = [d for d in ds if not d.timing.latest_start and d.item.get("stage") not in ("completed", "cancelled")]
        if untimed:
            L.append("\n无明确节点、需持续跟踪:" + ";".join(f"{d.item['id']} {d.item['name_zh']}" for d in untimed) + "。")
        L.append(
            "\n最晚启动日期 = 下一节点日期(区间取起点) − 内部流程周期(见附录 B);"
            "节点为'推断'时日期本身不确定,应以官方公告为准并随时更新。\n"
        )

    # 活动与窗口(公开事实)
    events = sorted(led.meta.get("events") or [], key=lambda e: parse_pdate(e["date"]).start)
    if events:
        L.append(sec("活动与窗口"))
        L.append(
            table(
                ["日期", "活动", "地点", "说明", "来源"],
                [[f"{e['date']}" + (f" 至 {e['end']}" if e.get("end") else ""), e.get("name"), e.get("place"), e.get("note"), refs(e.get("src"))]
                 for e in events],
            )
            + "\n"
        )

    # 建议动作清单(内部)
    if not public:
        acts = [dict(a, _id="国别") for a in overlay.get("actions") or []]
        for iid, ov in ov_items.items():
            acts += [dict(a, _id=iid) for a in (ov or {}).get("next_actions") or []]
        if acts:
            L.append(sec("建议动作清单(内部)"))
            acts.sort(key=lambda a: (parse_pdate(a["due"]).start if a.get("due") else date.max, a["_id"]))
            L.append(table(["建议完成", "对象", "动作", "责任"], [[a.get("due"), a["_id"], a.get("action"), a.get("owner")] for a in acts]) + "\n")

    # 投标资格规则
    L.append(sec("投标资格规则(本国涉及的采购制度)"))
    used: dict[str, list[str]] = {}
    for d in ds:
        pr = d.item.get("procurement") or {}
        for rg in [pr.get("regime")] + list(pr.get("other_regimes") or []):
            if rg:
                used.setdefault(rg, []).append(d.item["id"])
    regs = rules.get("regimes") or {}
    rows_ = []
    for rg, ids in used.items():
        r = regs.get(rg) or {}
        rows_.append(
            [r.get("name_zh", "采购制度未定" if rg == "unknown" else rg), NATIONALITY.get(r.get("nationality"), "待核实"),
             {"none": "不要求", "bid": "投标前", "contract": "签约前", "unknown": "待核实"}.get(r.get("local_registration"), "待核实"),
             r.get("platform"), r.get("grade", "—"), "、".join(ids)]
        )
    L.append(table(["采购制度", "国籍规则", "当地注册", "平台/入口", "等级", "涉及项目"], rows_))
    L.append("")
    for rg in used:
        r = regs.get(rg)
        if r and r.get("summary"):
            srcs = "; ".join(f"[{s.get('title', s.get('url'))}]({s['url']})" if s.get("url") else s.get("title", "") for s in r.get("src") or [])
            L.append(f"- **{r.get('name_zh', rg)}**:{r['summary']}" + (f" 来源:{srcs}" if srcs else ""))
    L.append("")
    access = led.meta.get("access") or []
    if access:
        L.append("**本国准入门槛(与出资方无关)**\n")
        L += [f"- {a['text']}({refs(a.get('src'))})" for a in access]
        L.append("")

    # 项目卡片
    L.append(sec("项目卡片"))
    for d in ds:
        L.extend(_card(d, led, rules, ov_items.get(d.item["id"]), public))

    # 待核实与过期提醒
    L.append(sec("待核实与过期提醒"))
    any_ = False
    for d in ds:
        v = d.item.get("verification") or {}
        lines = [f"待核实:{x}" for x in v.get("pending") or []] + [f"来源冲突:{x}" for x in v.get("conflicts") or []] + d.stale
        if lines:
            any_ = True
            L.append(f"- **{d.item['id']} {d.item['name_zh']}**:" + ";".join(lines))
    if not any_:
        L.append("无。")
    L.append("")

    # 附录
    L.append("## 附录 A 来源登记\n")
    used = cited(led)
    L.append(f"只列本报告引用的 {len(used)} 条;完整登记 {len(led.sources)} 条(含成稿来源登记表)见 countries/{led.country}/sources.yaml。\n")
    L.append(
        table(
            ["编号", "名称", "机构", "日期", "等级", "获取方式", "链接"],
            [[k, led.sources[k].get("title"), led.sources[k].get("publisher"), led.sources[k].get("date"), led.sources[k].get("grade"),
              VIA.get(led.sources[k].get("via"), led.sources[k].get("via")),
              f"<{led.sources[k]['url']}>" if led.sources[k].get("url") else None]
             for k in sorted(used)],
        )
    )
    L.append("\n来源等级:" + ";".join(f"{k} {v}" for k, v in GRADES.items()) + "。")
    L.append("获取方式:" + ";".join(f"{v}" for v in VIA.values()) + "——'检索摘要'表示只读到检索结果摘要,未打开原文,关键数字须复核。\n")
    L.append("## 附录 B 口径说明\n")
    L.append("汇率(折美元):" + ";".join(f"1 {CUR_ZH.get(k, k)} = {v} 美元" for k, v in led.fx.items() if k != "USD") + "。\n")
    if not public and profile.lead_times:
        L.append("内部流程周期(天,来自使用方 profile;用于最晚启动日期):\n")
        L.append(
            table(
                ["节点类型", "构成"],
                [[MILESTONES.get(k, {"local_registration": "本地实体注册与建筑许可(投标前须注册的制度才计入)"}.get(k, k)),
                  "; ".join(f"{a} {b}" for a, b in v.items()) if isinstance(v, dict) else v]
                 for k, v in profile.lead_times.items()],
            )
        )
        L.append("")
    L.append("状态定义:阶段——" + ";".join(STAGES.values()) + "。核实——" + ";".join(VERIFY_STATUS.values()) + "。\n")
    return "\n".join(L)


# 国别背景表:(WDI 代码, IMF 代码);WDI 取最新实际值,IMF 给同一指标的后续预测
BG_ROWS = [
    ("NY.GDP.MKTP.CD", None), ("NY.GDP.PCAP.CD", None), ("NY.GNP.PCAP.CD", None),
    ("NY.GDP.MKTP.KD.ZG", "NGDP_RPCH"), ("FP.CPI.TOTL.ZG", "PCPIPCH"), (None, "GGXWDG_NGDP"), (None, "BCA_NGDPD"),
    ("DT.DOD.DECT.GN.ZS", None), ("BX.KLT.DINV.CD.WD", None), ("PA.NUS.FCRF", None),
    ("SP.POP.TOTL", None), ("SP.URB.TOTL.IN.ZS", None),
    ("SH.H2O.BASW.ZS", None), ("SH.H2O.SMDW.ZS", None), ("SH.STA.BASS.ZS", None), ("SH.STA.SMSS.ZS", None),
    ("ER.H2O.FWST.ZS", None),
]
SRC_ZH = {"WB-WDI": "世行 WDI", "IMF-WEO": "IMF WEO"}


def _fmt(v: float, unit: str) -> str:
    if unit == "USD":
        return money(v, "USD")
    if unit == "%":
        return f"{v:.1f}%"
    if unit == "人":
        return f"{_num(v / 1e4, 0)} 万人"
    if unit == "百万人":
        return f"{v:.2f} 百万人"
    if unit == "LCU/USD":
        return f"{v:.3f}"
    return _num(v, 2)


def _background(indicators: dict | None) -> list[str]:
    """国别背景表:每个指标给最新实际值(标年份),IMF 指标另给后三年预测。无快照时不出这一节。"""
    rows_ = (indicators or {}).get("rows") or []
    if not rows_:
        return []
    series: dict[tuple[str, str], list[dict]] = {}
    for r in rows_:
        series.setdefault((r["source"], r["indicator"]), []).append(r)
    body = []
    for wdi, imf in BG_ROWS:
        base = sorted(series.get(("WB-WDI", wdi)) or [], key=lambda r: r["year"]) if wdi else []
        fc = sorted(series.get(("IMF-WEO", imf)) or [], key=lambda r: r["year"]) if imf else []
        actual = [r for r in base if not r["is_forecast"]] or [r for r in fc if not r["is_forecast"]]
        if not actual:
            continue
        last = actual[-1]
        src = SRC_ZH.get(last["source"], last["source"])
        ahead = [r for r in fc if r["is_forecast"] and r["year"] > last["year"]][:3]
        trend = "、".join(f"{r['year']} {_fmt(r['value'], r['unit'])}" for r in ahead)
        if ahead and last["source"] != "IMF-WEO":
            src += " / IMF WEO"
        label = last["indicator_label"].replace("(IMF,含预测)", "").replace("(IMF)", "")
        body.append([label, f"{_fmt(last['value'], last['unit'])}({last['year']})", trend or "—", src])
    if not body:
        return []
    return [
        table(["指标", "最新值(年份)", "IMF 预测", "来源"], body),
        f"\n数据取自 {indicators.get('as_of', '最近一期')} 的 ingest 快照(世行 WDI 与 IMF《世界经济展望》数据接口);"
        "年份为数据所属年份。人均国民收入决定世行工程国内优惠是否适用;供排水覆盖率为联合国/世卫联合监测口径。\n",
    ]


def _funder(code: object, rules: dict) -> str:
    f = (rules.get("funders") or {}).get(str(code)) or {}
    return f.get("short_zh") or f.get("name_zh") or str(code)


def _card(d: Derived, led: Ledger, rules: dict, ov: dict | None, public: bool) -> list[str]:
    it = d.item
    L = [f"### {it['id']} {it['name_zh']}\n"]
    if it.get("name_en") or it.get("name_local"):
        L.append("*" + " / ".join(x for x in (it.get("name_en"), it.get("name_local")) if x) + "*\n")
    if it.get("brief"):
        L.append(f"> {it['brief']}\n")
    loc, own, v = it.get("location") or {}, it.get("owner") or {}, it.get("verification") or {}
    tc = it.get("total_cost") or {}
    kv = [
        ["阶段", STAGES[it["stage"]] + (f":{it['stage_note']}" if it.get("stage_note") else "")],
        ["地点", (loc.get("city_zh") or "") + (f"({loc['detail']})" if loc.get("detail") else "")],
        ["业主与执行", ";".join(f"{k}:{own[k]}" for k in own)],
        ["编号", ";".join(f"{k} {val}" for k, val in (it.get("ids") or {}).items())],
        ["体量", _size_text(tc, d.size_usd)],
        ["下一节点", _node_text(d)],
    ]
    if not public:
        kv += [
            ["投标资格", f"{ELIG[d.elig.status]}——{d.elig.basis}"],
            ["最晚启动", _latest_text(d) + (f";{d.timing.lead_basis}" if d.timing.latest_start else "")],
        ]
    kv.append(["核实", f"{VERIFY_STATUS.get(v.get('status'), '—')};{v.get('checked_at', '')};{VIA.get(v.get('method'), v.get('method') or '')}"])
    L.append(table(["项", "内容"], kv) + "\n")

    fin = it.get("financing") or []
    if fin:
        L.append("**资金**\n")
        L.append(
            table(
                ["出资方", "方式", "金额", "状态", "日期", "说明", "来源"],
                [[_funder(f.get("funder"), rules), INSTRUMENTS.get(f.get("instrument"), f.get("instrument")), money(f.get("amount"), f.get("currency")),
                  FUNDING_STATUS.get(f.get("status"), f.get("status")), f.get("date"), f.get("note"), refs(f.get("src"))] for f in fin],
            )
            + "\n"
        )
    sc = it.get("scope") or []
    if sc:
        L.append("**工程内容**\n")
        L.append(table(["分项", "数量", "说明", "来源"], [[s.get("item"), f"{_q(s.get('qty'))} {s.get('unit', '')}".strip(), s.get("note"), refs(s.get("src"))] for s in sc]) + "\n")
    pr = it.get("procurement") or {}
    if pr:
        L.append("**采购安排**\n")
        rg = (rules.get("regimes") or {}).get(pr.get("regime")) or {}
        rname = rg.get("name_zh") or ("采购制度未定" if pr.get("regime") in (None, "unknown") else pr.get("regime"))
        head = ";".join(x for x in (f"制度:{rname}", f"方式:{pr['method']}" if pr.get("method") else "", pr.get("note") or "") if x)
        if head:
            L.append(head + "\n")
        pk = pr.get("packages") or []
        if pk:
            L.append(
                table(
                    ["标段/合同", "类型", "状态", "发布", "截止", "结果/说明", "来源"],
                    [[p.get("name"), PKG_TYPES.get(p.get("type"), p.get("type")), PKG_STATUS.get(p.get("status"), p.get("status")),
                      p.get("published"), p.get("deadline"), "; ".join(x for x in (p.get("awardee"), p.get("note")) if x), refs(p.get("src"))] for p in pk],
                )
                + "\n"
            )
    ms = sorted(it.get("milestones") or [], key=lambda m: parse_pdate(m["date"]).start)
    if ms:
        L.append("**时间线**\n")
        L.append(table(["日期", "节点", "状态", "说明", "来源"], [[m["date"], MILESTONES.get(m["type"]), MS_STATUS.get(m["status"]), m.get("note"), refs(m.get("src"))] for m in ms]) + "\n")
    parties = it.get("parties") or []
    if parties:
        L.append("**已知参与方**\n")
        L += [f"- {p.get('role')}:{p.get('name')}" + (f"——{p['note']}" if p.get("note") else "") + (f"({refs(p.get('src'))})" if p.get("src") else "") for p in parties]
        L.append("")
    notes = it.get("notes") or []
    if notes:
        L.append("**要点与风险(公开信息)**\n")
        L += [f"- {n}" if isinstance(n, str) else f"- {n.get('text')}({refs(n.get('src'))})" for n in notes]
        L.append("")
    if v.get("pending"):
        L.append("**待核实**:" + ";".join(v["pending"]) + "\n")
    if ov and not public:
        L.append("**内部叠加(不公开)**\n")
        for k, lab in (("priority", "内部优先级"), ("owner", "我方负责人"), ("business_line", "业务线"), ("role", "我方角色"),
                       ("judgement", "判断"), ("note", "内部备注")):
            if ov.get(k) is not None:
                L.append(f"- {lab}:{ov[k]}")
        for a in ov.get("next_actions") or []:
            L.append(f"- 建议动作:{a.get('action')}" + (f"(建议完成 {a['due']})" if a.get("due") else "") + (f"〔{a['owner']}〕" if a.get("owner") else ""))
        for x in ov.get("intel") or []:
            L.append(f"- 非公开信息({x.get('date', '')},{x.get('source', '')}):{x.get('content')}")
        L.append("")
    return L


def _size_text(tc: dict, v: float | None) -> str:
    if tc.get("amount") is None:
        return usd(v)
    s = money(tc["amount"], tc.get("currency"))
    return s if str(tc.get("currency")).upper() == "USD" else f"{s}({usd(v)})"


def _alloc(need: list[float], floor: list[float], total: float) -> list[float]:
    """先给每列下限,余下宽度按"内容比例宽度超出下限的部分"分配;下限之和超过总宽时才整体缩放。
    (旧做法先取 max(比例, 下限) 再整体缩放,会把编号列压回下限以下,AZE-013 被折成两行。)"""
    if sum(floor) >= total:
        k = total / sum(floor)
        return [f * k for f in floor]
    target = [total * a / sum(need) for a in need]
    gap = [max(0.0, x - f) for x, f in zip(target, floor)]
    rest, g = total - sum(floor), sum(gap)
    if g <= 0:
        return [f + rest / len(floor) for f in floor]
    return [f + rest * x / g for f, x in zip(floor, gap)]


def _fit_columns(tb, total_cm: float, Cm, qn) -> None:  # noqa: N803
    """按内容宽度分配列宽:长文本列按比例,短列至少放得下最长的英文/编号词,避免 AZE-001 被竖排。"""
    rows = [[c.text for c in r.cells] for r in tb.rows]
    n = len(tb.columns)
    if not n:
        return
    need, floor = [], []
    for j in range(n):
        col = [r[j] for r in rows if j < len(r)]
        widest = max((_w(x) for x in col), default=4)
        need.append(max(6, min(40, widest)))
        tokens = [t for x in col for t in re.findall(r"[A-Za-z0-9./%$()\-]+", x)]
        fl = max([0.21 * _w(t) + 0.45 for t in tokens] + [0.8])  # 最长的英文/编号词不断开
        if widest <= 10:  # 短列(≤5 个汉字)整行放下;长文本列不设整行下限,把宽度留给正文
            fl = max(fl, 0.19 * widest + 0.45)
        floor.append(fl)
    cm = _alloc(need, floor, total_cm)
    tbl = tb._tbl
    tblw = tbl.tblPr.find(qn("w:tblW"))
    if tblw is not None:
        tblw.set(qn("w:type"), "dxa")
        tblw.set(qn("w:w"), str(int(total_cm * 567)))
    for j, col in enumerate(tb.columns):
        col.width = Cm(cm[j])
    for r in tb.rows:
        for j, c in enumerate(r.cells):
            if j < n:
                c.width = Cm(cm[j])


def _q(x: object) -> str:
    if isinstance(x, float):
        return _num(x, 2)
    if isinstance(x, int):
        return f"{x:,}"
    return "" if x is None else str(x)


def to_docx(md_path: Path, docx_path: Path) -> Path:
    """Markdown → Word:一级标题作文档标题;横向版面,统一中文字体、表格边框与字号(目录用 Word 导航窗格)。"""
    try:
        import pypandoc
    except ImportError as e:  # pragma: no cover
        raise RuntimeError("出 Word 需要 pypandoc_binary:pip install pypandoc_binary") from e
    args = ["--quiet", "--shift-heading-level-by=-1", "-M", "lang=zh-CN"]
    pypandoc.convert_file(str(md_path), "docx", format="markdown", outputfile=str(docx_path), extra_args=args)
    _polish_docx(docx_path)
    return docx_path


def _polish_docx(path: Path) -> None:
    try:
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.enum.section import WD_ORIENT
        from docx.shared import Cm, Pt
    except ImportError:  # 没装 python-docx 就保留 pandoc 默认样式
        return
    doc = Document(str(path))
    for sec in doc.sections:  # 总表列多,A4 横向(pandoc 的 sectPr 不带纸张尺寸,须显式给出)
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
        sec.left_margin = sec.right_margin = Cm(1.8)
        sec.top_margin = sec.bottom_margin = Cm(1.6)
    text_cm = 29.7 - 2 * 1.8
    for st in doc.styles:
        if st.type != 1 or st.font is None:  # 1 = 段落样式
            continue
        rpr = st.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        east = "SimHei" if st.name.startswith(("Heading", "Title")) else "SimSun"
        fonts.set(qn("w:eastAsia"), east)
    after_borders = {qn(f"w:{t}") for t in ("shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription")}
    for tb in doc.tables:
        tblpr = tb._tbl.tblPr
        for old in tblpr.findall(qn("w:tblBorders")):
            tblpr.remove(old)
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), "A6A6A6")
            borders.append(el)
        # OOXML 规定 tblPr 子元素顺序:tblBorders 须在 shd/tblLayout/tblCellMar/tblLook 之前,否则 Word 报"内容无法读取"
        nxt = next((c for c in tblpr if c.tag in after_borders), None)
        if nxt is None:
            tblpr.append(borders)
        else:
            nxt.addprevious(borders)
        _fit_columns(tb, text_cm, Cm, qn)
        for i, row in enumerate(tb.rows):
            for c in row.cells:
                if i == 0:
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:fill"), "E8EEF1")
                    c._tc.get_or_add_tcPr().append(shd)
                for p in c.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9)
                        if i == 0:
                            r.font.bold = True
                        _nobreak_hyphens(r, OxmlElement, qn)
    doc.save(str(path))


# 编号与日期里的连字符(AZE-013、2026-10-08、2026-Q4)不作断行点,换成 Word 的不断行连字符
_HYPHEN = re.compile(r"(?<=[A-Z0-9])-(?=[0-9QH])")


def _nobreak_hyphens(run, OxmlElement, qn) -> None:  # noqa: N803
    txt = run.text
    if not _HYPHEN.search(txt):
        return
    r = run._r
    for el in r.findall(qn("w:t")):
        r.remove(el)
    for k, part in enumerate(_HYPHEN.split(txt)):
        if k:
            r.append(OxmlElement("w:noBreakHyphen"))
        el = OxmlElement("w:t")
        el.set(qn("xml:space"), "preserve")
        el.text = part
        r.append(el)
