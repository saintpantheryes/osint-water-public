"""项目报告:按使用方的《项目分析报告》模板(private/templates/项目分析报告模板.docx)为台账里的一个项目出 Word,另存一份 Markdown。

事实从哪里来:
- 台账(countries/{ISO3}/ledger.yaml):名称、编号、地点、业主、资金、建设内容、采购、节点、参与方、备注、待核实事项,每条带来源编号;
- 本系统的线索:挂到这个项目的(编号吻合或模型确认,● 标记),作为来源登记,也写进进展;
- 国别数据(core/country_facts.py:世行 WDI、IMF《世界经济展望》、世行国别分类):第二部分后附"国别基本情况";
- 采购规则(reference/procurement_rules.yaml):采购规则一行与资格要求的合格性一条,带规则原文链接。
判断与建议:内部叠加(private/overlay,优先级、判断、切入点、下一步)只进内部版;没有的由模型据以上数据起草(注明"推断"),
都没有时写"待研判"。每条事实句末标来源编号,第七部分按首次引用排号登记;模板要求只引用公开信息,
所以台账里的内部成稿(发布方含"内部")与 D 级来源不引用,内部叠加里的情报不写进报告。
模板里没有、系统也没有的(业主联系人、资格条件细项等)写"待核实"或"以招标文件为准"。
"""
from __future__ import annotations

import copy
import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from core import geo, opportunity
from core.country_facts import CountryFacts, load_facts
from core.ledger import (
    ELIG, _refs, FUNDING_STATUS, INSTRUMENTS, MILESTONES, NATIONALITY, PKG_STATUS, PKG_TYPES, ROOT, SEGMENTS, STAGES, Derived, Ledger, derive,
    load_ledger, load_overlay, load_rules, parse_pdate, template_file,
)
from core.ledger_auto import history
from core.markets import load_markets
from core.profile import Profile
from core.watch.llm import Classifier
from core.watch.query import load_rows
from core.watch.store import Store

TEMPLATE = template_file("项目分析报告模板.docx")
CUR_ZH = {"EUR": "欧元", "USD": "美元", "AZN": "马纳特", "MAD": "迪拉姆", "SAR": "沙特里亚尔", "CNY": "人民币", "GBP": "英镑", "JPY": "日元",
          "KWD": "科威特第纳尔", "XOF": "西非法郎", "EGP": "埃及镑", "TRY": "土耳其里拉"}
GRADE_TEXT = {"A": "官方一手", "B": "权威二手", "C": "媒体与企业"}
STATES = ("已完成", "进行中", "待核实", "未开始")
SHADE = {"已完成": "E2EFDA", "进行中": "FFF2CC", "待核实": "F8D7DA", "未开始": "EDEDED",  # 绿、黄、红、灰
         "Completed": "E2EFDA", "In progress": "FFF2CC", "Unconfirmed": "F8D7DA", "Not started": "EDEDED"}
STATES_EN = dict(zip(STATES, ("Completed", "In progress", "Unconfirmed", "Not started")))
# 英文(外部)模板里与中文模板对应的文字:认表、认段落用(模板见 scripts/make_project_template_en.py)
LABELS = {
    "zh": {"country": "国别", "core": ("项目：", "进展："), "basics": ["项目", "内容", "来源"], "phases": ["阶段", "状态"], "pros": "有利因素",
           "quals": "项目要求", "actions": ["序号", "动作"], "sources": ["编号", "来源与网址"], "howto": "填写说明", "title": ("【项目名称】", "【国别】"),
           "title_en": "外文名称", "background": "背景", "entry": ("首选", "次选", "补充", "潜在竞争者"), "colon": "：", "grades": "等级说明",
           "note": "编制说明:", "leads": "挂到本项目的系统线索:", "lang": "中文"},
    "en": {"country": "Country", "core": ("Project:", "Progress:"), "basics": ["Item", "Details", "Source"], "phases": ["Stage", "Status"],
           "pros": "Strengths", "quals": "Requirement", "actions": ["No.", "Action"], "sources": ["No.", "Source and URL"], "howto": "Notes for authors",
           "title": ("【Project name】",), "title_en": "official name", "background": "Background",
           "entry": ("First choice", "Second choice", "Supply", "Potential competitors"), "colon": ": ", "grades": "Grades",
           "note": "Notes: ", "leads": "", "lang": "English"},
}
MEDIA = ("新闻", "Trend", "APA", "GDELT")  # 信源名里有这些的线索按"本地媒体"登记,其余按官方一手
MANUAL = "manual"  # 人工入库的线索(core.watch.manual.SOURCE):等级待核实
MUNICIPAL = {"water_supply", "transmission", "nrw", "wastewater", "sludge", "reuse", "stormwater"}
PHASES = ("前期研究", "融资批准", "咨询招标", "融资签约", "详细设计", "施工招标", "施工建设")
PHASE_OF = {"concept": "前期研究", "mou": "前期研究", "board_approval": "融资批准", "loan_signing": "融资签约", "effectiveness": "融资签约",
            "gpn": "施工招标", "tender_launch": "施工招标", "eoi_deadline": "施工招标", "pq_deadline": "施工招标", "bid_deadline": "施工招标",
            "rfp_deadline": "施工招标", "award": "施工招标", "contract_signing": "施工招标", "construction_start": "施工建设",
            "completion": "施工建设", "cod": "施工建设"}
OTHER_PHASE = (("短名单|咨询|监理|实施支持", "咨询招标"), ("详细设计|设计", "详细设计"), ("开工|施工|完工|投运|竣工", "施工建设"),
               ("招标|采购|公告|标书", "施工招标"), ("董事会|批准", "融资批准"), ("签署|协议|贷款", "融资签约"))
QUALS = (
    ("合格性与注册", None),
    ("财务稳健", "一般要求提交近 3–5 年经审计的财务报表,无破产、清算与重大欠税"),
    ("平均年营业额", "一般按近 3–5 年建筑业务平均年营业额设门槛,联合体各方按比例合计满足,牵头方另有最低比例"),
    ("融资能力", "一般要求证明可动用的流动资金或银行授信不低于合同数月的现金流需要"),
    ("一般经验", "一般要求近 5–10 年以总包、分包或联合体成员身份完成若干份工程合同"),
    ("管理经验", "一般要求有总承包或联合体牵头的管理经验"),
    ("类似工程经验", "一般要求近 5–10 年完成规模与性质类似的合同(如同等规模的水厂、管网),并列出关键工作量"),
    ("履约与诉讼", "一般要求近年无因承包商原因的合同终止,未决诉讼占净资产的比例有上限"),
    ("环境健康安全", None),
    ("关键人员", "一般要求项目经理、各专业负责人与环境健康安全负责人的学历与类似项目年限"),
    ("本地许可", None),
)
PROS = ("资金", "采购", "匹配度", "信息与渠道")
CONS = ("技术与施工", "合同与接口", "不确定性", "我方短板")
ENTRY = ("首选", "次选", "补充", "潜在竞争者")
_MARK = re.compile(r"⟦([^⟦⟧]+)⟧")


def money(v: float | None, cur: str | None) -> str:
    """57000000 EUR → 5,700 万欧元;410000000 → 4.1 亿欧元。"""
    if v is None:
        return "金额未公布"
    unit = CUR_ZH.get(str(cur or ""), str(cur or ""))
    if v >= 1e8:
        n = f"{v / 1e8:,.2f}".rstrip("0").rstrip(".")
        return f"{n} 亿{unit}"
    if v >= 1e4:
        n = f"{v / 1e4:,.1f}".rstrip("0").rstrip(".")
        return f"{n} 万{unit}"
    return f"{v:,.0f} {unit}"


def plain(md: str) -> str:
    """模型写的 Markdown → 纯文字(去粗体、斜体与列表记号)。"""
    s = re.sub(r"\*\*([^*]+)\*\*|__([^_]+)__", lambda m: m.group(1) or m.group(2), str(md or ""))
    s = re.sub(r"(?m)^\s*(?:[-*•]|\d+[.、)])\s+", "", s)
    return s.strip()


@dataclass
class Sources:
    """来源登记。先按键登记(台账来源编号、线索、国别数据、采购规则),文中用 ⟦键⟧ 标;成稿时按首次出现排号。"""

    rows: dict[str, dict] = field(default_factory=dict)  # 键 → {title, url, publisher, date, grade}
    blocked: set[str] = field(default_factory=set)  # 台账里不能引用的来源(内部成稿、D 级)

    def add(self, key: str, title: str, url: str = "", publisher: str = "", when: str = "", grade: str = "") -> str:
        if key not in self.rows:
            self.rows[key] = {"title": title, "url": url or "", "publisher": publisher or "—", "date": when or "—", "grade": grade or "—"}
        return key

    def cite(self, keys) -> str:
        """来源标记;这条事实只有不能引用的来源时,标"待核实"(不把内部材料当公开事实写)。"""
        keys = list(dict.fromkeys(keys or ()))
        ok = "".join(f"⟦{k}⟧" for k in keys if k in self.rows)
        return ok or ("(待核实:仅有内部来源)" if any(k in self.blocked for k in keys) else "")


@dataclass
class Project:
    iso3: str
    item: dict
    led: Ledger
    derived: Derived | None
    overlay: dict  # 内部叠加里这一项(判断、优先级、切入点、下一步;情报不用)
    leads: list[dict]  # 挂到这个项目的线索(●),新的在前
    facts: CountryFacts | None
    rules: dict
    bidder: dict
    country_zh: str = ""  # profile 里的国名(没有时按国别参考表)


def load_project(profile: Profile, store: Store, iso3: str, item_id: str, today: date) -> Project:
    c = next((x for x in profile.countries if x.iso3 == iso3 and x.ledger), None)
    if c is None:
        raise ValueError(f"{iso3} 没有台账")
    rules = load_rules(profile.rules) if profile.rules else load_rules()
    led = load_ledger(c.ledger, rules)
    it = next((i for i in led.items if i["id"] == item_id), None)
    if it is None:
        raise ValueError(f"台账里没有 {item_id}")
    d = next((x for x in derive(led, rules, profile.bidder, profile.lead_times, today) if x.item["id"] == item_id), None)
    ov = ((load_overlay(c.overlay) if c.overlay else {}).get("items") or {}).get(item_id) or {}
    # 所属线索按全系统同一口径(core.opportunity:人并入的、写进台账的、编号吻合的)
    leds = {x.iso3: (led if x.iso3 == iso3 else load_ledger(x.ledger, rules)) for x in profile.countries if x.ledger}
    book = opportunity.build(store.load_signals(), leds, store.events(), store.ledger_reviewed(), history(store), store.triage())
    mine = set(book.leads.get(item_id, []))
    leads = [r for r in load_rows(store, today) if r["no"] in mine and citable(r)]
    leads.sort(key=lambda r: r["published"] or r["first_seen"], reverse=True)
    facts = load_facts(store.root.parent, load_markets(profile.markets), [iso3]).get(iso3)
    return Project(iso3, it, led, d, ov if isinstance(ov, dict) else {}, leads, facts, rules, profile.bidder or {},
                   getattr(c, "name_zh", "") or geo.name_zh(iso3))


def citable(r: dict) -> bool:
    """能当来源的线索:没被标为忽略(无关或重复),且有网页地址——人工粘贴的文字、上传的文件可能是内部材料,不引用。
    台账出处补成的线索(core.ledger_leads)不另列:报告直接引台账的出处,免得同一份材料列两遍。"""
    return (r.get("status") != "ignore" and r.get("source") != "ledger"
            and str(r.get("url") or "").lower().startswith(("https://", "http://")))


# ---------- 事实 → 各部分文字(带 ⟦来源⟧) ----------

def _register(p: Project, src: Sources) -> None:
    for k, s in p.led.sources.items():
        if s.get("grade") == "D" or "内部" in str(s.get("publisher") or ""):
            src.blocked.add(k)
            continue
        src.add(k, str(s.get("title") or k), str(s.get("url") or ""), str(s.get("publisher") or ""), str(s.get("date") or ""),
                GRADE_TEXT.get(str(s.get("grade")), str(s.get("grade") or "")))
    for r in p.leads:
        grade = ("待核实(人工提交)" if r.get("source") == MANUAL else "本地媒体" if any(w in r["source_name"] for w in MEDIA)
                 else "官方一手")
        src.add(f"lead:{r['no']}", f"{r['title']}(系统线索 {r['no']})", r["url"], r["source_name"], r["published"] or r["first_seen"], grade)
    proc = p.item.get("procurement") or {}
    reg = (p.rules.get("regimes") or {}).get(proc.get("regime") or "") or {}
    funder = ((p.rules.get("funders") or {}).get(reg.get("funder") or "") or {}).get("name_zh") or reg.get("funder") or ""
    for i, s in enumerate(reg.get("src") or []):
        src.add(f"rule:{i}", str(s.get("title") or reg.get("name_zh") or ""), str(s.get("url") or ""), funder,
                str(p.rules.get("checked_at") or ""), GRADE_TEXT.get(str(reg.get("grade") or "A"), "官方一手"))
    f = p.facts
    if f is not None:
        if any(v.source == "WB-WDI" for v in f.values.values()):
            src.add("wdi", f"世界发展指标(WDI):{f.name_zh}", f"https://data.worldbank.org/country/{f.iso2}" if f.iso2 else "", "世界银行",
                    f.as_of, "官方一手")
        if any(v.source == "IMF-WEO" or v.forecasts for v in f.values.values()):
            src.add("imf", f"《世界经济展望》数据库:{f.name_zh}", f"https://www.imf.org/en/Countries/{f.iso3}", "国际货币基金组织", f.as_of,
                    "官方一手")
        if f.meta:
            src.add("wbclass", "世界银行国别分类(收入组与贷款类别)", f"https://data.worldbank.org/country/{f.iso2}" if f.iso2 else "",
                    "世界银行", f.as_of, "官方一手")


def _when(m: dict) -> str:
    raw = str(m.get("date") or "")
    return raw + {"planned": "(计划)", "expected": "(估算)"}.get(str(m.get("status")), "")


def _ms_text(m: dict) -> str:
    what = MILESTONES.get(m.get("type"), "") if m.get("type") != "other" else ""
    note = str(m.get("note") or "")
    return (f"{what}({note})" if what and note else what or note) or "—"


def _first(text: object, n: int = 140) -> str:
    """长句只取到第一个句号(不在句中截断);第一句也超长时整句保留。"""
    s = str(text or "")
    if len(s) <= n:
        return s
    head = s.split("。")[0]
    return head + "。" if len(head) < len(s) else s


def _fin_line(fz: dict) -> str:
    inst = INSTRUMENTS.get(fz.get("instrument"), fz.get("instrument") or "")
    st = FUNDING_STATUS.get(fz.get("status"), fz.get("status") or "")
    when = f",{fz['date']}" if fz.get("date") else ""
    amount = money(fz.get("amount"), fz.get("currency")) if fz.get("amount") is not None else "金额未公布"
    note = f";{fz['note']}" if fz.get("note") else ""
    return f"{fz.get('funder') or '—'} {inst} {amount}({st}{when}){note}"


def _field(segs: list[str]) -> str:
    kinds = [k for k, hit in (("市政水务", MUNICIPAL & set(segs)), ("水利", {"reservoir", "irrigation", "flood"} & set(segs)),
                              ("海水淡化", "desalination" in segs), ("水环境", "environment" in segs)) if hit]
    return f"{'、'.join(kinds) or '水务'}({'、'.join(SEGMENTS.get(s, s) for s in segs)})" if segs else "水务"


def _notes(it: dict, pattern: str) -> list[dict]:
    """台账备注里合 pattern 的;备注可以只写一句话(没有来源),按 {text} 对待。"""
    notes = [n if isinstance(n, dict) else {"text": str(n)} for n in it.get("notes") or []]
    return [n for n in notes if re.search(pattern, str(n.get("text") or ""))]


STEPS = {"tender_launch", "eoi_deadline", "pq_deadline", "bid_deadline", "rfp_deadline", "award", "contract_signing"}
CONSULT = "咨询|监理|实施支持|短名单|顾问|consult"


def _phase_of(m: dict, consult_dates: set[str] = frozenset()) -> str:
    """节点 → 进展表的阶段。采购步骤(意向书截止、授标等)是咨询的——备注里写着咨询、监理,或日期就是咨询包的日期——归"咨询招标"。"""
    note = str(m.get("note") or "")
    if m.get("type") in STEPS and (re.search(CONSULT, note, re.I) or str(m.get("date")) in consult_dates):
        return "咨询招标"
    if m.get("type") != "other":
        return PHASE_OF.get(m.get("type"), "前期研究")
    return next((ph for pat, ph in OTHER_PHASE if re.search(pat, note)), "前期研究")


def _state(entries: list[tuple[str, str]], today: date) -> str:
    """一个阶段的状态:entries = [(节点状态 actual/planned/expected 或采购包状态, 日期)]。"""
    states = []
    for st, when in entries:
        if st in ("actual", "awarded"):
            states.append("已完成")
        elif st in ("open", "closed"):
            states.append("进行中")
        elif st == "cancelled":
            states.append("待核实")
        else:
            try:
                past = parse_pdate(when).end < today if when else False
            except (ValueError, TypeError):
                past = False
            states.append("待核实" if past and st == "planned" else "未开始")
    if not states:
        return "待核实"
    if all(s == "已完成" for s in states):
        return "已完成"
    if "进行中" in states or ("已完成" in states and "未开始" in states):
        return "进行中"
    return "待核实" if "待核实" in states else "未开始"


def phases(p: Project, src: Sources, today: date) -> list[tuple[str, str, str, str, str]]:
    """第三部分进展表:[(阶段, 状态, 时间, 事项, 来源)]。台账节点按类型归到阶段;咨询包归"咨询招标",工程与设计施工包归"施工招标";
    挂到项目的线索里最近 3 条采购公告也写进来(咨询类进"咨询招标",其余进"施工招标")。"""
    it = p.item
    rows: dict[str, list[tuple[str, str, str, list]]] = {ph: [] for ph in PHASES}  # 阶段 → [(状态, 时间, 事项, 来源)]
    consult = {str(pk.get(k)) for pk in (it.get("procurement") or {}).get("packages") or [] if pk.get("type") == "consultancy"
               for k in ("published", "deadline") if pk.get(k)}
    for m in it.get("milestones") or []:
        rows[_phase_of(m, consult)].append((str(m.get("status") or ""), _when(m), _ms_text(m), m.get("src") or []))
    for pk in (it.get("procurement") or {}).get("packages") or []:
        ph = "咨询招标" if pk.get("type") == "consultancy" else "施工招标"
        span = "–".join(str(x) for x in (pk.get("published"), pk.get("deadline")) if x)
        what = f"{pk.get('name')}({PKG_TYPES.get(pk.get('type'), '')},{PKG_STATUS.get(pk.get('status'), '')})" + (f";{pk['note']}" if pk.get("note") else "")
        rows[ph].append((str(pk.get("status") or "planned"), span or "—", what, pk.get("src") or []))
    for r in [x for x in p.leads if x.get("procurement")][:3]:  # 最近 3 条采购公告:咨询类归"咨询招标";授标算已完成,其余算进行中
        consult = re.search(r"consult|advis|supervis|咨询|监理|顾问", f"{r['title']} {r['summary']}", re.I)
        rows["咨询招标" if consult else "施工招标"].append(
            ("actual" if r.get("category") == "award" else "open", r["published"] or r["first_seen"],
             f"系统采集:{r['category_zh']}·{r['title'][:80]}", [f"lead:{r['no']}"]))
    out = []
    for ph in PHASES:
        es = rows[ph]
        if not es:
            out.append((ph, "待核实", "—", "未检索到公开信息", "—"))
            continue
        state = _state([(s, w.split("(")[0]) for s, w, _, _ in es], today)
        out.append((ph, state, ";".join(dict.fromkeys(w for _, w, _, _ in es if w and w != "—")) or "—",
                    ";".join(x for _, _, x, _ in es), src.cite([k for *_, ks in es for k in ks]) or "—"))
    return out


def basics(p: Project, src: Sources) -> list[tuple[str, list[str], str]]:
    """第二部分:[(项目, [内容各行], 来源)],行序与模板相同。"""
    it, cite = p.item, src.cite
    owner = it.get("owner") or {}
    fin = it.get("financing") or []
    proc = it.get("procurement") or {}
    reg = (p.rules.get("regimes") or {}).get(proc.get("regime") or "") or {}
    ids = ";".join(f"{k} {v}" for k, v in (it.get("ids") or {}).items())
    loc = it.get("location") or {}
    people = _notes(it, "人口|受益|服务范围")
    parties = it.get("parties") or []
    design = [x for x in parties if re.search("咨询|设计|监理", str(x.get("role") or ""))]
    cpk = [pk for pk in proc.get("packages") or [] if pk.get("type") == "consultancy"]
    period = [m for m in it.get("milestones") or [] if m.get("type") in ("construction_start", "completion", "cod")] + _notes(it, "工期")
    env = _notes(it, "环境|环评|环社|征地|移民|遗产|文物")
    ops = _notes(it, "运营|运维")
    total = it.get("total_cost") or {}
    budget_only = fin and all(f.get("instrument") == "budget" for f in fin)
    execs = [x for x in parties if "执行" in str(x.get("role") or "")]
    return [
        ("项目名称与编号", [it.get("name_zh") or "—", " / ".join(x for x in (it.get("name_en"), it.get("name_local")) if x) or "外文名称待核实",
                     ids or "出资方项目号或招标编号待核实"], "—"),
        ("项目地点", ["".join(x for x in (loc.get("city_zh"), f"({loc['detail']})" if loc.get("detail") else "") if x) or "待核实",
                  *[str(n["text"]) + cite(n.get("src")) for n in people[:2]]], cite(k for n in people[:2] for k in n.get("src") or []) or "—"),
        ("业主与执行机构", [";".join(f"{k}:{v}" for k, v in owner.items()) or "待核实",
                     *[f"{x.get('role')}:{x.get('name')}" + cite(x.get("src")) for x in execs[:2]]],
         cite(k for x in execs for k in x.get("src") or []) or "—"),
        ("借款人", [str(owner.get("借款人") or ("国家预算" if budget_only else "待核实"))], "—"),
        ("总投资与资金来源", [*([f"总投资约 {money(total.get('amount'), total.get('currency'))}" + cite(total.get("src"))] if total else []),
                      *[_fin_line(f) + cite(f.get("src")) for f in fin]] or ["待核实"],
         cite([*(total.get("src") or []), *[k for f in fin for k in f.get("src") or []]]) or "—"),
        ("建设内容", [f"{s.get('item')}" + (f"({s['qty']:,} {s.get('unit') or ''})".replace(" )", ")") if s.get("qty") else "")
                  + (f";{s['note']}" if s.get("note") else "") + cite(s.get("src")) for s in it.get("scope") or []] or ["待核实"],
         cite(k for s in it.get("scope") or [] for k in s.get("src") or []) or "—"),
        ("采购规则", [f"{reg.get('name_zh') or proc.get('regime') or '待核实'}:{NATIONALITY.get(reg.get('nationality'), '待核实')}"
                  + cite(["rule:0"]), *[x for x in (proc.get("method"), proc.get("note"), reg.get("platform")) if x]],
         cite(["rule:0"]) or "—"),
        ("资格要求", ["资格条件以招标文件为准,尚未公布的见第四部分(二)框架"
                  + (f";投标资格判断:{ELIG.get(p.derived.elig.status, '')}({p.derived.elig.basis})" if p.derived else "")], "—"),
        ("设计与监理", [*[f"{x.get('role')}:{x.get('name')}" + (f"({x['note']})" if x.get("note") else "") + cite(x.get("src")) for x in design],
                   *[f"{pk.get('name')}:{PKG_STATUS.get(pk.get('status'), '')}" + (f";{pk['note']}" if pk.get("note") else "") + cite(pk.get("src"))
                     for pk in cpk]] or ["待核实"],
         cite([*[k for x in design for k in x.get("src") or []], *[k for pk in cpk for k in pk.get("src") or []]]) or "—"),
        ("工期", [(f"{MILESTONES.get(m.get('type'), '')} {_when(m)}" + (f";{m['note']}" if m.get("note") else "") if "type" in m
                 else str(m.get("text"))) + cite(m.get("src")) for m in period] or ["待核实"],
         cite(k for m in period for k in m.get("src") or []) or "—"),
        ("环境与社会", [str(n["text"]) + cite(n.get("src")) for n in env] or ["待核实"], cite(k for n in env for k in n.get("src") or []) or "—"),
        ("建成后运营", [str(n["text"]) + cite(n.get("src")) for n in ops]
         or [f"推断:由{(execs[0].get('name') if execs else None) or next(iter(owner.values()), '业主')}或其下属运营单位运营(推断)"],
         cite(k for n in ops for k in n.get("src") or []) or "—"),
        ("业主联系人", ["待核实(公开渠道未检索到)"], "—"),
    ]


def country_rows(p: Project, src: Sources) -> list[tuple[str, str, str]]:
    """国别基本情况(附在第二部分后):[(指标, 数值, 来源)]。"""
    f = p.facts
    if f is None:
        return []
    out = []
    for k in ("pop", "gdp", "gdp_pc", "growth", "inflation", "gov_debt", "urban", "water_basic", "san_basic"):
        v = f.get(k)
        if v is None:
            continue
        text = v.text(f.currency_code)
        if k == "growth" and v.forecasts:
            text += f";IMF 预测 {v.forecast_text(2)}"
        out.append((v.metric.label, text, src.cite(["wdi" if v.source == "WB-WDI" else "imf", *(["imf"] if v.forecasts else [])])))
    if f.meta:
        out.append(("世行收入组 / 贷款类别", f"{f.income()} / {f.lending()}", src.cite(["wbclass"])))
    return out


def fallback(p: Project, src: Sources, today: date) -> dict:
    """没有模型(或模型没写出)时各判断部分的写法:只用台账、线索与内部叠加里已有的,缺的写"待研判"。"""
    it, cite, ov = p.item, src.cite, p.overlay
    fin = it.get("financing") or []
    total = it.get("total_cost") or {}
    scope = "、".join(f"{s['item']}" + (f" {s['qty']:,} {s.get('unit') or ''}".rstrip() if s.get("qty") else "") for s in (it.get("scope") or [])[:3])
    head = f"{it.get('name_zh')}" + (f":{scope}等" if scope else "")
    money_part = (f";总投资约 {money(total.get('amount'), total.get('currency'))}" if total else "")
    fins = ";".join(_fin_line(f).split(";")[0] for f in fin[:3])
    ms = sorted([m for m in it.get("milestones") or [] if m.get("status") == "actual"], key=lambda m: str(m.get("date")))
    nxt = p.derived.timing.node if p.derived else None
    stage = STAGES.get(it.get("stage"), it.get("stage") or "")
    lead = p.leads[0] if p.leads else None
    lead_cite = cite([f"lead:{lead['no']}"]) if lead else ""
    progress = (f"当前处于{stage}" + (f"({it['stage_note']})" if it.get("stage_note") else "") + "。"
                + (f"最近的节点:{ms[-1].get('date')} {_ms_text(ms[-1])}{cite(ms[-1].get('src'))};" if ms else "")
                + (f"下一个节点:{MILESTONES.get(nxt.type, nxt.type)} {nxt.when.raw}{'(估算)' if nxt.status == 'expected' else ''}。" if nxt else "")
                + (f"系统最近采集到相关公告:{lead['title'][:60]}({lead['published'] or lead['first_seen']}){lead_cite}。" if lead else ""))
    pending = (it.get("verification") or {}).get("pending") or []
    conflicts = (it.get("verification") or {}).get("conflicts") or []
    proc = it.get("procurement") or {}
    reg = (p.rules.get("regimes") or {}).get(proc.get("regime") or "") or {}
    access = [a for a in (getattr(p.led, "meta", {}) or {}).get("access") or [] if isinstance(a, dict)]
    compete = [a for a in access if re.search("竞争格局|对手|竞争者", str(a.get("text")))]
    gaps = [a for a in access if re.search("注册|许可|资质|劳务|用工", str(a.get("text")))]
    works = [pk for pk in proc.get("packages") or [] if pk.get("type") in ("works", "design_build", "ppp")]
    goods = [pk for pk in proc.get("packages") or [] if pk.get("type") == "goods"]
    firm = [f for f in fin if f.get("status") in ("approved", "signed", "effective", "disbursing")]
    judge = ov.get("judgement") or ov.get("judgment")
    acts = ov.get("next_actions") or []
    return {
        "project": f"{head}{money_part};资金:{fins or '待核实'}。{cite([*(total.get('src') or []), *[k for f in fin for k in f.get('src') or []]])}",
        "progress": progress,
        "judgement": (f"{judge}(内部判断,优先级 {ov.get('priority') or '未定'})" if judge else "待研判(内部叠加里还没有这一项的判断)"),
        "advice": (f"建议列为 {ov.get('priority')} 类跟踪项目" if ov.get("priority") else "建议列为跟踪项目")
                  + (f",切入方式:{ov['role']}" if ov.get("role") else "") + "。"
                  + ("下一步:" + ";".join(f"{a.get('action')}(完成时间 {a.get('due') or '待定'})" for a in acts[:3]) + "。" if acts else ""),
        "background": str(it.get("brief") or "待核实") + cite([*[k for f in fin for k in f.get("src") or []], *(total.get("src") or [])][:3]),
        "资金": (f"已落实:{';'.join(_fin_line(f).split(';')[0] for f in firm)}{cite(k for f in firm for k in f.get('src') or [])}"
               if firm else "资金尚未落实(推断)") + ";支付币种与付款保障待核实",
        "采购": f"{reg.get('name_zh') or '采购规则待核实'},{NATIONALITY.get(reg.get('nationality'), '国籍限制待核实')}{cite(['rule:0'])}"
               + (f";{proc['note']}" if proc.get("note") else ""),
        "匹配度": f"建设内容为{_field(it.get('segments') or [])},与我方水务工程总承包主业吻合(推断)",
        "信息与渠道": f"本项目有公开来源 {len({k for k in _refs(it) if k in src.rows})} 条、"
                   f"系统线索 {len(p.leads)} 条;核实状态:{(it.get('verification') or {}).get('status') or '待核实'}"
                   f"(核对于 {(it.get('verification') or {}).get('checked_at') or '—'});对接渠道待补充",
        "技术与施工": ";".join(str(n["text"]) + cite(n.get("src")) for n in _notes(it, "风险|续建|质量|难点|地质|技术")[:2]) or "待研判",
        "合同与接口": ("工程标段:" + ";".join(f"{pk.get('name')}({PKG_STATUS.get(pk.get('status'), '')})" for pk in works)
                   + cite(k for pk in works for k in pk.get("src") or []) if works else "标段划分与合同条件未公布(待核实)"),
        "不确定性": ("待核实事项:" + "、".join(str(x) for x in pending) if pending else "")
                + (";来源冲突:" + "、".join(str(x) for x in conflicts) if conflicts else "") or "待研判",
        "我方短板": ";".join(_first(a["text"]) + cite(a.get("src")) for a in gaps[:2]) or "资质、业绩与本地化方面的不足待评估",
        "首选": (f"{ov.get('role')}({ov.get('business_line') or ''})" if ov.get("role") else
               (f"{works[0].get('name')}(推断)" if works else "待研判")),
        "次选": ("以联合体成员或分包身份参与" + ("、".join(str(pk.get("name")) for pk in works[1:3]) if len(works) > 1 else "") + "(推断)")
               if works else "待研判",
        "补充": ("设备材料供应:" + "、".join(str(pk.get("name")) for pk in goods) if goods
               else "设备材料供应(水泵、阀门、管材、仪表与 SCADA 等)待结合标段研判(推断)"),
        "潜在竞争者": ";".join(_first(a["text"], 200) + cite(a.get("src")) for a in compete[:1]) or "待研判",
        "leader": [f"是否同意将本项目列为{ov.get('priority') or ''}{'类' if ov.get('priority') else ''}跟踪项目",
                   "前期费用预算(购买招标文件、现场考察、联合体洽谈等):待编制",
                   f"授权与业主({next(iter((it.get('owner') or {}).values()), '业主')})及出资方驻地机构接触",
                   "协调当地注册与许可、类似业绩与关键人员等资源"],
        "partner": [*(f"核实:{x}" for x in pending[:3]), *([] if pending else ["核实下一节点的时间与标段划分"]),
                    f"拜访或对接:{';'.join(f'{k}:{v}' for k, v in (it.get('owner') or {}).items()) or '业主与执行机构'}",
                    "收集现场条件、当地材料与人工价格、潜在合作伙伴信息"],
    }


FIELDS = {
    "project": "项目:一至两句,写项目内容、总投资与资金来源;句末标来源编号",
    "progress": "进展:当前所处阶段、最近的关键节点和下一个节点;句末标来源编号",
    "judgement": "判断:是否值得跟踪,最主要的有利因素和难点;编制方判断注明\"推断\"",
    "advice": "建议:是否列为跟踪项目及下一步动作;请领导批准和请前方跟进的核心事项",
    "background": "背景:项目由来与此前阶段的完成情况,二至四句;句末标来源编号",
    **{k: f"有利因素·{k}:一两句" for k in PROS},
    **{k: f"不利因素与风险·{k}:一两句" for k in CONS},
    "首选": "我方切入点·首选:最适合我方的标段或合同及理由,一两句",
    "次选": "我方切入点·次选:备选切入方式,一句",
    "补充": "我方切入点·补充:设备材料供应等机会,一句",
    "潜在竞争者": "潜在竞争者:已知或可能的竞争对手及其优势;有来源的标来源编号",
    "leader": "请领导决策与支持:4 条,每行一条,不要编号(依次:是否列为跟踪项目、前期费用预算、对外接触授权、资源协调)",
    "partner": "请前方伙伴跟进:3 条,每行一条,不要编号(依次:需要核实的关键信息、需要拜访或对接的机构、需要收集的信息)",
}
SYSTEM = """你是投标人国籍为 {bidder} 的水务工程承包与投资企业的项目分析员,按公司《项目分析报告》模板为管理层写一个项目的分析。
只根据用户给出的数据写:不编造数字、机构、日期与竞争对手;引用事实时在句末用方括号写数据里给出的来源编号,如 [3];
数据里的"内部判断"可以用来写判断与建议,但不能当来源引用;编制方自己的判断注明"推断"或"估算";数据不足时写"待核实"。
用简洁的中文书面语,不写客套话。"""


def dataset(p: Project, src: Sources, facts: dict) -> tuple[str, dict[str, int]]:
    """给模型的数据 → (文字, 来源键 → 临时编号)。只给这些事实用到的来源与挂到项目的线索编号,模型只能引用它们。"""
    it = p.item
    lines = [f"项目:{it.get('id')} {it.get('name_zh')}({it.get('name_en') or ''});国别 {p.country_zh or geo.name_zh(p.iso3)};"
             f"阶段 {STAGES.get(it.get('stage'), '')}"]
    lines += [f"{label}:" + ";".join(rows) for label, rows, _ in facts["basics"]]
    lines += ["进展:" + ";".join(f"{a} {b} {c} {d} {e}" for a, b, c, d, e in facts["phases"])]
    lines += ["国别:" + ";".join(f"{a} {b}{c}" for a, b, c in facts["country"])]
    for a in (getattr(p.led, "meta", {}) or {}).get("access") or []:
        if isinstance(a, dict) and a.get("text"):
            lines.append("国别准入与竞争:" + str(a["text"]) + src.cite(a.get("src")))
    ov = p.overlay
    if ov:
        lines.append("内部判断(不能当来源):" + ";".join(f"{k} {v}" for k, v in ov.items() if k in ("priority", "business_line", "role", "judgement")))
        for a in ov.get("next_actions") or []:
            lines.append(f"内部计划动作:{a.get('action')}(完成时间 {a.get('due') or '待定'})")
    for r in p.leads[:12]:
        lines.append(f"系统线索:{r['published'] or r['first_seen']} {r['title'][:100]} | {r['category_zh']} | {r['summary'][:160]}"
                     + src.cite([f"lead:{r['no']}"]))
    text = "\n".join(lines)
    num: dict[str, int] = {}
    for k in _MARK.findall(text):
        num.setdefault(k, len(num) + 1)
    text = _MARK.sub(lambda m: f"[{num[m.group(1)]}]", text)
    return text + "\n来源编号:" + ";".join(f"[{n}] {src.rows[k]['publisher']}《{src.rows[k]['title'][:40]}》" for k, n in num.items()), num


def compose(p: Project, today: date, clf: Classifier | None = None, org: str = "", to: str = "公司领导") -> dict:
    """→ 报告内容(各部分文字里的来源编号已按首次出现排好)。"""
    src = Sources()
    _register(p, src)
    it = p.item
    facts = {"basics": basics(p, src), "phases": phases(p, src, today), "country": country_rows(p, src)}
    fb = fallback(p, src, today)
    note = "本报告未使用模型,判断部分取自内部叠加,没有的写\"待研判\"。"
    text = dict(fb)
    if clf is not None:
        data, num = dataset(p, src, facts)
        back = {str(i): k for k, i in num.items()}
        system = SYSTEM.format(bidder=(p.bidder or {}).get("nationality") or "未设")
        out = clf.write(system, data, FIELDS)
        if out and not out.get("_error"):
            bad: list[str] = []

            def mark(m: re.Match) -> str:
                nums = [x for x in re.split(r"[,,、\s]+", m.group(1)) if x]
                bad.extend(x for x in nums if x not in back)
                return "".join(f"⟦{back[x]}⟧" for x in nums if x in back)

            for k, v in out.items():
                if not v.strip():
                    continue
                v = re.sub(r"\[(\d+(?:\s*[,,、]\s*\d+)*)\]", mark, v)
                text[k] = [plain(x) for x in v.splitlines() if plain(x)] if k in ("leader", "partner") else plain(v)
            note = (f"判断、分析与建议部分由模型({clf.label})据台账、线索与国别数据起草,编制方判断已注明\"推断\";请核对后使用。"
                    + (f"其中 {len(bad)} 处来源编号在数据里没有,已删去。" if bad else ""))
        else:
            note = f"模型没有写出({(out or {}).get('_error') or clf.fatal or '未返回'}),判断部分取自内部叠加,没有的写\"待研判\"。"
    ov = p.overlay
    acts = [(a.get("action") or "", "【待定】", str(a.get("due") or "待定"),
             "待核实" if a.get("due") and str(a.get("due")) < today.isoformat() else "未开始") for a in ov.get("next_actions") or []]
    if not acts:
        nxt = p.derived.timing.node if p.derived else None
        acts = [(f"跟踪下一节点:{MILESTONES.get(nxt.type, nxt.type)} {nxt.when.raw}" if nxt else "跟踪项目下一节点", "【待定】",
                 str(p.derived.timing.latest_start or "待定") if p.derived else "待定", "未开始"),
                ("核实待核实事项,补充台账来源", "【待定】", "待定", "未开始")]
    content = {
        "id": it["id"], "iso3": p.iso3, "country": p.country_zh or geo.name_zh(p.iso3), "name": it.get("name_zh") or it["id"],
        "name_en": " / ".join(x for x in (it.get("name_en"), it.get("name_local")) if x) or "外文名称待核实",
        "field": _field(it.get("segments") or []),
        "owner": next((v for k, v in (it.get("owner") or {}).items() if re.search("业主|执行|受益", k)), None)
                 or next(iter((it.get("owner") or {}).values()), "待核实"),
        "funding": ";".join(_fin_line(f).split(";")[0] for f in it.get("financing") or []) or "待核实",
        "doc_no": f"{p.iso3}-XM-{today.year}-{it['id'].split('-')[-1]}", "as_of": today.isoformat(), "org": org or "【部门】", "to": to,
        "core": [("项目", text["project"]), ("进展", text["progress"]), ("判断", text["judgement"]), ("建议", text["advice"])],
        "background": text["background"],
        "basics": facts["basics"], "country_rows": facts["country"], "phases": facts["phases"],
        "pros": [(k, text[k]) for k in PROS], "cons": [(k, text[k]) for k in CONS],
        "quals": _quals(p, src), "entry": [(k, text[k]) for k in ENTRY], "actions": acts,
        "leader": list(text["leader"])[:6], "partner": list(text["partner"])[:6],
        "leads": [r["no"] for r in p.leads], "note": note,
    }
    return _number(content, src)


def _quals(p: Project, src: Sources) -> list[tuple[str, str]]:
    proc = p.item.get("procurement") or {}
    reg = (p.rules.get("regimes") or {}).get(proc.get("regime") or "") or {}
    access = [a for a in (getattr(p.led, "meta", {}) or {}).get("access") or [] if isinstance(a, dict)]
    lic = [a for a in access if re.search("许可|注册", str(a.get("text")))]
    env = _notes(p.item, "环境|环评|类别")
    out = []
    for k, generic in QUALS:
        if k == "合格性与注册":
            t = ";".join(x for x in (NATIONALITY.get(reg.get("nationality"), ""), str(reg.get("summary") or ""), str(reg.get("platform") or "")) if x)
            out.append((k, (t or "待核实") + src.cite(["rule:0"])))
        elif k == "环境健康安全":
            out.append((k, "一般要求环境健康安全管理体系认证与类似项目的环境社会管理经验(框架)"
                        + (";" + ";".join(str(n["text"]) + src.cite(n.get("src")) for n in env[:1]) if env else "")))
        elif k == "本地许可":
            out.append((k, ";".join(_first(a["text"], 200) + src.cite(a.get("src")) for a in lic[:1]) or "当地注册、执照与许可要求待核实"))
        else:
            out.append((k, f"{generic}(框架,以招标文件为准)"))
    return out


def _number(content: dict, src: Sources) -> dict:
    """按首次出现给来源排号:⟦键⟧ → [n];第七部分只列用到的。"""
    order: dict[str, int] = {}

    def walk(x):
        if isinstance(x, str):
            for k in _MARK.findall(x):
                order.setdefault(k, len(order) + 1)
        elif isinstance(x, (list, tuple)):
            for y in x:
                walk(y)

    keys = ("core", "basics", "country_rows", "background", "phases", "pros", "cons", "quals", "entry", "actions", "leader", "partner")
    for k in keys:
        walk(content[k])

    def fix(x):
        if isinstance(x, str):
            return _MARK.sub(lambda m: f"[{order[m.group(1)]}]", x)
        if isinstance(x, tuple):
            return tuple(fix(y) for y in x)
        if isinstance(x, list):
            return [fix(y) for y in x]
        return x

    for k in keys:
        content[k] = fix(content[k])
    content["sources"] = [(n, src.rows[k]) for k, n in sorted(order.items(), key=lambda kv: kv[1])]
    return content


# ---------- Markdown ----------

def _md_cell(s: str) -> str:
    return str(s).replace("|", "/").replace("\n", "<br>")


def to_markdown(c: dict) -> str:
    L = [f"# {c['country']}{c['name']}", "", f"**{c['name_en']}** · 项目分析报告", "",
         "| 国别 | 项目领域 | 业主机构 | 资金来源 | 文件编号 | 信息截至 | 编制单位 | 报送对象 |", "|" + "---|" * 8,
         "| " + " | ".join(_md_cell(c[k]) for k in ("country", "field", "owner", "funding", "doc_no", "as_of", "org", "to")) + " |", "",
         "## 一、核心结论", "", *[f"- **{k}:**{v}" for k, v in c["core"]], "",
         "## 二、项目基本情况", "", "| 项目 | 内容 | 来源 |", "|---|---|---|",
         *[f"| {k} | {'<br>'.join(_md_cell(x) for x in v)} | {s} |" for k, v, s in c["basics"]], ""]
    if c["country_rows"]:
        L += ["**国别基本情况**(括号内为数据年份)", "", "| 指标 | 数值 | 来源 |", "|---|---|---|",
              *[f"| {a} | {_md_cell(b)} | {s} |" for a, b, s in c["country_rows"]], ""]
    L += ["## 三、项目进展", "", f"**背景:**{c['background']}", "", "| 阶段 | 状态 | 时间 | 事项 | 来源 |", "|---|---|---|---|---|",
          *[f"| {a} | {b} | {_md_cell(t)} | {_md_cell(w)} | {s} |" for a, b, t, w, s in c["phases"]], "",
          "## 四、分析判断", "", "### (一)有利因素与不利因素", "", "**有利因素**", "", *[f"{i}. **{k}。**{v}" for i, (k, v) in enumerate(c["pros"], 1)],
          "", "**不利因素与风险**", "", *[f"{i}. **{k}。**{v}" for i, (k, v) in enumerate(c["cons"], 1)], "",
          "### (二)资格要求对照", "", "| 项目要求 | 我方已满足 | 需联合体补充 |", "|---|---|---|",
          *[f"| **{k}:**{_md_cell(v)} |  |  |" for k, v in c["quals"]], "",
          "### (三)我方切入点与竞争", "", *[f"**{k}:**{v}\n" for k, v in c["entry"]],
          "## 五、计划动作", "", "| 序号 | 动作 | 责任方 | 时间 | 状态 |", "|---|---|---|---|---|",
          *[f"| {i} | {_md_cell(a)} | {o} | {d} | {s} |" for i, (a, o, d, s) in enumerate(c["actions"], 1)], "",
          "## 六、需要的支持", "", "### (一)请领导决策与支持", "", *[f"{i}. {x}" for i, x in enumerate(c["leader"], 1)], "",
          "### (二)请前方伙伴跟进", "", *[f"{i}. {x}" for i, x in enumerate(c["partner"], 1)], "",
          "## 七、信息来源", "", "| 编号 | 来源与网址 | 发布方 | 日期 | 等级 |", "|---|---|---|---|---|",
          *[f"| [{n}] | {_md_cell(r['title'])}" + (f"<br>{r['url']}" if r["url"] else "") + f" | {_md_cell(r['publisher'])} | {r['date']} | {r['grade']} |"
            for n, r in c["sources"]], "",
          f"编制说明:{c['note']}" + (f"挂到本项目的系统线索:{'、'.join(c['leads'])}。" if c["leads"] else ""), ""]
    return "\n".join(L)


# ---------- Word(在模板上填) ----------

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_CITE = re.compile(r"(\[\d+\])")


def _q(tag: str) -> str:
    return f"{{{W}}}{tag}"


def _base_rpr(par):
    """占位符那一段的字体设置(去掉斜体、灰色与上标),新写的字用它。"""
    runs = [r for r in par.runs if r.italic or (r.font.color is not None and r.font.color.type is not None)] or list(par.runs)[:1]
    if not runs or runs[0]._r.rPr is None:
        return None
    rpr = copy.deepcopy(runs[0]._r.rPr)
    for t in ("i", "iCs", "color", "vertAlign"):
        for el in rpr.findall(_q(t)):
            rpr.remove(el)
    return rpr


def _write(par, text: str, keep_labels: bool = True, cite: bool = True) -> None:
    """段落:保留开头的粗体标签(如"项目:"),其余(连同模板里的链接)换成 text;cite 时 [n] 写成上标(正文里),
    表格的"来源"列与编号不上标。"""
    rpr = _base_rpr(par)
    runs = list(par.runs)
    keep = 0
    if keep_labels:
        while keep < len(runs) and runs[keep].bold and not runs[keep].italic:
            keep += 1
    kept = {r._r for r in runs[:keep]}
    for el in list(par._p):
        if el.tag in (_q("r"), _q("hyperlink"), _q("fldSimple"), _q("smartTag")) and el not in kept:
            par._p.remove(el)
    for part in (_CITE.split(str(text)) if cite else [str(text)]):
        if not part:
            continue
        run = par.add_run(part)
        if rpr is not None:
            run._r.insert(0, copy.deepcopy(rpr))
        run.italic = False
        if cite and _CITE.fullmatch(part):
            run.font.superscript = True


def _clean(el):
    """复制出来的行、段落去掉书签(模板里的书签只能有一个)。"""
    for bm in el.iter(_q("bookmarkStart"), _q("bookmarkEnd")):
        bm.getparent().remove(bm)
    return el


def _set_lines(cell, lines: list[str]) -> None:
    """单元格写几行:沿用第一段的格式,多的行复制段落。"""
    pars = cell.paragraphs
    proto = _clean(copy.deepcopy(pars[0]._p))
    for p in pars[1:]:
        p._p.getparent().remove(p._p)
    first = cell.paragraphs[0]
    _write(first, lines[0] if lines else "—", keep_labels=False)
    anchor = first._p
    from docx.text.paragraph import Paragraph

    for line in lines[1:]:
        el = copy.deepcopy(proto)
        anchor.addnext(el)
        anchor = el
        _write(Paragraph(el, first._parent), line, keep_labels=False)


_TCPR_AFTER_SHD = ("noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark", "headers", "cellIns", "cellDel", "cellMerge",
                   "tcPrChange")


def _shade(cell, fill: str) -> None:
    """单元格底色;w:shd 按 OOXML 的顺序放(在 tcMar、vAlign 等之前),Word 才不报"内容有问题"。"""
    from docx.oxml import OxmlElement

    tcpr = cell._tc.get_or_add_tcPr()
    for old in tcpr.findall(_q("shd")):
        tcpr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(_q("val"), "clear")
    shd.set(_q("color"), "auto")
    shd.set(_q("fill"), fill)
    nxt = next((c for c in tcpr if c.tag in {_q(t) for t in _TCPR_AFTER_SHD}), None)
    if nxt is None:
        tcpr.append(shd)
    else:
        nxt.addprevious(shd)


def _fit_rows(table, n: int, proto_index: int = -1) -> None:
    """表格数据行调成 n 行(表头之外):多的从后面删,少的照 proto_index 那一行(先存下,去掉书签)复制。"""
    proto = _clean(copy.deepcopy(table.rows[proto_index]._tr))
    while len(table.rows) - 1 > max(n, 1):
        table.rows[-1]._tr.getparent().remove(table.rows[-1]._tr)
    while len(table.rows) - 1 < n:
        table.rows[-1]._tr.addnext(copy.deepcopy(proto))


def _fit_pars(pars: list, n: int) -> list:
    """一组连续的段落(如编号列表)调成 n 段。"""
    from docx.text.paragraph import Paragraph

    pars = list(pars)
    while len(pars) > n and len(pars) > 1:
        pars[-1]._p.getparent().remove(pars[-1]._p)
        pars.pop()
    while len(pars) < n:
        el = _clean(copy.deepcopy(pars[-1]._p))
        pars[-1]._p.addnext(el)
        pars.append(Paragraph(el, pars[-1]._parent))
    return pars


def _hyperlink(par, url: str, text: str) -> None:
    from docx.oxml import OxmlElement
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    rid = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    for tag, val in (("color", "1F4FD8"), ("sz", "16"), ("u", "single")):  # OOXML 顺序:color、sz、u
        el = OxmlElement(f"w:{tag}")
        el.set(_q("val"), val)
        rpr.append(el)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r.append(t)
    h.append(r)
    par._p.append(h)


ROLES = {"info": "首页信息表(国别、项目领域……)", "core": "核心结论框", "basics": "项目基本情况表", "phases": "项目进展表",
         "pros": "有利与不利因素表", "quals": "资格要求对照表", "actions": "计划动作表", "sources": "信息来源表"}


def _tables(doc, lang: str = "zh") -> dict:
    """按表头认出模板里的各张表(不按位置,模板删了"填写说明"框、挪了表也认得)。"""
    L = LABELS[lang]
    out: dict = {}
    for t in doc.tables:
        if not t.rows:
            continue
        head = [c.text.strip() for c in t.rows[0].cells]
        first = head[0] if head else ""
        role = ("info" if first == L["country"] else "howto" if first.startswith(L["howto"]) else
                "core" if all(x in first for x in L["core"]) else "basics" if head[:3] == L["basics"] else
                "phases" if head[:2] == L["phases"] else "pros" if first == L["pros"] else "quals" if first == L["quals"] else
                "actions" if head[:2] == L["actions"] else "sources" if head[:2] == L["sources"] else None)
        if role:
            out.setdefault(role, t)
    missing = [v for k, v in ROLES.items() if k not in out]
    if missing:
        raise ValueError(f"项目报告模板里找不到:{'、'.join(missing)}(模板的表头改过?)")
    return out


def to_docx(c: dict, out: Path, template: Path = TEMPLATE) -> Path:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(template))
    lang = c.get("lang", "zh")
    L = LABELS[lang]
    tb = _tables(doc, lang)
    body = [Paragraph(e, doc) for e in doc.element.body.iterchildren() if e.tag == _q("p")]
    pars = {p.text.strip(): p for p in body}

    def par_starting(prefix: str):
        hit = next((p for t, p in pars.items() if t.startswith(prefix)), None)
        if hit is None:
            raise ValueError(f"项目报告模板里找不到以「{prefix}」开头的段落(模板改过?)")
        return hit

    def title_text() -> str:
        return f"{c['country']}{c['name']}" if lang == "zh" else c["name"]

    # 封面(标题两行)与首页信息表
    title = next((p for p in body if any(x in p.text for x in L["title"])), body[0])
    title_en = next((p for p in body if L["title_en"] in p.text), body[1])
    _write(title, title_text(), keep_labels=False)
    _write(title_en, c["name_en"], keep_labels=False)
    info = tb["info"]
    for (r, col), key in {(0, 1): "country", (0, 3): "field", (1, 1): "owner", (1, 3): "funding", (2, 1): "doc_no", (2, 3): "as_of",
                          (3, 1): "org", (3, 3): "to"}.items():
        _write(info.cell(r, col).paragraphs[0], c[key], keep_labels=False)
    if "howto" in tb:  # 填写说明框(正式报送前删除;模板里已删掉时不管)
        tb["howto"]._tbl.getparent().remove(tb["howto"]._tbl)
    # 一、核心结论
    for par, (_, text) in zip(tb["core"].cell(0, 0).paragraphs, c["core"]):
        _write(par, text)
    # 二、项目基本情况(国别基本情况照这张表的空白样子复制一张,所以先复制再填)
    bt = tb["basics"]
    ct = _clean(copy.deepcopy(bt._tbl))
    if c.get("basics_labels"):  # 行名也换(英文版、或数据给出的行与模板不同时),行数跟着数据
        _fit_rows(bt, len(c["basics"]), proto_index=-1)
        for row, (k, _, _) in zip(bt.rows[1:], c["basics"]):
            _write(row.cells[0].paragraphs[0], k, keep_labels=False, cite=False)
    for row, (_, lines, s) in zip(bt.rows[1:], c["basics"]):
        _set_lines(row.cells[1], lines)
        _write(row.cells[2].paragraphs[0], s, keep_labels=False, cite=False)
    if c["country_rows"]:  # 前面一段说明
        note = _clean(copy.deepcopy(par_starting("背景")._p))
        bt._tbl.addnext(note)
        note.addnext(ct)
        np_ = Paragraph(note, doc)
        for r in np_.runs:
            if r.bold:
                r.text = "国别基本情况："
        _write(np_, f"{c['country']}的主要经济、人口与供排水指标(最新值,括号内为数据年份)。")
        ctab = Table(ct, doc)
        _write(ctab.cell(0, 0).paragraphs[0], "指标", keep_labels=False)
        _write(ctab.cell(0, 1).paragraphs[0], "数值", keep_labels=False)
        _fit_rows(ctab, len(c["country_rows"]), proto_index=2)
        for row, (a, b, s) in zip(ctab.rows[1:], c["country_rows"]):
            _set_lines(row.cells[0], [a])
            _set_lines(row.cells[1], [b])
            _write(row.cells[2].paragraphs[0], s or "—", keep_labels=False, cite=False)
    # 三、项目进展
    _write(par_starting(L["background"]), c["background"])
    pt = tb["phases"]
    _fit_rows(pt, len(c["phases"]), proto_index=2)
    for row, (ph, st, t, w, s) in zip(pt.rows[1:], c["phases"]):
        for cell, v in zip(row.cells[:4], (ph, st, t, w)):
            _set_lines(cell, [v])
        _write(row.cells[4].paragraphs[0], s, keep_labels=False, cite=False)
        for cell in row.cells[:2]:
            _shade(cell, SHADE.get(st, "FFFFFF"))
    # 四、分析判断
    ft = tb["pros"]
    for col, items in ((0, c["pros"]), (1, c["cons"])):
        for par, (_, text) in zip(ft.cell(1, col).paragraphs, items):
            _write(par, text)
    qt = tb["quals"]
    for row, (_, text) in zip(qt.rows[1:], c["quals"]):
        _write(row.cells[0].paragraphs[0], text)
    for (k, text), label in zip(c["entry"], L["entry"]):
        _write(par_starting(f"{label}{L['colon']}".strip()), text)
    # 五、计划动作
    at = tb["actions"]
    _fit_rows(at, len(c["actions"]), proto_index=1)
    for i, (row, (a, o, d, s)) in enumerate(zip(at.rows[1:], c["actions"]), 1):
        for cell, v in zip(row.cells, (str(i), a, o, d, s)):
            _write(cell.paragraphs[0], v, keep_labels=False)
    # 六、需要的支持
    heads = [Paragraph(e, doc) for e in doc.element.body.iterchildren() if e.tag == _q("p")]  # 插过国别表说明后重新取
    lists = []
    cur: list = []
    for p in heads:
        if p.style is not None and p.style.name == "List Paragraph":
            cur.append(p)
        elif cur:
            lists.append(cur)
            cur = []
    for grp, items in zip(lists[-2:] if lang == "zh" else lists[-1:], (c["leader"], c["partner"]) if lang == "zh" else (c["partner"],)):
        for par, text in zip(_fit_pars(grp, max(len(items), 1)), items or ["待研判"]):
            _write(par, text, keep_labels=False)
    # 七、信息来源
    st_ = tb["sources"]
    _fit_rows(st_, max(len(c["sources"]), 1), proto_index=-1)
    for row, (n, r) in zip(st_.rows[1:], c["sources"]):
        _write(row.cells[0].paragraphs[0], f"[{n}]", keep_labels=False, cite=False)
        cell = row.cells[1]
        _set_lines(cell, [r["title"]])
        if r["url"].lower().startswith(("https://", "http://")):  # 只有网页地址做成链接
            pnew = copy.deepcopy(cell.paragraphs[0]._p)
            for rr in pnew.findall(_q("r")):
                pnew.remove(rr)
            cell.paragraphs[0]._p.addnext(pnew)
            _hyperlink(Paragraph(pnew, cell.paragraphs[0]._parent), r["url"], r["url"])
        elif r["url"]:
            _set_lines(cell, [r["title"], r["url"]])
        for cell, v in zip(row.cells[2:], (r["publisher"], r["date"], r["grade"])):
            _write(cell.paragraphs[0], v, keep_labels=False)
    if not c["sources"]:
        for cell, v in zip(st_.rows[1].cells, ("—", "(没有可引用的公开来源)", "—", "—", "—")):
            _set_lines(cell, [v])
    # 编制说明:附在等级说明后
    last = par_starting(L["grades"])
    extra = _clean(copy.deepcopy(last._p))
    last._p.addnext(extra)
    _write(Paragraph(extra, doc), f"{L['note']}{c['note']}" + (f"{L['leads']}{'、'.join(c['leads'])}。" if c["leads"] and L["leads"] else ""),
           keep_labels=False)
    # 页眉:项目名称与报告语言
    for sec in doc.sections:
        for p in sec.header.paragraphs:
            for r in p.runs:
                if "项目名称" in r.text or "Project name" in r.text:
                    r.text, r.italic = c["name"], False
                    r.font.color.rgb = None
                elif "中文 / English" in r.text or r.text.strip() == "English":
                    r.text, r.italic = L["lang"], False
                    r.font.color.rgb = None
    doc.core_properties.title = f"{title_text()}{'项目分析报告' if lang == 'zh' else ' – Project Analysis Report'}"
    doc.core_properties.author = c["org"] if not c["org"].startswith("【") else doc.core_properties.author
    doc.core_properties.last_modified_by = "Water Tender Watch"
    doc.save(str(out))
    return out


def default_org(template: Path = TEMPLATE) -> str:
    """模板文件里的编制单位(文档属性"作者");没有模板或没写时空。"""
    if not Path(template).exists():
        return ""
    from docx import Document

    return str(Document(str(template)).core_properties.author or "")


def write(profile: Profile, store: Store, iso3: str, item_id: str, today: date, clf: Classifier | None = None, org: str = "",
          to: str = "公司领导", out_dir: Path | None = None, template: Path = TEMPLATE) -> list[Path]:
    """生成项目报告 → [md, docx],存在数据目录 reports/custom/(或 out_dir)。"""
    p = load_project(profile, store, iso3, item_id, today)
    c = compose(p, today, clf, org, to)
    out_dir = out_dir or store.root.parent / "reports" / "custom"
    out_dir.mkdir(parents=True, exist_ok=True)
    short = re.sub(r"[^\w一-鿿-]+", "_", c["name"]).strip("_")[:24]
    fp = out_dir / f"项目报告_{item_id}_{short}_{today}.md"
    docx = to_docx(c, fp.with_suffix(".docx"), template)  # Word 先成功,再写 Markdown(不留只有 Markdown 的半份报告)
    fp.write_text(to_markdown(c), encoding="utf-8")
    return [fp, docx]
